# -*- coding: utf-8 -*-
"""Điền cột Bác sĩ của sheet tieuphau từ EMR.

Quy tắc:
- Chỉ xử lý sheet tieuphau.
- Giữ nguyên ô đã là bí danh/họ tên của Bác sĩ hoặc BSNT đang hoạt động.
- Ô trống hoặc đang chứa Điều dưỡng/KTV/người không phải bác sĩ sẽ được kiểm tra lại.
- Tìm ở D/s Thủ thuật và D/s Phẫu thuật, ưu tiên khớp họ tên + thời gian.
- Sau khi mở chi tiết, lấy Bác sĩ chỉ định ở thủ thuật hoặc BS mổ chính ở phẫu thuật.
- Đối chiếu họ tên EMR với nhan_su_web.json và ghi bí danh vào Excel.
"""
from __future__ import annotations

import argparse
import json
import math
import os
import re
import sys
import time as time_module
import traceback
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime, time
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from emr_integration.emr_list_search import EmrListSearcher, load_config

try:
    from chuyen_thu_thuat_sang_tieuphau_t5 import HEADER_ALIASES
except Exception:
    HEADER_ALIASES = {}


APP_ROOT = Path(__file__).resolve().parent
STAFF_FILE = APP_ROOT / "nhan_su_web.json"
DEFAULT_CONFIG = APP_ROOT / "emr_config.json"
DOCTOR_ROLES = {"bac_si", "bsnt"}
DEBUG_ROOT = APP_ROOT / "emr_debug"


class DebugRecorder:
    """Ghi nhật ký EMR theo từng bước để dễ xác định nơi bị dừng.

    Nhật ký được ghi ngay xuống đĩa nên vẫn còn khi trình duyệt hoặc tiến trình
    gặp lỗi giữa chừng. Mỗi phiên có thư mục riêng gồm debug.log, events.jsonl,
    ảnh chụp màn hình và HTML tại thời điểm lỗi.
    """

    def __init__(self, root: Path, *, enabled: bool = True) -> None:
        self.enabled = bool(enabled)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")[:-3]
        self.folder = root / stamp
        self.log_path = self.folder / "debug.log"
        self.events_path = self.folder / "events.jsonl"
        if self.enabled:
            self.folder.mkdir(parents=True, exist_ok=True)
            self.log("START", "Bắt đầu phiên đối chiếu EMR", {"folder": str(self.folder)})

    @staticmethod
    def _safe_name(value: Any) -> str:
        text = norm(value) if 'norm' in globals() else str(value or '').lower()
        text = re.sub(r"[^a-z0-9]+", "_", text).strip("_")
        return (text or "debug")[:80]

    def log(self, event: str, message: str, data: Optional[Dict[str, Any]] = None) -> None:
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S.%f")[:-3]
        line = f"[{timestamp}] [{event}] {message}"
        print(line, file=sys.stderr, flush=True)
        if not self.enabled:
            return
        with self.log_path.open("a", encoding="utf-8") as fh:
            fh.write(line + "\n")
            if data:
                fh.write(json.dumps(data, ensure_ascii=False, default=str, indent=2) + "\n")
        payload = {
            "time": timestamp,
            "event": event,
            "message": message,
            "data": data or {},
        }
        with self.events_path.open("a", encoding="utf-8") as fh:
            fh.write(json.dumps(payload, ensure_ascii=False, default=str) + "\n")

    def dump_browser(self, client: Any, label: str, extra: Optional[Dict[str, Any]] = None) -> Dict[str, str]:
        if not self.enabled:
            return {}
        driver = getattr(client, "driver", None)
        if driver is None:
            return {}
        base = f"{datetime.now().strftime('%H%M%S_%f')[:-3]}_{self._safe_name(label)}"
        png_path = self.folder / f"{base}.png"
        html_path = self.folder / f"{base}.html"
        meta_path = self.folder / f"{base}.json"
        result: Dict[str, str] = {}
        meta: Dict[str, Any] = dict(extra or {})
        try:
            meta["current_url"] = str(driver.current_url or "")
        except Exception as exc:
            meta["current_url_error"] = str(exc)
        try:
            meta["title"] = str(driver.title or "")
        except Exception as exc:
            meta["title_error"] = str(exc)
        try:
            driver.save_screenshot(str(png_path))
            result["screenshot"] = str(png_path)
        except Exception as exc:
            meta["screenshot_error"] = str(exc)
        try:
            html_path.write_text(str(driver.page_source or ""), encoding="utf-8")
            result["html"] = str(html_path)
        except Exception as exc:
            meta["html_error"] = str(exc)
        meta_path.write_text(json.dumps(meta, ensure_ascii=False, default=str, indent=2), encoding="utf-8")
        result["meta"] = str(meta_path)
        self.log("DUMP", f"Đã lưu trạng thái trình duyệt: {label}", result)
        return result

    def finish(self, summary: Dict[str, Any]) -> None:
        if self.enabled:
            (self.folder / "summary.json").write_text(
                json.dumps(summary, ensure_ascii=False, default=str, indent=2),
                encoding="utf-8",
            )
        self.log("FINISH", "Kết thúc phiên đối chiếu EMR", {
            "updatedRows": summary.get("updatedRows"),
            "unmatchedRows": summary.get("unmatchedRows"),
        })


def norm(value: Any) -> str:
    text = unicodedata.normalize("NFD", str(value or ""))
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    text = text.replace("đ", "d").replace("Đ", "D").lower()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def norm_person(value: Any) -> str:
    text = norm(value)
    text = re.sub(r"\b(?:bs|bsi|bac si|ths|ts|cki|ckii|bsnt)\b", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def norm_patient_name(value: Any) -> str:
    """Chuẩn hóa tên NB lấy từ danh sách EMR.

    D/s Phẫu thuật có thể ghép chú thích ngay trong ô Họ tên, ví dụ
    ``NGHI THỊ ANH ĐÀO (Phẫu thuật tại khoa)``. Nếu so sánh tuyệt đối,
    chương trình nhìn thấy người bệnh nhưng vẫn bỏ qua dòng đó.
    """
    text = norm_person(value)
    annotation_phrases = (
        "phau thuat tai khoa",
        "thu thuat tai khoa",
        "phau thuat ngoai vien",
        "thu thuat ngoai vien",
    )
    for phrase in annotation_phrases:
        text = re.sub(rf"\b{re.escape(phrase)}\b", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def patient_names_match(expected: Any, actual: Any) -> bool:
    """Khớp tên NB nhưng cho phép chú thích phụ trong ô tên EMR."""
    target = norm_patient_name(expected)
    candidate = norm_patient_name(actual)
    if not target or not candidate:
        return False
    if target == candidate:
        return True

    target_tokens = target.split()
    candidate_tokens = candidate.split()
    # Tên đầy đủ của NB xuất hiện nguyên vẹn trong ô có thêm chú thích.
    if len(target_tokens) >= 2 and re.search(rf"(?:^| ){re.escape(target)}(?: |$)", candidate):
        return True
    if len(candidate_tokens) >= 2 and re.search(rf"(?:^| ){re.escape(candidate)}(?: |$)", target):
        return True

    # Dự phòng khi EMR chèn xuống dòng hoặc ký hiệu nhưng vẫn giữ đủ các từ tên.
    target_set = set(target_tokens)
    candidate_set = set(candidate_tokens)
    return len(target_tokens) >= 2 and target_set.issubset(candidate_set)


def row_matches_patient(row: Dict[str, Any], patient_name: str) -> bool:
    known = row.get("known_fields") or {}
    values: List[Any] = [known.get("ho_ten")]
    values.extend((link or {}).get("text") for link in (row.get("links") or []))
    return any(patient_names_match(patient_name, value) for value in values if value)


def header_key(value: Any) -> str:
    return norm(value).replace(" ", "")


# Tên gọi khác của cùng một cột (vd "Ngày chỉ định" thay cho "Ngày"), lấy
# chung từ cau_hinh_alias_cot.json qua HEADER_ALIASES để không phải khai báo
# lặp lại ở đây.
ALIAS_KEYS_BY_CANONICAL: Dict[str, List[str]] = {
    header_key(canonical): [header_key(name) for name in names]
    for canonical, names in HEADER_ALIASES.items()
}


def is_blank(value: Any) -> bool:
    return value is None or str(value).strip() == ""


def parse_datetime(value: Any, epoch=None) -> Optional[datetime]:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number = float(value)
        if math.isfinite(number) and 1 <= number <= 2958465:
            try:
                converted = from_excel(number, epoch=epoch)
                if isinstance(converted, datetime):
                    return converted
                if isinstance(converted, date):
                    return datetime.combine(converted, time.min)
            except Exception:
                pass
    text = str(value).strip()
    if not text:
        return None
    text = re.sub(r"\s+", " ", text)
    formats = (
        "%Y/%m/%d %H:%M", "%Y-%m-%d %H:%M", "%d/%m/%Y %H:%M",
        "%H:%M %d/%m/%Y", "%d-%m-%Y %H:%M", "%H:%M %d-%m-%Y",
        "%Y/%m/%d", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y",
    )
    for fmt in formats:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def minute_distance(a: Optional[datetime], b: Optional[datetime]) -> Optional[int]:
    if not a or not b:
        return None
    return int(abs((a - b).total_seconds()) // 60)


def same_day(a: Optional[datetime], b: Optional[datetime]) -> bool:
    return bool(a and b and a.date() == b.date())


def text_similarity(a: Any, b: Any) -> float:
    left = norm(a)
    right = norm(b)
    if not left or not right:
        return 0.0
    if left == right:
        return 1.0
    seq = SequenceMatcher(None, left, right).ratio()
    la, lb = set(left.split()), set(right.split())
    jac = len(la & lb) / max(1, len(la | lb))
    containment = 1.0 if left in right or right in left else 0.0
    return max(seq, jac, containment * 0.92)


def read_staff() -> Tuple[List[Dict[str, Any]], Dict[str, Dict[str, Any]], Dict[str, Dict[str, Any]]]:
    if not STAFF_FILE.is_file():
        raise FileNotFoundError(f"Không tìm thấy {STAFF_FILE.name}")
    payload = json.loads(STAFF_FILE.read_text(encoding="utf-8"))
    staff = payload.get("staff", payload if isinstance(payload, list) else [])
    if not isinstance(staff, list):
        raise ValueError("Danh sách nhân sự không hợp lệ")

    doctors: List[Dict[str, Any]] = []
    by_name: Dict[str, Dict[str, Any]] = {}
    by_alias: Dict[str, Dict[str, Any]] = {}
    for item in staff:
        if not isinstance(item, dict) or not item.get("active", True):
            continue
        if str(item.get("vaiTro") or "") not in DOCTOR_ROLES:
            continue
        name = str(item.get("hoTen") or "").strip()
        alias = str(item.get("biDanh") or "").strip()
        if not name or not alias:
            continue
        doctors.append(item)
        by_name[norm_person(name)] = item
        by_alias[norm(alias)] = item
    return doctors, by_name, by_alias


def map_doctor(name: str, doctors: List[Dict[str, Any]], by_name: Dict[str, Dict[str, Any]]) -> Optional[Dict[str, Any]]:
    key = norm_person(name)
    if not key:
        return None
    if key in by_name:
        return by_name[key]
    # Dự phòng cho trường hợp EMR có tiền tố chức danh hoặc viết tắt nhỏ.
    candidates = []
    cleaned_name = norm_person(name)
    for item in doctors:
        score = text_similarity(cleaned_name, norm_person(item.get("hoTen")))
        if score >= 0.90:
            candidates.append((score, item))
    candidates.sort(key=lambda x: x[0], reverse=True)
    if not candidates:
        return None
    if len(candidates) > 1 and candidates[0][0] - candidates[1][0] < 0.04:
        return None
    return candidates[0][1]


def existing_is_valid(value: Any, by_name: Dict[str, Dict[str, Any]], by_alias: Dict[str, Dict[str, Any]]) -> bool:
    alias_key = norm(value)
    name_key = norm_person(value)
    return bool((alias_key and alias_key in by_alias) or (name_key and name_key in by_name))


def find_sheet(wb, wanted: str):
    key = header_key(wanted)
    for name in wb.sheetnames:
        if header_key(name) == key:
            return wb[name]
    raise ValueError(f"Không tìm thấy sheet {wanted}")


def find_headers(ws) -> Tuple[int, Dict[str, int], Dict[int, str]]:
    required = {"ngay", "hovaten", "tencls", "bacsi"}
    max_col = ws.max_column  # ws.max_column quét lại cả sheet mỗi lần gọi, tính một lần ở đây.
    for row_idx in range(1, min(ws.max_row, 40) + 1):
        by_key: Dict[str, int] = {}
        by_col: Dict[int, str] = {}
        for col_idx in range(1, max_col + 1):
            raw = ws.cell(row_idx, col_idx).value
            key = header_key(raw)
            if key:
                by_key.setdefault(key, col_idx)
                by_col[col_idx] = str(raw).strip()

        # Bổ sung tên chuẩn (vd "ngay") trỏ về đúng cột nếu dòng này chỉ có
        # tên gọi khác (vd "ngaychidinh").
        for canonical_key, alias_keys in ALIAS_KEYS_BY_CANONICAL.items():
            if canonical_key in by_key:
                continue
            for alias_key in alias_keys:
                if alias_key in by_key:
                    by_key[canonical_key] = by_key[alias_key]
                    break

        if required.issubset(by_key):
            return row_idx, by_key, by_col
    raise ValueError("Không nhận diện được các cột Ngày, Họ và tên, Tên CLS, Bác sĩ")


def find_total_row(ws, header_row: int) -> int:
    max_col = min(ws.max_column, 12)  # tránh gọi lại ws.max_column trong vòng lặp theo dòng.
    for row_idx in range(header_row + 1, ws.max_row + 1):
        for col_idx in range(1, max_col + 1):
            if norm(ws.cell(row_idx, col_idx).value) == "tong cong":
                return row_idx
    return ws.max_row + 1




def atomic_save_workbook(wb, output_path: Path, *, retries: int = 3) -> None:
    """Lưu workbook an toàn sau từng dòng đã tìm được.

    File tạm được ghi cùng thư mục rồi thay thế file đích. Nếu người dùng đang
    mở file bằng Excel, thao tác thay thế có thể tạm thất bại nên thử lại vài lần.
    """
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = output_path.with_name(f".{output_path.stem}.autosave{output_path.suffix}")
    last_error: Optional[Exception] = None
    for attempt in range(max(1, retries)):
        try:
            if temp_path.exists():
                temp_path.unlink()
            wb.save(temp_path)
            os.replace(temp_path, output_path)
            return
        except Exception as exc:
            last_error = exc
            try:
                if temp_path.exists():
                    temp_path.unlink()
            except Exception:
                pass
            if attempt + 1 < retries:
                time_module.sleep(0.5 * (attempt + 1))
    raise RuntimeError(
        f"Không thể tự động lưu file {output_path.name}. "
        "Hãy đóng file này trong Excel rồi chạy lại."
    ) from last_error


def detail_url(row: Dict[str, Any], list_type: str) -> str:
    """Trả liên kết hoặc dấu hiệu có thể mở hồ sơ chi tiết.

    D/s Phẫu thuật mới dùng ``javascript:void(0)`` và
    ``onclick=checkBanGiaoPT(access_id)`` nên không được xem là thiếu liên kết.
    """
    needle = "thuthuatid=" if list_type == "procedure" else "phauthuatid="
    for link in row.get("links") or []:
        href = str((link or {}).get("href") or "")
        onclick = str((link or {}).get("onclick") or "")
        if needle in href.lower():
            return href
        if list_type == "surgery" and "checkbangiaopt" in onclick.lower():
            return f"javascript:{onclick}"
    access_id = str(row.get("access_id") or "").strip()
    if list_type == "surgery" and access_id:
        return f"access_id:{access_id}"
    return ""


@dataclass
class Candidate:
    source: str
    row: Dict[str, Any]
    list_datetime: Optional[datetime]
    list_method: str
    list_score: float
    list_time_source: str = ""
    list_method_source: str = ""
    detail: Optional[Dict[str, Any]] = None
    detail_datetime: Optional[datetime] = None
    method_score: float = 0.0
    total_score: float = 0.0
    staff: Optional[Dict[str, Any]] = None


def candidate_from_row(
    source: str,
    row: Dict[str, Any],
    patient_name: str,
    target_dt: Optional[datetime],
    target_method: str,
) -> Optional[Candidate]:
    known = row.get("known_fields") or {}
    row_name = str(known.get("ho_ten") or "")
    if not row_matches_patient(row, patient_name):
        return None

    if source == "procedure":
        # Sheet tieuphau lấy mốc Ngày từ thời gian CHỈ ĐỊNH. Trên D/s Thủ thuật
        # cột "Ngày thủ thuật" có thể là ngày thực hiện sau đó nhiều tuần, nên
        # bắt buộc ưu tiên "Thời gian chỉ định" để không loại nhầm đúng hồ sơ.
        dt_text = known.get("thoi_gian_chi_dinh") or known.get("ngay_thu_thuat")
        time_source = "Thời gian chỉ định" if known.get("thoi_gian_chi_dinh") else "Ngày thủ thuật"
        # Tên CLS phải đối chiếu đúng cột Chỉ định thủ thuật.
        method = str(known.get("chi_dinh_thu_thuat") or "")
        method_source = "Chỉ định thủ thuật"
    else:
        dt_text = known.get("thoi_gian")
        time_source = "Thời gian"
        # Với D/s Phẫu thuật phải đối chiếu đúng cột Nội dung phẫu thuật.
        method = str(known.get("noi_dung_phau_thuat") or "")
        method_source = "Nội dung phẫu thuật"
    row_dt = parse_datetime(dt_text)
    score = 20.0
    dist = minute_distance(target_dt, row_dt)
    if dist is not None:
        if dist == 0:
            score += 130
        elif dist <= 5:
            score += 110
        elif dist <= 30:
            score += 70
        elif same_day(target_dt, row_dt):
            score += 45
    elif same_day(target_dt, row_dt):
        score += 45
    method_score = text_similarity(target_method, method)
    score += method_score * 55
    status = norm(known.get("trang_thai"))
    if "hoan tat" in status:
        score += 5
    if not detail_url(row, source):
        score -= 100
    return Candidate(
        source, row, row_dt, method, score,
        list_time_source=time_source,
        list_method_source=method_source,
    )


def evaluate_detail(
    candidate: Candidate,
    detail: Dict[str, Any],
    target_dt: Optional[datetime],
    target_method: str,
    doctors: List[Dict[str, Any]],
    by_name: Dict[str, Dict[str, Any]],
) -> Candidate:
    candidate.detail = detail
    candidate.detail_datetime = parse_datetime(detail.get("start_time")) or candidate.list_datetime
    candidate.method_score = text_similarity(target_method, detail.get("method") or candidate.list_method)
    candidate.staff = map_doctor(str(detail.get("doctor") or ""), doctors, by_name)

    score = candidate.list_score
    dist = minute_distance(target_dt, candidate.detail_datetime)
    if dist is not None:
        if dist == 0:
            score += 260
        elif dist <= 5:
            score += 220
        elif dist <= 15:
            score += 160
        elif same_day(target_dt, candidate.detail_datetime):
            score += 65
        else:
            score -= 120
    elif same_day(target_dt, candidate.detail_datetime):
        score += 65
    score += candidate.method_score * 110
    if candidate.staff:
        score += 80
    else:
        score -= 140
    candidate.total_score = score
    return candidate


def confidence_ok(candidate: Candidate, target_dt: Optional[datetime]) -> bool:
    if not candidate.staff:
        return False
    dist = minute_distance(target_dt, candidate.detail_datetime)
    if dist is not None and dist <= 15:
        return True
    return same_day(target_dt, candidate.detail_datetime) and candidate.method_score >= 0.82


def source_order(method: str) -> List[str]:
    return ["surgery", "procedure"] if "phau thuat" in norm(method) else ["procedure", "surgery"]


def _debug_row_summary(row: Dict[str, Any]) -> Dict[str, Any]:
    known = row.get("known_fields") or {}
    return {
        "access_id": row.get("access_id"),
        "known_fields": known,
        "raw_text": row.get("raw_text"),
        "links": [
            {
                "text": (link or {}).get("text"),
                "href": (link or {}).get("href"),
                "onclick": (link or {}).get("onclick"),
            }
            for link in (row.get("links") or [])
        ],
    }


def find_doctor_for_row(
    client: EmrListSearcher,
    patient_name: str,
    target_dt: Optional[datetime],
    target_method: str,
    doctors: List[Dict[str, Any]],
    by_name: Dict[str, Dict[str, Any]],
    *,
    debug: Optional[DebugRecorder] = None,
    excel_row: Optional[int] = None,
) -> Tuple[Optional[Candidate], str]:
    reasons: List[str] = []
    row_label = f"dòng {excel_row}" if excel_row else "dòng chưa xác định"

    for source in source_order(target_method):
        source_label = "D/s Phẫu thuật" if source == "surgery" else "D/s Thủ thuật"
        if debug:
            debug.log("SEARCH_BEGIN", f"{row_label}: tìm {patient_name} tại {source_label}", {
                "patient": patient_name,
                "target_time": target_dt.isoformat(sep=" ") if target_dt else "",
                "target_method": target_method,
                "source": source,
            })
        try:
            result = client.search(source, patient_name, months=3).to_dict()
            if debug:
                debug.log("SEARCH_RESULT", f"{source_label}: nhận {len(result.get('rows') or [])} dòng", {
                    "current_url": result.get("current_url"),
                    "search_input_id": result.get("search_input_id"),
                    "search_trigger": result.get("search_trigger"),
                    "headers": result.get("table_headers"),
                    "rows": [_debug_row_summary(x) for x in (result.get("rows") or [])],
                })
        except Exception as exc:
            message = f"{source}: lỗi tìm kiếm {exc}"
            reasons.append(message)
            if debug:
                debug.log("SEARCH_ERROR", message, {"traceback": traceback.format_exc()})
                debug.dump_browser(client, f"{excel_row}_{source}_search_error")
            continue

        candidates: List[Candidate] = []
        for row in result.get("rows") or []:
            item = candidate_from_row(source, row, patient_name, target_dt, target_method)
            if item:
                candidates.append(item)
        candidates.sort(key=lambda c: c.list_score, reverse=True)
        if debug:
            debug.log("CANDIDATES", f"{source_label}: có {len(candidates)} ứng viên cùng người bệnh", {
                "candidates": [
                    {
                        "list_score": c.list_score,
                        "list_datetime": c.list_datetime.isoformat(sep=" ") if c.list_datetime else "",
                        "list_time_source": c.list_time_source,
                        "list_method": c.list_method,
                        "list_method_source": c.list_method_source,
                        "method_score": text_similarity(target_method, c.list_method),
                        "detail_marker": detail_url(c.row, source),
                        "row": _debug_row_summary(c.row),
                    }
                    for c in candidates
                ]
            })
        if not candidates:
            reasons.append(f"{source}: không có dòng cùng người bệnh")
            continue

        same_date_candidates = [c for c in candidates if same_day(target_dt, c.list_datetime)]
        source_candidates = same_date_candidates or candidates
        evaluated: List[Candidate] = []

        for candidate_index, candidate in enumerate(source_candidates[:3], start=1):
            marker = detail_url(candidate.row, source)
            if debug:
                debug.log("DETAIL_BEGIN", f"{source_label}: mở ứng viên {candidate_index}", {
                    "detail_marker": marker,
                    "list_datetime": candidate.list_datetime.isoformat(sep=" ") if candidate.list_datetime else "",
                    "list_method": candidate.list_method,
                    "row": _debug_row_summary(candidate.row),
                })
            try:
                detail = client.open_detail(source, candidate.row)
                evaluated_candidate = evaluate_detail(
                    candidate, detail, target_dt, target_method, doctors, by_name
                )
                evaluated.append(evaluated_candidate)
                if debug:
                    debug.log("DETAIL_OK", f"{source_label}: đọc được hồ sơ chi tiết", {
                        "doctor_from_emr": detail.get("doctor"),
                        "mapped_alias": (evaluated_candidate.staff or {}).get("biDanh"),
                        "mapped_name": (evaluated_candidate.staff or {}).get("hoTen"),
                        "start_time": detail.get("start_time"),
                        "start_time_field": detail.get("start_time_field"),
                        "method": detail.get("method"),
                        "method_field": detail.get("method_field"),
                        "list_time_source": evaluated_candidate.list_time_source,
                        "list_method_source": evaluated_candidate.list_method_source,
                        "current_url": detail.get("current_url"),
                        "total_score": evaluated_candidate.total_score,
                        "method_score": evaluated_candidate.method_score,
                    })
            except Exception as exc:
                message = f"{source}: không đọc được chi tiết ({exc})"
                reasons.append(message)
                if debug:
                    debug.log("DETAIL_ERROR", message, {
                        "detail_marker": marker,
                        "row": _debug_row_summary(candidate.row),
                        "traceback": traceback.format_exc(),
                    })
                    debug.dump_browser(client, f"{excel_row}_{source}_detail_error", {
                        "patient": patient_name,
                        "detail_marker": marker,
                    })

        evaluated.sort(key=lambda c: c.total_score, reverse=True)
        if debug:
            debug.log("EVALUATED", f"{source_label}: đánh giá {len(evaluated)} hồ sơ", {
                "items": [
                    {
                        "doctor": (c.detail or {}).get("doctor"),
                        "alias": (c.staff or {}).get("biDanh"),
                        "detail_datetime": c.detail_datetime.isoformat(sep=" ") if c.detail_datetime else "",
                        "method_score": c.method_score,
                        "total_score": c.total_score,
                        "confidence_ok": confidence_ok(c, target_dt),
                    }
                    for c in evaluated
                ]
            })

        good = [c for c in evaluated if confidence_ok(c, target_dt)]
        if good:
            if len(good) == 1:
                if debug:
                    debug.log("SOURCE_MATCH", f"{source_label}: đã chọn bác sĩ, không tìm nguồn còn lại", {
                        "alias": (good[0].staff or {}).get("biDanh")
                    })
                return good[0], ""
            alias0 = norm(good[0].staff.get("biDanh")) if good[0].staff else ""
            alias1 = norm(good[1].staff.get("biDanh")) if good[1].staff else ""
            if alias0 and alias0 == alias1:
                return good[0], ""
            if good[0].total_score - good[1].total_score >= 35:
                return good[0], ""
            return None, f"{source}: có nhiều hồ sơ khớp nhưng bác sĩ khác nhau"

        mapped = [c for c in evaluated if c.staff]
        mapped.sort(key=lambda c: c.total_score, reverse=True)
        if same_date_candidates and mapped:
            if len(mapped) == 1:
                if debug:
                    debug.log("SOURCE_MATCH_DATE", f"{source_label}: chọn hồ sơ cùng ngày và dừng", {
                        "alias": (mapped[0].staff or {}).get("biDanh")
                    })
                return mapped[0], ""
            alias0 = norm(mapped[0].staff.get("biDanh")) if mapped[0].staff else ""
            alias1 = norm(mapped[1].staff.get("biDanh")) if mapped[1].staff else ""
            if alias0 and alias0 == alias1:
                return mapped[0], ""
            if mapped[0].total_score - mapped[1].total_score >= 35:
                return mapped[0], ""
            return None, f"{source}: có nhiều ca cùng ngày nhưng bác sĩ khác nhau"

        if same_date_candidates:
            reason = "; ".join(reasons) or f"{source}: đã có ca cùng ngày nhưng chưa lấy được bác sĩ"
            if debug:
                debug.log("STOP_AT_SOURCE", f"{source_label}: đã thấy ca cùng ngày nên không sang nguồn khác", {
                    "reason": reason
                })
            return None, reason

        reasons.append(f"{source}: có kết quả nhưng không khớp ngày")

    return None, "; ".join(reasons) or "Không tìm được hồ sơ khớp thời gian"

def iter_target_rows(ws, header_row: int, total_row: int, cols: Dict[str, int], only_row: Optional[int]) -> Iterable[int]:
    if only_row is not None:
        if only_row <= header_row or only_row >= total_row:
            raise ValueError("Dòng được chọn không nằm trong vùng dữ liệu tieuphau")
        yield only_row
        return
    for row_idx in range(header_row + 1, total_row):
        if all(is_blank(ws.cell(row_idx, cols[key]).value) for key in ("ngay", "hovaten", "tencls")):
            continue
        yield row_idx


def run(args: argparse.Namespace) -> Dict[str, Any]:
    input_path = Path(args.file).expanduser().resolve()
    output_path = Path(args.output).expanduser().resolve()
    config_path = Path(args.config or DEFAULT_CONFIG).expanduser().resolve()
    debug_root = Path(args.debug_dir).expanduser().resolve() if args.debug_dir else DEBUG_ROOT
    debug = DebugRecorder(debug_root, enabled=not bool(args.no_debug))
    debug.log("CONFIG", "Thông tin phiên chạy", {
        "input_file": str(input_path),
        "output_file": str(output_path),
        "config_file": str(config_path),
        "only_row": args.row,
        "headless_arg": args.headless,
    })

    if not input_path.is_file():
        raise FileNotFoundError(f"Không tìm thấy file Excel: {input_path}")
    if not config_path.is_file():
        raise FileNotFoundError("Chưa có cấu hình EMR")

    config = load_config(config_path)
    if not str(config.get("username") or "").strip() or not str(config.get("password") or "").strip():
        raise ValueError("Chưa cấu hình tài khoản EMR")

    keep_vba = input_path.suffix.lower() == ".xlsm"
    wb = load_workbook(input_path, keep_vba=keep_vba, data_only=False)
    ws = find_sheet(wb, "tieuphau")
    header_row, cols, headers = find_headers(ws)
    total_row = find_total_row(ws, header_row)
    doctors, by_name, by_alias = read_staff()
    debug.log("WORKBOOK", "Đã đọc sheet tieuphau", {
        "sheet": ws.title,
        "header_row": header_row,
        "total_row": total_row,
        "columns": cols,
        "headers": headers,
        "doctor_count": len(doctors),
    })

    rows_to_check: List[int] = []
    skipped_valid = 0
    for row_idx in iter_target_rows(ws, header_row, total_row, cols, args.row):
        current = ws.cell(row_idx, cols["bacsi"]).value
        if existing_is_valid(current, by_name, by_alias):
            skipped_valid += 1
            debug.log("SKIP_VALID", f"Dòng {row_idx}: đã là bác sĩ hợp lệ", {"value": current})
            continue
        rows_to_check.append(row_idx)

    debug.log("TARGET_ROWS", f"Có {len(rows_to_check)} dòng cần kiểm tra EMR", {
        "rows": rows_to_check,
        "skipped_valid": skipped_valid,
    })

    output_path.parent.mkdir(parents=True, exist_ok=True)
    atomic_save_workbook(wb, output_path)
    debug.log("AUTOSAVE_INIT", "Đã tạo/cập nhật file làm việc trước khi tìm EMR", {
        "output": str(output_path)
    })

    results: List[Dict[str, Any]] = []
    updated = 0
    unmatched = 0
    client: Optional[EmrListSearcher] = None
    try:
        if rows_to_check:
            headless = bool(config.get("headless", True)) if args.headless is None else bool(args.headless)
            client = EmrListSearcher(config, headless=headless, log=lambda message: debug.log("BROWSER", message))
            debug.log("BROWSER_START", "Khởi động trình duyệt EMR", {"headless": headless})
            client.start()
            debug.log("BROWSER_READY", "Trình duyệt EMR đã sẵn sàng")

            for position, row_idx in enumerate(rows_to_check, start=1):
                patient = str(ws.cell(row_idx, cols["hovaten"]).value or "").strip()
                raw_time = ws.cell(row_idx, cols["ngay"]).value
                target_dt = parse_datetime(raw_time, wb.epoch)
                method = str(ws.cell(row_idx, cols["tencls"]).value or "").strip()
                current = str(ws.cell(row_idx, cols["bacsi"]).value or "").strip()
                record: Dict[str, Any] = {
                    "rowNumber": row_idx,
                    "patient": patient,
                    "time": target_dt.strftime("%H:%M %d/%m/%Y") if target_dt else "",
                    "method": method,
                    "oldValue": current,
                }
                debug.log("ROW_BEGIN", f"[{position}/{len(rows_to_check)}] Bắt đầu dòng {row_idx}: {patient}", {
                    "raw_time": raw_time,
                    "parsed_time": record["time"],
                    "method": method,
                    "old_value": current,
                })

                if not patient or not target_dt:
                    reason = "Thiếu họ tên hoặc thời gian hợp lệ"
                    record.update({"status": "unmatched", "reason": reason})
                    unmatched += 1
                    results.append(record)
                    debug.log("ROW_SKIP", f"Dòng {row_idx}: {reason}")
                    continue

                try:
                    candidate, reason = find_doctor_for_row(
                        client, patient, target_dt, method, doctors, by_name,
                        debug=debug, excel_row=row_idx,
                    )
                    if candidate and candidate.staff:
                        alias = str(candidate.staff.get("biDanh") or "").strip()
                        ws.cell(row_idx, cols["bacsi"]).value = alias
                        record.update({
                            "status": "updated",
                            "newValue": alias,
                            "doctorName": candidate.staff.get("hoTen"),
                            "emrDoctor": (candidate.detail or {}).get("doctor"),
                            "source": "D/s Thủ thuật" if candidate.source == "procedure" else "D/s Phẫu thuật",
                            "matchedTime": candidate.detail_datetime.strftime("%H:%M %d/%m/%Y") if candidate.detail_datetime else "",
                            "matchedTimeSource": (candidate.detail or {}).get("start_time_field") or candidate.list_time_source,
                            "matchedMethod": (candidate.detail or {}).get("method") or candidate.list_method,
                            "matchedMethodSource": (candidate.detail or {}).get("method_field") or candidate.list_method_source,
                        })
                        atomic_save_workbook(wb, output_path)
                        updated += 1
                        debug.log("AUTOSAVE_ROW", f"Dòng {row_idx}: đã điền {alias} và lưu ngay", {
                            "patient": patient,
                            "source": record["source"],
                            "doctor_name": record["doctorName"],
                            "emr_doctor": record["emrDoctor"],
                            "matched_time": record["matchedTime"],
                            "matched_method": record["matchedMethod"],
                            "progress": f"{updated}/{len(rows_to_check)}",
                            "output": str(output_path),
                        })
                    else:
                        final_reason = reason or "Không tìm được bác sĩ phù hợp"
                        record.update({"status": "unmatched", "reason": final_reason})
                        unmatched += 1
                        debug.log("ROW_UNMATCHED", f"Dòng {row_idx}: {final_reason}", record)
                except Exception as exc:
                    record.update({"status": "unmatched", "reason": str(exc)})
                    unmatched += 1
                    debug.log("ROW_ERROR", f"Dòng {row_idx}: lỗi chưa xử lý {exc}", {
                        "record": record,
                        "traceback": traceback.format_exc(),
                    })
                    debug.dump_browser(client, f"{row_idx}_unexpected_error", {"patient": patient})
                results.append(record)
                debug.log("ROW_END", f"Kết thúc dòng {row_idx}", {
                    "status": record.get("status"),
                    "new_value": record.get("newValue"),
                    "reason": record.get("reason"),
                })
    finally:
        if client is not None:
            try:
                client.close()
                debug.log("BROWSER_CLOSE", "Đã đóng trình duyệt EMR")
            except Exception as exc:
                debug.log("BROWSER_CLOSE_ERROR", f"Không đóng được trình duyệt: {exc}")

    atomic_save_workbook(wb, output_path)
    summary = {
        "ok": True,
        "file": str(output_path),
        "sheet": ws.title,
        "message": f"Đã kiểm tra {len(rows_to_check)} dòng cần đối chiếu EMR.",
        "checkedRows": len(rows_to_check),
        "updatedRows": updated,
        "skippedValidRows": skipped_valid,
        "unmatchedRows": unmatched,
        "debugFolder": str(debug.folder) if debug.enabled else "",
        "debugLog": str(debug.log_path) if debug.enabled else "",
        "debugEvents": str(debug.events_path) if debug.enabled else "",
        "results": results,
        "unmatched": [item for item in results if item.get("status") != "updated"],
        "updated": [item for item in results if item.get("status") == "updated"],
    }
    debug.finish(summary)
    return summary

def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Điền bác sĩ sheet tieuphau từ EMR")
    parser.add_argument("--file", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--config", default=str(DEFAULT_CONFIG))
    parser.add_argument("--row", type=int)
    parser.add_argument("--debug-dir", default=str(DEBUG_ROOT))
    parser.add_argument("--no-debug", action="store_true")
    group = parser.add_mutually_exclusive_group()
    group.add_argument("--headless", dest="headless", action="store_true")
    group.add_argument("--show-browser", dest="headless", action="store_false")
    parser.set_defaults(headless=None)
    return parser


def main() -> int:
    try:
        data = run(build_parser().parse_args())
        print(json.dumps(data, ensure_ascii=False), flush=True)
        return 0
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False), flush=True)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
