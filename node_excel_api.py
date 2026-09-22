# -*- coding: utf-8 -*-
"""
API Python phục vụ web app Node.js nhập liệu Excel.
Node.js gọi file này bằng child_process, dữ liệu trao đổi qua JSON stdout.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
import uuid
from copy import copy
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel


# Chỉ các cột nhập liệu chính mới được dùng để xác định "Thiếu dữ liệu".
# Các cột công thức, số tiền và cột phụ không bắt buộc sẽ không làm cả dòng bị
# đánh dấu thiếu. Có thể chỉnh danh sách này nếu mẫu Excel thay đổi.
REQUIRED_HEADER_KEYS_BY_SHEET = {
    "phauthuat": {
        "ngay",
        "hovatenbenhnhan",
        "tuoi",
        "chandoanvaphuongphapphauthuat",
        "loaiphauthuat",
        "ptvchinh",
    },
    "thuthuat": {
        "ngay",
        "hovaten",
        "tuoi",
        "tencls",
        "soluong",
    },
    "tieuphau": {
        "ngay",
        "hovaten",
        "tuoi",
        "tencls",
        "soluong",
        "bacsi",
        "dieuduong",
    },
}

# Sheet phẫu thuật chỉ được phép cập nhật bốn cột nhân sự này. Đây là lớp
# bảo vệ ở phía máy chủ, nên ngay cả khi trình duyệt gửi nhầm cột khác thì
# dữ liệu bệnh nhân/kỹ thuật trong Excel vẫn không bị thay đổi.
SURGERY_EDITABLE_HEADER_KEYS = {"ptvchinh", "phumo1", "phumo2", "phumo3"}

HELPER_IMPORT_WARNING = ""

try:
    from chuyen_thu_thuat_sang_tieuphau_t5 import (
        copy_row_style,
        find_header_row_and_columns as find_header_row_required,
        find_sheet_case_insensitive,
        find_total_row,
        insert_rows_preserve_merges,
        delete_rows_preserve_merges,
        normalize_header,
        normalize_text,
        update_total_formulas_thu_thuat,
        update_total_formulas_tieuphau,
        update_formula_lien_ket_tong_cong,
        translate_formula,
        HEADER_ALIASES,
    )
except Exception as exc:
    HELPER_IMPORT_WARNING = (
        "Không nạp được các hàm hỗ trợ giữ định dạng/công thức từ "
        f"chuyen_thu_thuat_sang_tieuphau_t5.py: {exc}"
    )
    HEADER_ALIASES = {}
    # Fallback tối thiểu nếu import lỗi.
    def normalize_text(value: Any) -> str:
        if value is None:
            return ""
        text = str(value).strip().lower()
        text = unicodedata.normalize("NFD", text)
        text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
        text = text.replace("đ", "d")
        text = re.sub(r"[^a-z0-9]+", " ", text)
        return re.sub(r"\s+", " ", text).strip()

    def normalize_header(value: Any) -> str:
        return normalize_text(value).replace(" ", "")

    def find_sheet_case_insensitive(wb, wanted_name: str) -> str:
        key = normalize_header(wanted_name)
        for name in wb.sheetnames:
            if normalize_header(name) == key:
                return name
        raise ValueError(f"Không tìm thấy sheet {wanted_name}")

    def find_total_row(ws) -> int:
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
            for cell in row:
                if normalize_text(cell.value) == "tong cong":
                    return cell.row
        return ws.max_row + 1

    def copy_row_style(ws, src_row: int, dst_row: int, max_col: int) -> None:
        pass

    def insert_rows_preserve_merges(ws, insert_at: int, amount: int) -> None:
        ws.insert_rows(insert_at, amount)

    def delete_rows_preserve_merges(ws, rows_to_delete: list[int]) -> None:
        for row_idx in sorted(rows_to_delete, reverse=True):
            ws.delete_rows(row_idx, 1)

    def update_total_formulas_thu_thuat(ws, header_row: int, total_row: int) -> None:
        pass

    def update_total_formulas_tieuphau(ws, header_row: int, total_row: int) -> None:
        pass

    def update_formula_lien_ket_tong_cong(ws_tieu, total_row_tieu: int, total_row_thu: int) -> None:
        pass

    def translate_formula(formula: Any, src_cell, dst_cell) -> Any:
        return formula


def json_default(value: Any):
    if isinstance(value, datetime):
        if value.time() == time.min:
            return value.strftime("%d/%m/%Y")
        return value.strftime("%d/%m/%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    return str(value)


def respond(data: Any) -> None:
    print(json.dumps(data, ensure_ascii=False, default=json_default))


def fail(message: str, code: int = 1) -> None:
    respond({"ok": False, "error": message})
    raise SystemExit(code)


def value_is_blank(value: Any) -> bool:
    if value is None:
        return True
    text = str(value).strip()
    return text == "" or text.lower() in {"nan", "none"}


def cell_value_to_display(value: Any, header: str | None = None, epoch=None) -> Any:
    """Đổi giá trị Excel sang dạng hiển thị ổn định cho giao diện.

    Với cột ngày/giờ, hàm luôn thử nhận diện số seri Excel (ví dụ 46182)
    ngay cả khi ô bị mất định dạng ngày. Nhờ vậy file cũ vẫn hiển thị đúng
    trong lần mở đầu tiên, trước cả khi workbook được lưu chuẩn hóa lại.
    """
    if _is_date_header(header):
        normalized, _kind = _coerce_excel_date(value, epoch)
        if normalized is not None:
            if normalized.time() == time.min:
                return normalized.strftime("%d/%m/%Y")
            return normalized.strftime("%d/%m/%Y %H:%M")

    if isinstance(value, datetime):
        if value.time() == time.min:
            return value.strftime("%d/%m/%Y")
        return value.strftime("%d/%m/%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    return value


def parse_date_text(text: str) -> datetime | str:
    text = text.strip()
    for fmt in ["%d/%m/%Y %H:%M", "%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%Y-%m-%d %H:%M"]:
        try:
            return datetime.strptime(text, fmt)
        except Exception:
            pass
    return text


def coerce_value(header: str, text_value: Any, old_value: Any = None) -> Any:
    if text_value is None:
        return None
    if isinstance(text_value, str):
        text = text_value.strip()
    else:
        return text_value
    if text == "":
        return None

    header_key = normalize_header(header)
    if isinstance(old_value, (datetime, date)) or any(k in header_key for k in ["ngay", "thoigian"]):
        parsed = parse_date_text(text)
        return parsed

    numeric_headers = {
        "stt", "tuoi", "soluong", "thanhtien", "sotien", "thuclanh", "tong", "dongia"
    }
    if header_key in numeric_headers or any(k in header_key for k in ["tien", "soluong"]):
        clean = text.replace(".", "").replace(",", ".")
        try:
            number = float(clean)
            if number.is_integer():
                return int(number)
            return number
        except Exception:
            return text

    return text


def find_header_row_generic(ws) -> tuple[int, list[str], dict[str, int]]:
    # ws.max_column quét lại toàn bộ sheet mỗi lần gọi (không cache nội bộ),
    # nên chỉ tính một lần ở đây thay vì gọi lại trong từng vòng lặp dòng/cột.
    max_col = ws.max_column
    best = None
    for row_idx in range(1, min(ws.max_row, 50) + 1):
        values = [ws.cell(row_idx, col).value for col in range(1, max_col + 1)]
        nonblank = [(i + 1, str(v).strip()) for i, v in enumerate(values) if not value_is_blank(v)]
        if len(nonblank) < 3:
            continue
        keys = {normalize_header(v) for _, v in nonblank}
        score = len(nonblank)
        if "stt" in keys:
            score += 20
        if any(k in keys for k in ["hovaten", "hovatenbenhnhan"]):
            score += 10
        if any(k in keys for k in ["tencls", "chandoanvaphuongphapphauthuat"]):
            score += 10
        if best is None or score > best[0]:
            best = (score, row_idx, nonblank)
    if not best:
        raise ValueError(f"Không tìm thấy dòng tiêu đề trong sheet '{ws.title}'.")

    header_row = best[1]
    headers: list[str] = []
    header_map: dict[str, int] = {}
    used: dict[str, int] = {}
    for col in range(1, max_col + 1):
        raw = ws.cell(header_row, col).value
        if value_is_blank(raw):
            continue
        name = str(raw).strip()
        base = name
        key = normalize_header(name)
        if key in used:
            used[key] += 1
            name = f"{base}_{used[key]}"
        else:
            used[key] = 1
        headers.append(name)
        header_map[name] = col
    return header_row, headers, header_map


def header_span_columns(ws, header_row: int, col_idx: int) -> list[int]:
    """Các cột dữ liệu nằm dưới một ô tiêu đề gộp.

    Mẫu phẫu thuật gộp tiêu đề ``Tuổi`` qua hai cột D:E và dữ liệu thực tế có
    thể nằm ở một trong hai cột. Đọc theo cả vùng gộp giúp không báo thiếu sai.
    """
    for merged in ws.merged_cells.ranges:
        if merged.min_row <= header_row <= merged.max_row and merged.min_col <= col_idx <= merged.max_col:
            return list(range(merged.min_col, merged.max_col + 1))
    return [col_idx]


def data_columns_for_header(
    ws, header_row: int, col_idx: int, header: str | None = None, max_col: int | None = None
) -> list[int]:
    columns = header_span_columns(ws, header_row, col_idx)
    # Một số mẫu phẫu thuật để tiêu đề Tuổi ở cột D nhưng dữ liệu có thể nằm
    # ở D hoặc E (E không có tiêu đề). Xem hai cột như một trường duy nhất.
    if normalize_header(header or "") == "tuoi" and len(columns) == 1:
        next_col = col_idx + 1
        # ws.max_column quét lại toàn bộ sheet mỗi lần gọi (không cache), nên
        # bên gọi trong vòng lặp theo dòng cần tự tính một lần và truyền vào
        # qua max_col để tránh lặp lại hàng trăm nghìn lần trên sheet lớn.
        effective_max_col = ws.max_column if max_col is None else max_col
        if next_col <= effective_max_col and value_is_blank(ws.cell(header_row, next_col).value):
            columns.append(next_col)
    return columns


def read_header_value(
    ws, header_row: int, row_idx: int, col_idx: int, header: str | None = None, max_col: int | None = None
) -> Any:
    columns = data_columns_for_header(ws, header_row, col_idx, header, max_col)
    values = [ws.cell(row_idx, col).value for col in columns]
    epoch = getattr(ws.parent, "epoch", None)
    for value in values:
        if not value_is_blank(value):
            return cell_value_to_display(value, header=header, epoch=epoch)
    return cell_value_to_display(values[0] if values else None, header=header, epoch=epoch)


def target_column_for_header(
    ws, header_row: int, row_idx: int, col_idx: int, header: str | None = None, max_col: int | None = None
) -> int:
    """Chọn cột ghi phù hợp, ưu tiên cột đang có dữ liệu."""
    columns = data_columns_for_header(ws, header_row, col_idx, header, max_col)
    for col in columns:
        if not value_is_blank(ws.cell(row_idx, col).value):
            return col
    return columns[0]


def load_sheet(file_path: Path, sheet_name: str | None):
    if not file_path.exists():
        raise FileNotFoundError(f"Không tìm thấy file: {file_path}")
    wb = load_workbook(file_path, keep_vba=file_path.suffix.lower() == ".xlsm")
    if sheet_name:
        real_name = find_sheet_case_insensitive(wb, sheet_name)
    else:
        real_name = wb.sheetnames[0]
    return wb, wb[real_name], real_name


def _required_header_alias_keys(aliases: dict[str, list[str]]) -> dict[str, set[str]]:
    """Suy ra bảng alias (key chuẩn hóa -> tập key thay thế) từ HEADER_ALIASES
    (nạp từ cau_hinh_alias_cot.json), dùng chung với
    chuyen_thu_thuat_sang_tieuphau_t5.py và public/app.js để chỉ cần sửa một
    chỗ khi gặp mẫu file mới."""
    result: dict[str, set[str]] = {}
    for canonical_name, alias_names in aliases.items():
        canonical_key = normalize_header(canonical_name)
        alias_keys = {normalize_header(name) for name in alias_names if normalize_header(name) != canonical_key}
        if alias_keys:
            result[canonical_key] = alias_keys
    return result


# Một số file "BẢNG KÊ TIỀN TIỂU PHẪU" đặt tên cột khác nhưng cùng ý nghĩa
# (vd "Ngày chỉ định" thay vì "Ngày"). Dùng thêm để không bỏ sót các cột này
# khi kiểm tra dòng thiếu dữ liệu.
REQUIRED_HEADER_ALIAS_KEYS = _required_header_alias_keys(HEADER_ALIASES)


def required_headers_for_sheet(sheet_name: str, headers: list[str]) -> list[str]:
    """Trả về đúng tên header thực tế cần kiểm tra bắt buộc cho từng sheet."""
    configured_keys = REQUIRED_HEADER_KEYS_BY_SHEET.get(normalize_header(sheet_name), set())
    if configured_keys:
        expanded_keys = set(configured_keys)
        for key in configured_keys:
            expanded_keys |= REQUIRED_HEADER_ALIAS_KEYS.get(key, set())
        return [header for header in headers if normalize_header(header) in expanded_keys]

    # Sheet chưa cấu hình: chỉ suy luận các trường nhận diện cơ bản, không coi
    # toàn bộ hàng trăm cột là bắt buộc.
    common_keys = {
        "ngay",
        "thoigian",
        "hovaten",
        "hovatenbenhnhan",
        "tuoi",
        "tencls",
        "chandoanvaphuongphapphauthuat",
    }
    return [header for header in headers if normalize_header(header) in common_keys]



STAFF_CONFIG_FILE = Path(__file__).resolve().with_name("nhan_su_web.json")
CLS_CONFIG_FILE = Path(__file__).resolve().with_name("ten_cls_tieuphau.json")
VALID_STAFF_ROLES = {"bac_si", "bsnt", "dieu_duong", "ktv"}


def _default_staff_items() -> list[dict[str, Any]]:
    try:
        from danh_sach_nhan_su import BAC_SI, DIEU_DUONG, KTV, BSNT, BI_DANH_NHAN_SU
    except Exception as exc:
        raise ValueError(f"Không đọc được danh sách nhân sự: {exc}") from exc

    items: list[dict[str, Any]] = []

    def add_item(full_name: str, alias: str, group: str, role: str):
        full_name = str(full_name or "").strip()
        alias = str(alias or "").strip()
        if not alias and full_name in BI_DANH_NHAN_SU and BI_DANH_NHAN_SU[full_name]:
            alias = str(BI_DANH_NHAN_SU[full_name][0]).strip()
        if not alias:
            alias = full_name
        if not full_name or not alias:
            return
        items.append({
            "id": uuid.uuid4().hex,
            "hoTen": full_name,
            "biDanh": alias,
            "nhom": group,
            "vaiTro": role,
            "active": True,
        })

    for name in BAC_SI:
        add_item(name, "", "Bác sĩ", "bac_si")
    for item in BSNT:
        add_item(item.get("ho_ten"), item.get("bi_danh"), item.get("nhom", "BSNT"), "bsnt")
    for name in DIEU_DUONG:
        add_item(name, "", "Điều dưỡng", "dieu_duong")
    for name in KTV:
        add_item(name, "", "KTV", "ktv")
    return items


def _normalize_staff_item(item: dict[str, Any], index: int = 0) -> dict[str, Any]:
    if not isinstance(item, dict):
        raise ValueError(f"Nhân sự thứ {index + 1} không đúng định dạng.")
    full_name = str(item.get("hoTen") or item.get("ho_ten") or "").strip()
    alias = str(item.get("biDanh") or item.get("bi_danh") or "").strip()
    role = str(item.get("vaiTro") or item.get("vai_tro") or "").strip()
    group = str(item.get("nhom") or "").strip()
    if not full_name:
        raise ValueError(f"Nhân sự thứ {index + 1} chưa có họ tên.")
    if not alias:
        raise ValueError(f"Nhân sự '{full_name}' chưa có bí danh.")
    if role not in VALID_STAFF_ROLES:
        raise ValueError(f"Vai trò của '{full_name}' không hợp lệ.")
    if not group:
        group = {
            "bac_si": "Bác sĩ",
            "bsnt": "BSNT",
            "dieu_duong": "Điều dưỡng",
            "ktv": "KTV",
        }[role]
    return {
        "id": str(item.get("id") or uuid.uuid4().hex),
        "hoTen": full_name,
        "biDanh": alias,
        "nhom": group,
        "vaiTro": role,
        "active": bool(item.get("active", True)),
        "label": f"{alias} — {full_name} ({group})",
    }


def _load_staff_config() -> list[dict[str, Any]]:
    if not STAFF_CONFIG_FILE.exists():
        return [_normalize_staff_item(item, idx) for idx, item in enumerate(_default_staff_items())]
    try:
        payload = json.loads(STAFF_CONFIG_FILE.read_text(encoding="utf-8"))
        rows = payload.get("staff", []) if isinstance(payload, dict) else []
        return [_normalize_staff_item(item, idx) for idx, item in enumerate(rows)]
    except Exception as exc:
        raise ValueError(f"Không đọc được file quản lý nhân sự: {exc}") from exc


def command_staff_list(args):
    items = _load_staff_config()
    role_counts = {role: 0 for role in VALID_STAFF_ROLES}
    active_count = 0
    for item in items:
        role_counts[item["vaiTro"]] += 1
        if item.get("active", True):
            active_count += 1
    respond({
        "ok": True,
        "staff": items,
        "count": len(items),
        "activeCount": active_count,
        "roleCounts": role_counts,
        "configFile": str(STAFF_CONFIG_FILE),
    })


def _role_family(vai_tro: str) -> str:
    return "bac_si_bsnt" if vai_tro in {"bac_si", "bsnt"} else vai_tro


def _chu_cai_ten_dem(ho_ten: str) -> str:
    """Chữ cái đầu của tên đệm (từ ngay trước tên chính), dùng để phân biệt
    hai người trùng tên — vd 'Phạm Việt Tân' -> 'V'. Trả về rỗng nếu họ tên
    chỉ có 1 từ, không đủ để lấy tên đệm."""
    words = str(ho_ten or "").strip().split()
    return words[-2][0].upper() if len(words) >= 2 else ""


def _tu_dong_tach_biet_trung_ten(items: list[dict[str, Any]]) -> list[str]:
    """Hai bác sĩ/BSNT khác người nhưng trùng tên 100% (cùng bí danh y hệt
    nhau, không phải chỉ gần giống) sẽ tự được thêm chữ cái đầu tên đệm để
    phân biệt, vd hai người cùng tên 'Tân' thành 'V.TÂN' và 'M.TÂN' — đúng
    theo cách đặt bí danh đã dùng sẵn trong danh sách (vd 'M.THIỆN').
    Nếu không đủ tên đệm để tách rõ ràng thì bỏ qua, để bước kiểm tra trùng
    bí danh phía sau báo lỗi cho người dùng tự xử lý.

    Trả về danh sách mô tả các thay đổi đã tự làm, để báo lại cho người dùng."""
    groups: dict[tuple[str, str], list[dict[str, Any]]] = {}
    for item in items:
        if not item.get("active", True):
            continue
        groups.setdefault((_role_family(item["vaiTro"]), item["biDanh"]), []).append(item)

    ghi_chu: list[str] = []
    for (_role_family_key, alias), group in groups.items():
        if len(group) < 2:
            continue

        prefixes = [_chu_cai_ten_dem(item["hoTen"]) for item in group]
        if not all(prefixes) or len(set(prefixes)) != len(prefixes):
            continue

        for item, prefix in zip(group, prefixes):
            new_alias = f"{prefix}.{alias}"
            ghi_chu.append(f"{item['hoTen']}: '{alias}' -> '{new_alias}' (trùng tên với người khác)")
            item["biDanh"] = new_alias

    return ghi_chu


def command_save_staff_list(args):
    payload = json.loads(args.staff_json or "[]")
    if not isinstance(payload, list):
        raise ValueError("Danh sách nhân sự không đúng định dạng.")
    normalized = [_normalize_staff_item(item, idx) for idx, item in enumerate(payload)]

    auto_adjusted = _tu_dong_tach_biet_trung_ten(normalized)
    for item in normalized:
        item["label"] = f"{item['biDanh']} — {item['hoTen']} ({item['nhom']})"

    # Bí danh chỉ khác dấu (vd "Thanh" so với "Thạnh") vẫn được coi là hai
    # người khác nhau và cho lưu bình thường, không chặn. Chỉ khi bí danh
    # giống hệt nhau từng ký tự (kể cả dấu) mới cần xử lý — bước
    # _tu_dong_tach_biet_trung_ten ở trên đã tự thêm tên đệm để phân biệt.
    # Nếu vẫn còn trùng y hệt sau khi đã thử tách (không đủ tên đệm khác
    # nhau) thì mới chặn, vì lúc đó không có cách nào phân biệt hai người
    # trên màn hình gợi ý nhập liệu.
    seen: dict[tuple[str, str], str] = {}
    for item in normalized:
        if not item.get("active", True):
            continue
        key = (_role_family(item["vaiTro"]), item["biDanh"])
        if key in seen:
            raise ValueError(
                f"Bí danh '{item['biDanh']}' bị trùng y hệt giữa '{seen[key]}' và '{item['hoTen']}' — không đủ "
                f"tên đệm khác nhau để tự phân biệt. Hãy đặt bí danh khác cho một trong hai (vd thêm tên đệm "
                f"đầy đủ hơn)."
            )
        seen[key] = item["hoTen"]

    data_to_write = {
        "version": 1,
        "updatedAt": datetime.now().isoformat(timespec="seconds"),
        "staff": [
            {
                "id": item["id"],
                "hoTen": item["hoTen"],
                "biDanh": item["biDanh"],
                "nhom": item["nhom"],
                "vaiTro": item["vaiTro"],
                "active": item["active"],
            }
            for item in normalized
        ],
    }
    STAFF_CONFIG_FILE.parent.mkdir(parents=True, exist_ok=True)
    tmp = STAFF_CONFIG_FILE.with_suffix(STAFF_CONFIG_FILE.suffix + ".tmp")
    tmp.write_text(json.dumps(data_to_write, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(STAFF_CONFIG_FILE)
    message = "Đã lưu danh sách nhân sự."
    if auto_adjusted:
        message += " Đã tự thêm tên đệm để phân biệt: " + "; ".join(auto_adjusted) + "."
    respond({
        "ok": True,
        "staff": normalized,
        "count": len(normalized),
        "activeCount": sum(1 for item in normalized if item.get("active", True)),
        "file": str(STAFF_CONFIG_FILE),
        "message": message,
        "autoAdjusted": auto_adjusted,
    })




def _default_cls_items() -> list[dict[str, Any]]:
    names = [
        "Hút ổ viêm/áp xe phần mềm",
        "Tiêm khớp gối",
        "Khâu vết thương phần mềm dài dưới 10 cm [tổn thương nông]",
        "Khâu vết thương phần mềm nông dài < 5cm",
        "Tiêm gân gấp ngón tay",
        "Chích rạch nhọt, Apxe nhỏ dẫn lưu",
        "Tiêm gân duỗi ngón tay",
        "Tiêm gân lồi cầu cánh tay trong/ngoài",
        "Mắt cá 1-2 thương tổn",
        "Khâu vết thương phần mềm sâu dài > 5 cm",
        "Tiêm gân Dequervein",
        "Khâu vết thương phần mềm nông dài > 5 cm",
        "Hút dịch khớp gối",
        "Điều trị thoái hóa khớp bằng huyết tương giàu tiểu cầu",
        "Khâu vết thương phần mềm sâu dài < 5cm",
        "Tiêm khớp vai",
        "Mắt cá 3-4 thương tổn",
        "Khâu vết thương phần mềm dài dưới 10 cm [tổn thương sâu]",
    ]
    return [
        {"id": f"cls_{idx + 1:03d}", "tenCls": name, "active": True}
        for idx, name in enumerate(names)
    ]


def _normalize_cls_item(item: dict[str, Any], index: int = 0) -> dict[str, Any]:
    if isinstance(item, str):
        item = {"tenCls": item}
    if not isinstance(item, dict):
        raise ValueError(f"Tên CLS thứ {index + 1} không đúng định dạng.")
    name = str(item.get("tenCls") or item.get("ten_cls") or item.get("name") or "").strip()
    if not name:
        raise ValueError(f"Tên CLS thứ {index + 1} đang để trống.")
    return {
        "id": str(item.get("id") or uuid.uuid4().hex),
        "tenCls": name,
        "active": bool(item.get("active", True)),
    }


def _load_cls_config() -> list[dict[str, Any]]:
    if not CLS_CONFIG_FILE.exists():
        return [_normalize_cls_item(item, idx) for idx, item in enumerate(_default_cls_items())]
    try:
        payload = json.loads(CLS_CONFIG_FILE.read_text(encoding="utf-8"))
        rows = payload.get("cls", []) if isinstance(payload, dict) else payload
        if not isinstance(rows, list):
            raise ValueError("Trường cls phải là một danh sách.")
        return [_normalize_cls_item(item, idx) for idx, item in enumerate(rows)]
    except Exception as exc:
        raise ValueError(f"Không đọc được file quản lý tên CLS: {exc}") from exc


def command_cls_list(args):
    items = _load_cls_config()
    respond({
        "ok": True,
        "cls": items,
        "count": len(items),
        "activeCount": sum(1 for item in items if item.get("active", True)),
        "configFile": str(CLS_CONFIG_FILE),
    })


def command_save_cls_list(args):
    payload = json.loads(args.cls_json or "[]")
    if not isinstance(payload, list):
        raise ValueError("Danh sách tên CLS không đúng định dạng.")
    normalized = [_normalize_cls_item(item, idx) for idx, item in enumerate(payload)]
    seen: dict[str, str] = {}
    for item in normalized:
        key = normalize_text(item["tenCls"])
        if key in seen:
            raise ValueError(f"Tên CLS '{item['tenCls']}' bị trùng với '{seen[key]}'.")
        seen[key] = item["tenCls"]

    data_to_write = {
        "version": 1,
        "updatedAt": datetime.now().isoformat(timespec="seconds"),
        "cls": normalized,
    }
    CLS_CONFIG_FILE.parent.mkdir(parents=True, exist_ok=True)
    tmp = CLS_CONFIG_FILE.with_suffix(CLS_CONFIG_FILE.suffix + ".tmp")
    tmp.write_text(json.dumps(data_to_write, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(CLS_CONFIG_FILE)
    respond({
        "ok": True,
        "cls": normalized,
        "count": len(normalized),
        "activeCount": sum(1 for item in normalized if item.get("active", True)),
        "file": str(CLS_CONFIG_FILE),
        "message": "Đã lưu danh mục tên CLS chuyển sang tiểu phẫu.",
    })


def command_list_files(args):
    root = Path(args.root or Path(__file__).resolve().parent)
    files = []
    workbook_paths = list(root.rglob("*.xlsx")) + list(root.rglob("*.xlsm"))
    for path in workbook_paths:
        if "__pycache__" in path.parts:
            continue
        name_key = normalize_text(path.name)
        # Ẩn các file báo cáo và file output cũ để danh sách file đỡ rối.
        if name_key.startswith("bao cao") or " bao cao" in name_key:
            continue
        if name_key.startswith("thong ke phan loai"):
            continue
        if "bao cao" in name_key or normalize_header(path.stem).endswith("baocao"):
            continue
        if "_da_" in path.stem or "%20" in path.name:
            continue
        try:
            rel = path.relative_to(root)
        except Exception:
            rel = path
        files.append({
            "name": path.name,
            "path": str(path),
            "relativePath": str(rel),
            "size": path.stat().st_size,
            "modified": datetime.fromtimestamp(path.stat().st_mtime),
        })
    files.sort(key=lambda item: item["modified"], reverse=True)
    respond({"ok": True, "files": files[:200]})




def _is_date_header(header: Any) -> bool:
    """Chỉ nhận diện các cột ngày/giờ thật, tránh nhầm với ``Số ngày``."""
    key = normalize_header(header)
    if not key or key.startswith("so"):
        return False
    return key == "ngay" or key.startswith("ngay") or key.startswith("thoigian")


def _parse_date_string(value: str) -> datetime | None:
    text = str(value or "").strip()
    if not text or text.startswith("="):
        return None

    # Một số file nguồn lưu số seri Excel dưới dạng chuỗi, ví dụ ``46182``.
    if re.fullmatch(r"\d+(?:[.,]\d+)?", text):
        try:
            number = float(text.replace(",", "."))
            converted = from_excel(number)
            if isinstance(converted, datetime):
                return converted
            if isinstance(converted, date):
                return datetime.combine(converted, time.min)
        except Exception:
            return None

    formats = (
        "%d/%m/%Y %H:%M:%S",
        "%d/%m/%Y %H:%M",
        "%d/%m/%Y",
        "%d-%m-%Y %H:%M:%S",
        "%d-%m-%Y %H:%M",
        "%d-%m-%Y",
        "%d.%m.%Y",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%d/%m/%y",
        "%d-%m-%y",
    )
    for fmt in formats:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def _coerce_excel_date(value: Any, epoch) -> tuple[datetime | None, str | None]:
    """Đổi giá trị ngày lỗi thành ``datetime`` và cho biết loại đã chuyển."""
    if value_is_blank(value) or isinstance(value, bool):
        return None, None

    converted: datetime | None = None
    conversion_type: str | None = None

    if isinstance(value, datetime):
        converted = value
        conversion_type = "format"
    elif isinstance(value, date):
        converted = datetime.combine(value, time.min)
        conversion_type = "format"
    elif isinstance(value, (int, float)):
        try:
            result = from_excel(value, epoch=epoch)
            if isinstance(result, datetime):
                converted = result
            elif isinstance(result, date):
                converted = datetime.combine(result, time.min)
            if converted is not None:
                conversion_type = "serial"
        except Exception:
            return None, None
    elif isinstance(value, str):
        stripped = value.strip()
        if stripped.startswith("="):
            return None, None
        # Dùng epoch của workbook cho số seri dạng chuỗi.
        if re.fullmatch(r"\d+(?:[.,]\d+)?", stripped):
            try:
                result = from_excel(float(stripped.replace(",", ".")), epoch=epoch)
                if isinstance(result, datetime):
                    converted = result
                elif isinstance(result, date):
                    converted = datetime.combine(result, time.min)
                if converted is not None:
                    conversion_type = "serial_text"
            except Exception:
                return None, None
        else:
            converted = _parse_date_string(stripped)
            if converted is not None:
                conversion_type = "text"

    if converted is None:
        return None, None

    # Chỉ tự sửa các ngày hợp lý. Điều này ngăn số lượng hoặc mã số bị đổi nhầm.
    if not (date(1990, 1, 1) <= converted.date() <= date(2100, 12, 31)):
        return None, None
    return converted, conversion_type


def _normalize_date_cells_in_sheet(wb, ws) -> dict[str, Any]:
    """Chuẩn hóa ngày/giờ của một sheet ngay trong workbook đang mở."""
    try:
        header_row, headers, header_map = find_header_row_generic(ws)
    except Exception:
        return {"changed": 0, "converted": 0, "formatted": 0, "headers": {}}

    date_headers = [header for header in headers if _is_date_header(header)]
    if not date_headers:
        return {"changed": 0, "converted": 0, "formatted": 0, "headers": {}}

    try:
        total_row = find_total_row(ws)
    except Exception:
        total_row = ws.max_row + 1
    last_row = min(total_row - 1, ws.max_row)

    converted_count = 0
    formatted_count = 0
    header_details: dict[str, int] = {}

    for header in date_headers:
        base_col = header_map[header]
        columns = data_columns_for_header(ws, header_row, base_col, header)
        header_changed = 0
        for row_idx in range(header_row + 1, last_row + 1):
            for col_idx in columns:
                cell = ws.cell(row_idx, col_idx)
                original = cell.value
                if value_is_blank(original) or (
                    isinstance(original, str) and original.strip().startswith("=")
                ):
                    continue

                normalized, _kind = _coerce_excel_date(original, wb.epoch)
                if normalized is None:
                    continue

                has_time = normalized.time() != time.min or normalize_header(header).startswith("thoigian")
                desired_format = "dd/mm/yyyy hh:mm" if has_time else "dd/mm/yyyy"
                value_changed = not isinstance(original, (datetime, date))
                format_changed = cell.number_format != desired_format

                if value_changed:
                    cell.value = normalized
                    converted_count += 1
                    header_changed += 1
                if format_changed:
                    cell.number_format = desired_format
                    formatted_count += 1
                    if not value_changed:
                        header_changed += 1

        if header_changed:
            header_details[header] = header_changed

    return {
        "changed": converted_count + formatted_count,
        "converted": converted_count,
        "formatted": formatted_count,
        "headers": header_details,
    }


def _is_working_file(file_path: Path) -> bool:
    """Chỉ tự ghi đè file làm việc, không làm thay đổi file nguồn tháng gốc."""
    stem_key = normalize_header(file_path.stem)
    return (
        "nhaplieu" in stem_key
        or file_path.parent.name.lower() == "node_outputs"
    )


def command_normalize_dates(args):
    """Chuẩn hóa các cột ngày/giờ và lưu trực tiếp vào file làm việc.

    File gốc upload luôn được giữ nguyên. Các số seri Excel như ``46182`` và
    ngày dạng chữ như ``8/6/2026`` được đổi thành ngày Excel thật để giao diện
    hiển thị thống nhất và các bước nhập liệu sau đó tiếp tục ghi đúng kiểu dữ liệu.
    """
    input_file = Path(args.file)
    if not input_file.exists():
        raise FileNotFoundError(f"Không tìm thấy file: {input_file}")
    output_file = output_path_from_args(input_file, args.output, "_NHAP_LIEU")

    wb = load_workbook(input_file, keep_vba=input_file.suffix.lower() == ".xlsm")
    converted_count = 0
    formatted_count = 0
    skipped_invalid_count = 0
    sheet_results: list[dict[str, Any]] = []

    for ws in wb.worksheets:
        try:
            header_row, headers, header_map = find_header_row_generic(ws)
        except Exception:
            continue

        date_headers = [header for header in headers if _is_date_header(header)]
        if not date_headers:
            continue

        try:
            total_row = find_total_row(ws)
        except Exception:
            total_row = ws.max_row + 1
        last_row = min(total_row - 1, ws.max_row)

        per_sheet_converted = 0
        per_sheet_formatted = 0
        header_details: dict[str, int] = {}

        for header in date_headers:
            base_col = header_map[header]
            columns = data_columns_for_header(ws, header_row, base_col, header)
            header_changed = 0
            for row_idx in range(header_row + 1, last_row + 1):
                for col_idx in columns:
                    cell = ws.cell(row_idx, col_idx)
                    original = cell.value
                    if value_is_blank(original) or (isinstance(original, str) and original.strip().startswith("=")):
                        continue
                    normalized, kind = _coerce_excel_date(original, wb.epoch)
                    if normalized is None:
                        # Chỉ đếm giá trị số có vẻ là ngày nhưng nằm ngoài phạm vi an toàn.
                        if isinstance(original, (int, float)) or (
                            isinstance(original, str) and re.fullmatch(r"\d+(?:[.,]\d+)?", original.strip())
                        ):
                            skipped_invalid_count += 1
                        continue

                    has_time = normalized.time() != time.min or normalize_header(header).startswith("thoigian")
                    desired_format = "dd/mm/yyyy hh:mm" if has_time else "dd/mm/yyyy"
                    value_changed = not isinstance(original, (datetime, date))
                    format_changed = cell.number_format != desired_format

                    cell.value = normalized
                    cell.number_format = desired_format
                    if value_changed:
                        converted_count += 1
                        per_sheet_converted += 1
                        header_changed += 1
                    elif format_changed:
                        formatted_count += 1
                        per_sheet_formatted += 1
                        header_changed += 1

            if header_changed:
                header_details[header] = header_changed

        if per_sheet_converted or per_sheet_formatted:
            sheet_results.append({
                "sheet": ws.title,
                "converted": per_sheet_converted,
                "formatted": per_sheet_formatted,
                "headers": header_details,
            })

    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_file)
    respond({
        "ok": True,
        "file": str(output_file),
        "sourceFile": str(input_file),
        "convertedCount": converted_count,
        "formattedCount": formatted_count,
        "fixedCount": converted_count + formatted_count,
        "skippedInvalidCount": skipped_invalid_count,
        "sheets": sheet_results,
        "message": (
            f"Đã chuẩn hóa {converted_count + formatted_count} ô ngày/giờ."
            if converted_count + formatted_count
            else "Không phát hiện ô ngày/giờ cần sửa."
        ),
    })


def command_workbook_info(args):
    file_path = Path(args.file)
    wb = load_workbook(file_path, read_only=True, data_only=False, keep_vba=file_path.suffix.lower() == ".xlsm")
    respond({
        "ok": True,
        "file": str(file_path),
        "sheets": wb.sheetnames,
    })


def command_read_sheet(args):
    file_path = Path(args.file)
    wb, ws, real_name = load_sheet(file_path, args.sheet)

    # Tự sửa cả các file làm việc được tạo từ phiên bản cũ. Đây là lớp bảo vệ
    # thứ hai ngoài bước chuẩn hóa khi upload: chỉ cần mở sheet, số seri như
    # 46182 sẽ được đổi thành ngày thật và lưu ngay vào file _NHAP_LIEU.
    date_fix = _normalize_date_cells_in_sheet(wb, ws)
    if date_fix["changed"] and _is_working_file(file_path):
        wb.save(file_path)

    header_row, headers, header_map = find_header_row_generic(ws)

    try:
        total_row = find_total_row(ws)
    except Exception:
        total_row = ws.max_row + 1

    max_rows = int(args.limit or 0)
    query = normalize_text(args.query or "")
    missing_only = bool(args.missing_only)
    required_headers = required_headers_for_sheet(real_name, headers)
    max_col = ws.max_column

    rows = []
    matched_count = 0
    complete_count = 0
    blank_rows = 0
    template_rows = 0
    missing_by_header = {header: 0 for header in required_headers}

    for row_idx in range(header_row + 1, min(total_row, ws.max_row + 1)):
        values = {}
        haystack_parts = []
        has_any = False
        has_non_formula_value = False
        has_required_value = False
        missing = []
        for header in headers:
            col = header_map[header]
            val = read_header_value(ws, header_row, row_idx, col, header, max_col)
            values[header] = val
            if not value_is_blank(val):
                has_any = True
                haystack_parts.append(str(val))
                if not (isinstance(val, str) and val.strip().startswith("=")):
                    has_non_formula_value = True
                if header in required_headers:
                    has_required_value = True
            elif header in required_headers:
                missing.append(header)

        if not has_any:
            blank_rows += 1
            continue
        # Các dòng mẫu chỉ có công thức nhưng chưa có dữ liệu nhập thật.
        if required_headers and not has_required_value:
            template_rows += 1
            continue
        if not has_non_formula_value:
            template_rows += 1
            continue
        if query and query not in normalize_text(" ".join(haystack_parts)):
            continue
        if missing_only and not missing:
            continue

        matched_count += 1
        if not missing:
            complete_count += 1
        for header in missing:
            missing_by_header[header] = missing_by_header.get(header, 0) + 1
        if max_rows <= 0 or len(rows) < max_rows:
            rows.append({"rowNumber": row_idx, "values": values, "missing": missing})

    data_slots = max(0, min(total_row, ws.max_row + 1) - header_row - 1)
    respond({
        "ok": True,
        "file": str(file_path),
        "sheet": real_name,
        "headerRow": header_row,
        "totalRow": total_row,
        "headers": headers,
        "requiredHeaders": required_headers,
        "rows": rows,
        "rowCount": len(rows),
        "matchedRowCount": matched_count,
        "completeCount": complete_count,
        "missingCount": max(0, matched_count - complete_count),
        "missingByHeader": missing_by_header,
        "dataSlots": data_slots,
        "blankRows": blank_rows,
        "templateRows": template_rows,
        "truncated": max_rows > 0 and matched_count > len(rows),
        "maxRow": ws.max_row,
        "maxColumn": ws.max_column,
        "dateFixCount": date_fix["changed"],
        "dateFixConverted": date_fix["converted"],
        "dateFixFormatted": date_fix["formatted"],
    })


def output_path_from_args(input_file: Path, output: str | None, suffix: str | None) -> Path:
    if output:
        return Path(output)
    safe_suffix = suffix or "_NHAP_LIEU"
    if not safe_suffix.startswith("_"):
        safe_suffix = "_" + safe_suffix
    return input_file.with_name(input_file.stem + safe_suffix + input_file.suffix)


def report_path_for(output_file: Path) -> Path:
    # Chỉ dùng một file báo cáo phụ, ghi đè lại mỗi lần chạy để tránh rác file output.
    return output_file.with_name(output_file.stem + "_BAO_CAO.xlsx")


def update_totals_for_sheet(wb, ws, header_row: int) -> list[str]:
    warnings: list[str] = [HELPER_IMPORT_WARNING] if HELPER_IMPORT_WARNING else []
    try:
        total_row = find_total_row(ws)
    except Exception as exc:
        return [f"Không xác định được dòng Tổng Cộng của sheet {ws.title}: {exc}"]
    sheet_key = normalize_header(ws.title)
    try:
        if sheet_key == normalize_header("thu thuat"):
            update_total_formulas_thu_thuat(ws, header_row, total_row)
        elif sheet_key == normalize_header("tieuphau"):
            update_total_formulas_tieuphau(ws, header_row, total_row)
            # Nếu có sheet thu thuat thì cập nhật công thức liên kết tổng cộng.
            for other in wb.sheetnames:
                if normalize_header(other) == normalize_header("thu thuat"):
                    total_thu = find_total_row(wb[other])
                    update_formula_lien_ket_tong_cong(ws, total_row, total_thu)
                    break
    except Exception as exc:
        warnings.append(f"Không cập nhật được công thức tổng của sheet {ws.title}: {exc}")
    return warnings


def restrict_row_data_for_sheet(sheet_name: str, data: dict[str, Any]) -> tuple[dict[str, Any], list[str]]:
    """Lọc các cột được phép ghi theo loại sheet.

    Với sheet phẫu thuật, chỉ PTV chính và Phụ mổ 1–3 được giữ lại. Các cột
    khác bị bỏ qua để bảo vệ dữ liệu gốc của bệnh nhân và kỹ thuật.
    """
    if "phauthuat" not in normalize_header(sheet_name):
        return dict(data or {}), []

    allowed: dict[str, Any] = {}
    blocked: list[str] = []
    for header, value in (data or {}).items():
        if normalize_header(header) in SURGERY_EDITABLE_HEADER_KEYS:
            allowed[header] = value
        else:
            blocked.append(str(header))
    return allowed, blocked


def apply_row_changes(ws, header_map: dict[str, int], row_number: int, data: dict[str, Any], header_row: int | None = None) -> list[str]:
    """Áp dụng dữ liệu một dòng và trả về danh sách header không tồn tại.

    Ưu tiên tên header chính xác; nếu phía gọi khác hoa/thường hoặc dấu tiếng
    Việt thì dò lại bằng tên đã chuẩn hóa. Điều này giúp API ổn định khi file
    Excel đổi cách viết ``Số Lượng``/``Số lượng`` nhưng vẫn là cùng một cột.
    """
    unknown_headers = []
    normalized_header_map: dict[str, int | None] = {}
    for actual_header, col_idx in header_map.items():
        key = normalize_header(actual_header)
        if key in normalized_header_map and normalized_header_map[key] != col_idx:
            # Không tự đoán nếu workbook có hai cột trùng tên sau chuẩn hóa.
            normalized_header_map[key] = None
        else:
            normalized_header_map[key] = col_idx

    max_col = ws.max_column
    for header, value in data.items():
        col = header_map.get(header)
        if col is None:
            col = normalized_header_map.get(normalize_header(header))
        if col is None:
            unknown_headers.append(str(header))
            continue
        target_col = target_column_for_header(ws, header_row, row_number, col, str(header), max_col) if header_row else col
        old = ws.cell(row_number, target_col).value
        ws.cell(row_number, target_col).value = coerce_value(str(header), value, old)
    return unknown_headers


def command_update_row(args):
    input_file = Path(args.file)
    output_file = output_path_from_args(input_file, args.output, args.suffix)
    data = json.loads(args.data_json or "{}")
    row_number = int(args.row)

    wb, ws, real_name = load_sheet(input_file, args.sheet)
    data, blocked_headers = restrict_row_data_for_sheet(real_name, data)
    if not data:
        if blocked_headers:
            raise ValueError("Sheet phẫu thuật chỉ cho phép sửa PTV chính và Phụ mổ 1–3.")
        raise ValueError("Không có dữ liệu hợp lệ để cập nhật.")
    header_row, headers, header_map = find_header_row_generic(ws)
    try:
        total_row = find_total_row(ws)
    except Exception:
        total_row = ws.max_row + 1
    if row_number <= header_row or row_number >= total_row:
        raise ValueError(
            f"Dòng {row_number} không hợp lệ cho sheet {real_name}; "
            "không được sửa header hoặc dòng Tổng Cộng."
        )

    unknown_headers = apply_row_changes(ws, header_map, row_number, data, header_row)

    warnings = update_totals_for_sheet(wb, ws, header_row)
    if blocked_headers:
        warnings.append("Đã bảo vệ và bỏ qua cột không được phép sửa: " + ", ".join(sorted(set(blocked_headers))))
    if unknown_headers:
        warnings.append("Bỏ qua cột không tồn tại: " + ", ".join(sorted(set(unknown_headers))))
    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_file)
    respond({
        "ok": True,
        "file": str(output_file),
        "sheet": real_name,
        "rowNumber": row_number,
        "warnings": warnings,
    })


def command_update_rows(args):
    """Cập nhật nhiều dòng trong một lần mở/lưu workbook."""
    input_file = Path(args.file)
    output_file = output_path_from_args(input_file, args.output, args.suffix)
    payload = json.loads(args.rows_json or "[]")
    if not isinstance(payload, list) or not payload:
        raise ValueError("Danh sách dòng cập nhật đang trống.")

    wb, ws, real_name = load_sheet(input_file, args.sheet)
    header_row, _headers, header_map = find_header_row_generic(ws)
    try:
        total_row = find_total_row(ws)
    except Exception:
        total_row = ws.max_row + 1

    updated_rows: list[int] = []
    unknown_headers: list[str] = []
    blocked_headers: list[str] = []
    merged_changes: dict[int, dict[str, Any]] = {}

    # Gộp payload trùng dòng để một dòng chỉ được xử lý một lần.
    for item in payload:
        if not isinstance(item, dict):
            continue
        row_number = int(item.get("rowNumber") or item.get("row") or 0)
        data = item.get("data") or {}
        if row_number <= header_row or row_number >= total_row:
            raise ValueError(
                f"Dòng {row_number} không hợp lệ cho sheet {real_name}; "
                "không được sửa header hoặc dòng Tổng Cộng."
            )
        if not isinstance(data, dict):
            raise ValueError(f"Dữ liệu dòng {row_number} không đúng định dạng.")
        safe_data, blocked = restrict_row_data_for_sheet(real_name, data)
        blocked_headers.extend(blocked)
        if safe_data:
            merged_changes.setdefault(row_number, {}).update(safe_data)

    if not merged_changes:
        if blocked_headers:
            raise ValueError("Sheet phẫu thuật chỉ cho phép sửa PTV chính và Phụ mổ 1–3.")
        raise ValueError("Không có dữ liệu hợp lệ để cập nhật.")

    for row_number, data in sorted(merged_changes.items()):
        unknown_headers.extend(apply_row_changes(ws, header_map, row_number, data, header_row))
        updated_rows.append(row_number)

    warnings = update_totals_for_sheet(wb, ws, header_row)
    if blocked_headers:
        warnings.append("Đã bảo vệ và bỏ qua cột không được phép sửa: " + ", ".join(sorted(set(blocked_headers))))
    if unknown_headers:
        warnings.append("Bỏ qua cột không tồn tại: " + ", ".join(sorted(set(unknown_headers))))

    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_file)
    respond({
        "ok": True,
        "file": str(output_file),
        "sheet": real_name,
        "updatedRows": updated_rows,
        "updatedCount": len(updated_rows),
        "warnings": warnings,
    })



def command_clear_columns(args):
    """Xem trước hoặc xóa dữ liệu trong các cột được chọn của một sheet.

    Chức năng này được dùng khi file tháng mới vẫn mang dữ liệu mẫu ở các cột
    Phụ mổ 1–3. Chỉ các dòng dữ liệu giữa header và Tổng Cộng được xử lý;
    công thức được giữ nguyên để tránh phá cấu trúc workbook.
    """
    input_file = Path(args.file)
    requested_columns = json.loads(args.columns_json or "[]")
    if not isinstance(requested_columns, list) or not requested_columns:
        raise ValueError("Chưa chọn cột cần xóa.")

    wb, ws, real_name = load_sheet(input_file, args.sheet)
    header_row, headers, header_map = find_header_row_generic(ws)
    try:
        total_row = find_total_row(ws)
    except Exception:
        total_row = ws.max_row + 1

    normalized_map: dict[str, tuple[str, int] | None] = {}
    for actual_header, col_idx in header_map.items():
        key = normalize_header(actual_header)
        if key in normalized_map and normalized_map[key] != (actual_header, col_idx):
            normalized_map[key] = None
        else:
            normalized_map[key] = (actual_header, col_idx)

    matched: list[tuple[str, str, int]] = []
    missing_columns: list[str] = []
    for requested in requested_columns:
        requested_name = str(requested).strip()
        hit = normalized_map.get(normalize_header(requested_name))
        if not hit:
            missing_columns.append(requested_name)
            continue
        actual_header, col_idx = hit
        matched.append((requested_name, actual_header, col_idx))

    if not matched:
        raise ValueError("Không tìm thấy cột cần xóa trong sheet hiện tại.")

    column_counts: dict[str, int] = {requested: 0 for requested, _actual, _col in matched}
    affected_rows: set[int] = set()
    skipped_formula_cells = 0

    for row_idx in range(header_row + 1, min(total_row, ws.max_row + 1)):
        for requested, _actual, col_idx in matched:
            cell = ws.cell(row_idx, col_idx)
            value = cell.value
            if value_is_blank(value):
                continue
            if isinstance(value, str) and value.startswith("="):
                skipped_formula_cells += 1
                continue
            column_counts[requested] += 1
            affected_rows.add(row_idx)
            if not args.preview:
                cell.value = None

    available_columns = [requested for requested, _actual, _col in matched]
    result = {
        "ok": True,
        "file": str(input_file),
        "sheet": real_name,
        "preview": bool(args.preview),
        "columns": available_columns,
        "availableColumns": available_columns,
        "missingColumns": missing_columns,
        "columnCounts": column_counts,
        "clearedCells": sum(column_counts.values()),
        "clearedRows": len(affected_rows),
        "skippedFormulaCells": skipped_formula_cells,
    }

    if args.preview:
        respond(result)
        return

    output_file = output_path_from_args(input_file, args.output, args.suffix)
    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_file)
    result["file"] = str(output_file)
    respond(result)


def next_stt(ws, header_row: int, total_row: int, stt_col: int) -> int:
    values = []
    for r in range(header_row + 1, total_row):
        v = ws.cell(r, stt_col).value
        try:
            values.append(int(float(str(v).strip())))
        except Exception:
            pass
    return (max(values) if values else 0) + 1


def command_insert_row(args):
    input_file = Path(args.file)
    output_file = output_path_from_args(input_file, args.output, args.suffix)
    data = json.loads(args.data_json or "{}")

    wb, ws, real_name = load_sheet(input_file, args.sheet)
    if "phauthuat" in normalize_header(real_name):
        raise ValueError("Sheet phẫu thuật chỉ dùng để nhập PTV chính và Phụ mổ 1–3; không cho phép thêm hoặc xóa dòng.")
    header_row, headers, header_map = find_header_row_generic(ws)
    try:
        total_row_before = find_total_row(ws)
    except Exception:
        total_row_before = ws.max_row + 1

    insert_at = int(args.insert_at or total_row_before)
    if insert_at <= header_row:
        insert_at = header_row + 1

    warnings: list[str] = [HELPER_IMPORT_WARNING] if HELPER_IMPORT_WARNING else []
    style_row = max(header_row + 1, insert_at - 1)
    insert_rows_preserve_merges(ws, insert_at, 1)
    try:
        copy_row_style(ws, style_row, insert_at, ws.max_column)
    except Exception as exc:
        warnings.append(f"Không sao chép được định dạng dòng mẫu: {exc}")

    # Sao chép công thức từ dòng mẫu sang dòng mới, tự dịch tham chiếu theo dòng mới.
    try:
        for col_idx in range(1, ws.max_column + 1):
            src = ws.cell(style_row, col_idx)
            dst = ws.cell(insert_at, col_idx)
            if isinstance(src.value, str) and src.value.startswith('='):
                dst.value = translate_formula(src.value, src, dst)
    except Exception as exc:
        warnings.append(f"Không sao chép được công thức dòng mẫu: {exc}")

    stt_header = None
    for h in headers:
        if normalize_header(h) == "stt":
            stt_header = h
            break
    if stt_header:
        ws.cell(insert_at, header_map[stt_header]).value = next_stt(ws, header_row, total_row_before, header_map[stt_header])

    insert_data = dict(data)
    if stt_header and stt_header in insert_data and value_is_blank(insert_data[stt_header]):
        insert_data.pop(stt_header, None)
    unknown_headers = apply_row_changes(ws, header_map, insert_at, insert_data, header_row)
    if unknown_headers:
        warnings.append("Bỏ qua cột không tồn tại: " + ", ".join(sorted(set(unknown_headers))))

    warnings.extend(update_totals_for_sheet(wb, ws, header_row))
    warnings = list(dict.fromkeys(item for item in warnings if item))
    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_file)
    respond({
        "ok": True,
        "file": str(output_file),
        "sheet": real_name,
        "rowNumber": insert_at,
        "warnings": warnings,
    })


def command_delete_row(args):
    input_file = Path(args.file)
    output_file = output_path_from_args(input_file, args.output, args.suffix)
    row_number = int(args.row)

    wb, ws, real_name = load_sheet(input_file, args.sheet)
    if "phauthuat" in normalize_header(real_name):
        raise ValueError("Sheet phẫu thuật chỉ dùng để nhập PTV chính và Phụ mổ 1–3; không cho phép thêm hoặc xóa dòng.")
    header_row, headers, header_map = find_header_row_generic(ws)
    try:
        total_row = find_total_row(ws)
    except Exception:
        total_row = ws.max_row + 1
    if row_number <= header_row or row_number >= total_row:
        raise ValueError(f"Chỉ được xóa dòng dữ liệu, không xóa header hoặc Tổng Cộng. Dòng: {row_number}")
    delete_rows_preserve_merges(ws, [row_number])
    warnings = update_totals_for_sheet(wb, ws, header_row)
    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_file)
    respond({
        "ok": True,
        "file": str(output_file),
        "sheet": real_name,
        "deletedRow": row_number,
        "warnings": warnings,
    })




SPECIAL_PAYMENT_ALIASES = {"v.tan", "nguyen", "toan"}


def _find_header_row_by_keys(ws, required_keys: set[str], max_scan_rows: int = 30) -> int:
    for row_idx in range(1, min(ws.max_row, max_scan_rows) + 1):
        keys = {
            normalize_header(ws.cell(row_idx, col_idx).value)
            for col_idx in range(1, ws.max_column + 1)
            if not value_is_blank(ws.cell(row_idx, col_idx).value)
        }
        if required_keys.issubset(keys):
            return row_idx
    raise ValueError(
        f"Không tìm thấy dòng tiêu đề có các cột: {', '.join(sorted(required_keys))} trong sheet {ws.title}."
    )


def _header_positions(ws, header_row: int) -> dict[str, list[int]]:
    result: dict[str, list[int]] = {}
    for col_idx in range(1, ws.max_column + 1):
        key = normalize_header(ws.cell(header_row, col_idx).value)
        if key:
            result.setdefault(key, []).append(col_idx)
    return result


def _first_col(positions: dict[str, list[int]], key: str, *, required: bool = True) -> int | None:
    cols = positions.get(normalize_header(key), [])
    if cols:
        return cols[0]
    if required:
        raise ValueError(f"Không tìm thấy cột '{key}'.")
    return None


def _row_is_data(ws, row_idx: int, columns: list[int | None]) -> bool:
    return any(col and not value_is_blank(ws.cell(row_idx, col).value) for col in columns)


def _copy_formula_cell_style(ws, template_row: int | None, row_idx: int, columns: list[int]) -> None:
    if not template_row or template_row == row_idx:
        return
    for col_idx in columns:
        src = ws.cell(template_row, col_idx)
        dst = ws.cell(row_idx, col_idx)
        if dst.style_id == 0 and src.style_id != 0:
            dst._style = copy(src._style)
        if src.number_format and dst.number_format == "General":
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.border:
            dst.border = copy(src.border)
        if src.fill:
            dst.fill = copy(src.fill)
        if src.font:
            dst.font = copy(src.font)


def _find_formula_template_row(ws, start_row: int, end_row: int, formula_columns: list[int]) -> int | None:
    for row_idx in range(start_row, end_row + 1):
        if any(
            isinstance(ws.cell(row_idx, col_idx).value, str)
            and str(ws.cell(row_idx, col_idx).value).startswith("=")
            for col_idx in formula_columns
        ):
            return row_idx
    return None


def _set_workbook_recalculation(wb) -> None:
    try:
        wb.calculation.calcMode = "auto"
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
    except Exception:
        pass


def _finalize_tieuphau_sheet(ws) -> dict[str, Any]:
    header_row = _find_header_row_by_keys(ws, {"stt", "bacsi", "dieuduong", "thuclanh"})
    positions = _header_positions(ws, header_row)
    total_row = find_total_row(ws)

    col_stt = _first_col(positions, "STT")
    col_name = _first_col(positions, "Họ và tên", required=False)
    col_cls = _first_col(positions, "Tên CLS", required=False)
    col_amount = _first_col(positions, "Thành tiền")
    col_doctor = _first_col(positions, "Bác Sĩ")
    col_nurse = _first_col(positions, "Điều Dưỡng")
    col_take_home = _first_col(positions, "Thực Lãnh")

    # Hai cột "Số Tiền" được xác định theo vị trí: ngay sau Bác sĩ và Điều dưỡng.
    col_doctor_money = col_doctor + 1
    col_nurse_money = col_nurse + 1
    formula_cols = [col_doctor_money, col_nurse_money, col_take_home]
    template_row = _find_formula_template_row(ws, header_row + 1, total_row - 1, formula_cols)

    changed_rows = 0
    nurse_filled = 0
    formulas_written = 0
    data_rows: list[int] = []
    for row_idx in range(header_row + 1, total_row):
        if not _row_is_data(ws, row_idx, [col_stt, col_name, col_cls, col_amount, col_doctor]):
            continue
        data_rows.append(row_idx)
        _copy_formula_cell_style(ws, template_row, row_idx, formula_cols)
        row_changed = False
        if value_is_blank(ws.cell(row_idx, col_nurse).value):
            ws.cell(row_idx, col_nurse).value = "ĐD"
            nurse_filled += 1
            row_changed = True

        # Cập nhật lại toàn bộ công thức để sửa cả công thức cũ dùng "TÂN" thay vì "V.TÂN".
        ws.cell(row_idx, col_nurse_money).value = f"={get_column_letter(col_amount)}{row_idx}*0.05"
        ws.cell(row_idx, col_take_home).value = (
            f'=IF(OR(UPPER(TRIM({get_column_letter(col_doctor)}{row_idx}))="V.TÂN",'
            f'UPPER(TRIM({get_column_letter(col_doctor)}{row_idx}))="NGUYỆN",'
            f'UPPER(TRIM({get_column_letter(col_doctor)}{row_idx}))="TOẢN"),'
            f'{get_column_letter(col_amount)}{row_idx}*0.1875,'
            f'{get_column_letter(col_amount)}{row_idx}*0.25)'
        )
        ws.cell(row_idx, col_doctor_money).value = (
            f"={get_column_letter(col_take_home)}{row_idx}-{get_column_letter(col_nurse_money)}{row_idx}"
        )
        for col_idx in formula_cols:
            ws.cell(row_idx, col_idx).number_format = "#,##0"
        formulas_written += len(formula_cols)
        row_changed = True
        if row_changed:
            changed_rows += 1

    if data_rows:
        start, end = data_rows[0], data_rows[-1]
        total_columns = [
            _first_col(positions, "Số Lượng", required=False),
            col_amount,
            col_doctor_money,
            col_nurse_money,
            col_take_home,
        ]
        for col_idx in [col for col in total_columns if col]:
            ws.cell(total_row, col_idx).value = (
                f"=SUM({get_column_letter(col_idx)}{start}:{get_column_letter(col_idx)}{end})"
            )
            ws.cell(total_row, col_idx).number_format = "#,##0"

    return {
        "sheet": ws.title,
        "dataRows": len(data_rows),
        "changedRows": changed_rows,
        "nurseFilled": nurse_filled,
        "formulasWritten": formulas_written,
        "totalRow": total_row,
    }


def _finalize_phauthuat_sheet(ws) -> dict[str, Any]:
    header_row = _find_header_row_by_keys(ws, {"stt", "ptvchinh", "phumo1", "phumo2", "phumo3", "thuclanh"})
    positions = _header_positions(ws, header_row)
    total_row = find_total_row(ws)

    col_stt = _first_col(positions, "STT")
    col_name = _first_col(positions, "Họ và tên bệnh nhân", required=False)
    col_method = _first_col(positions, "Chẩn đoán và phương pháp phẫu thuật", required=False)
    col_ptv = _first_col(positions, "PTV chính")
    col_phu1 = _first_col(positions, "Phụ mổ 1")
    col_phu2 = _first_col(positions, "Phụ mổ 2")
    col_phu3 = _first_col(positions, "Phụ mổ 3")
    col_ddhl = _first_col(positions, "ĐD + HL")
    col_take_home = _first_col(positions, "Thực lãnh")
    col_amount = _first_col(positions, "Số tiền thực tính thù lao cho khoa", required=False) or (col_ptv - 1)

    col_ptv_money = col_ptv + 1
    col_phu1_money = col_phu1 + 1
    col_phu2_money = col_phu2 + 1
    col_phu3_money = col_phu3 + 1
    formula_cols = [col_ptv_money, col_phu1_money, col_phu2_money, col_phu3_money, col_ddhl, col_take_home]
    template_row = _find_formula_template_row(ws, header_row + 1, total_row - 1, formula_cols)

    data_rows: list[int] = []
    assistants_defaulted = 0
    formulas_written = 0
    for row_idx in range(header_row + 1, total_row):
        if not _row_is_data(ws, row_idx, [col_stt, col_name, col_method, col_amount, col_ptv]):
            continue
        data_rows.append(row_idx)
        _copy_formula_cell_style(ws, template_row, row_idx, formula_cols)

        for col_idx in [col_phu1, col_phu2, col_phu3]:
            if value_is_blank(ws.cell(row_idx, col_idx).value):
                ws.cell(row_idx, col_idx).value = "-"
                assistants_defaulted += 1

        amount_ref = f"{get_column_letter(col_amount)}{row_idx}"
        ptv_ref = f"{get_column_letter(col_ptv)}{row_idx}"
        ws.cell(row_idx, col_phu1_money).value = (
            f'=IF(OR({get_column_letter(col_phu1)}{row_idx}="-",TRIM({get_column_letter(col_phu1)}{row_idx})=""),0,{amount_ref}*0.03)'
        )
        ws.cell(row_idx, col_phu2_money).value = (
            f'=IF(OR({get_column_letter(col_phu2)}{row_idx}="-",TRIM({get_column_letter(col_phu2)}{row_idx})=""),0,{amount_ref}*0.02)'
        )
        ws.cell(row_idx, col_phu3_money).value = (
            f'=IF(OR({get_column_letter(col_phu3)}{row_idx}="-",TRIM({get_column_letter(col_phu3)}{row_idx})=""),0,{amount_ref}*0.01)'
        )
        ws.cell(row_idx, col_ddhl).value = f"={amount_ref}*0.005"
        ws.cell(row_idx, col_take_home).value = (
            f'=IF(OR(UPPER(TRIM({ptv_ref}))="V.TÂN",UPPER(TRIM({ptv_ref}))="NGUYỆN",'
            f'UPPER(TRIM({ptv_ref}))="TOẢN"),{amount_ref}*0.1375,{amount_ref}*0.17)'
        )
        ws.cell(row_idx, col_ptv_money).value = (
            f"={get_column_letter(col_take_home)}{row_idx}"
            f"-{get_column_letter(col_phu1_money)}{row_idx}"
            f"-{get_column_letter(col_phu2_money)}{row_idx}"
            f"-{get_column_letter(col_phu3_money)}{row_idx}"
            f"-{get_column_letter(col_ddhl)}{row_idx}"
        )
        for col_idx in formula_cols:
            ws.cell(row_idx, col_idx).number_format = "#,##0"
        formulas_written += len(formula_cols)

    if data_rows:
        start, end = data_rows[0], data_rows[-1]
        for col_idx in [col_amount, col_ptv_money, col_phu1_money, col_phu2_money, col_phu3_money, col_ddhl, col_take_home]:
            ws.cell(total_row, col_idx).value = (
                f"=SUM({get_column_letter(col_idx)}{start}:{get_column_letter(col_idx)}{end})"
            )
            ws.cell(total_row, col_idx).number_format = "#,##0"

    return {
        "sheet": ws.title,
        "dataRows": len(data_rows),
        "assistantsDefaulted": assistants_defaulted,
        "formulasWritten": formulas_written,
        "blankMainSurgeonRows": sum(1 for row_idx in data_rows if value_is_blank(ws.cell(row_idx, col_ptv).value)),
        "totalRow": total_row,
    }


def finalize_workbook_for_download(input_file: Path, output_file: Path | None = None) -> dict[str, Any]:
    if not input_file.exists():
        raise FileNotFoundError(f"Không tìm thấy file: {input_file}")
    output_file = output_file or input_file
    keep_vba = input_file.suffix.lower() == ".xlsm"
    wb = load_workbook(input_file, keep_vba=keep_vba)
    reports: list[dict[str, Any]] = []

    for wanted, handler in (("phauthuat", _finalize_phauthuat_sheet), ("tieuphau", _finalize_tieuphau_sheet)):
        try:
            real_name = find_sheet_case_insensitive(wb, wanted)
        except Exception:
            continue
        reports.append(handler(wb[real_name]))

    _set_workbook_recalculation(wb)
    output_file.parent.mkdir(parents=True, exist_ok=True)
    if output_file.resolve() == input_file.resolve():
        tmp = output_file.with_name(f".{output_file.stem}.dang_hoan_tat{output_file.suffix}")
        wb.save(tmp)
        tmp.replace(output_file)
    else:
        wb.save(output_file)

    return {
        "ok": True,
        "file": str(output_file),
        "sheets": reports,
        "message": "Đã điền công thức, giá trị mặc định và dòng Tổng Cộng trước khi tải Excel.",
    }


def command_finalize_workbook(args):
    input_file = Path(args.file)
    output_file = Path(args.output) if getattr(args, "output", None) else input_file
    respond(finalize_workbook_for_download(input_file, output_file))


def command_task_bo_sung_phu_mo(args):
    import bo_sung_phu_mo_t5_tu_so_phau_thuat as mod
    input_file = Path(args.file)
    output_file = output_path_from_args(input_file, args.output, "_da_bo_sung_phu_mo")
    report_file = report_path_for(output_file)
    mod.FILE_PM_T5 = input_file
    mod.FILE_DAU_RA = output_file
    mod.FILE_BAO_CAO = report_file
    mod.GHI_DE_FILE_GOC = False
    mod.bo_sung_phu_mo()
    respond({"ok": True, "file": str(output_file), "report": str(report_file), "message": "Đã bổ sung PTV chính/phụ mổ bằng bí danh."})


def command_task_chuyen_tieu_phau(args):
    import chuyen_thu_thuat_sang_tieuphau_t5 as mod
    input_file = Path(args.file)
    output_file = output_path_from_args(input_file, args.output, "_da_chuyen_tieu_phau")
    report_file = report_path_for(output_file)
    mod.FILE_NGUON = input_file
    mod.FILE_DAU_RA = output_file
    mod.FILE_BAO_CAO = report_file
    result = mod.chuyen_thu_thuat_sang_tieuphau()
    result.update({"ok": True, "file": str(output_file), "report": str(report_file)})
    respond(result)


def command_task_nhap_phau_thuat_so_bo(args):
    """Chuẩn hóa (tách 4 cột BS) file 'Danh sách sơ bộ' trích từ ảnh sổ phẫu
    thuật, rồi nhập các dòng đó vào sheet phauthuat của file đang dùng."""
    import xu_ly_so_phau_thuat as xu_ly_mod
    import nhap_phau_thuat_tu_so_bo as nhap_mod

    input_file = Path(args.file)
    so_bo_file = Path(args.so_bo)
    zip_file = Path(args.zip)
    sheet_so_bo = args.sheet_so_bo or nhap_mod.SHEET_SO_BO_MAC_DINH
    output_file = output_path_from_args(input_file, args.output, "_NHAP_LIEU")
    report_file = report_path_for(output_file).with_name(
        output_file.stem + "_bao_cao_nhap_phau_thuat.xlsx"
    )
    chuan_hoa_file = output_file.with_name(output_file.stem + "_so_bo_chuan_hoa_tmp.xlsx")

    xu_ly_mod.normalize_workbook(so_bo_file, zip_file, chuan_hoa_file, sheet_so_bo)
    result = nhap_mod.nhap_phau_thuat_tu_so_bo(input_file, chuan_hoa_file, output_file, sheet_so_bo)
    nhap_mod.tao_bao_cao(result["bao_cao_rows"], report_file)
    try:
        chuan_hoa_file.unlink()
    except Exception:
        pass

    result.pop("bao_cao_rows", None)
    result.update({"ok": True, "file": str(output_file), "report": str(report_file)})
    respond(result)


def command_extract_zip_images(args):
    """Giải nén ảnh (.jpg/.jpeg/.png) từ ZIP vào một thư mục cố định, chỉ giữ
    đúng 1 bộ ảnh đang dùng (xóa ảnh cũ trước khi giải nén) để khớp với đúng
    một 'quyển sổ phẫu thuật' đang xem trên web."""
    import zipfile

    zip_path = Path(args.zip)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    for existing in output_dir.glob("*"):
        if existing.is_file():
            try:
                existing.unlink()
            except Exception:
                pass

    extracted: list[str] = []
    with zipfile.ZipFile(zip_path) as zf:
        for name in zf.namelist():
            if not name.lower().endswith((".jpg", ".jpeg", ".png")):
                continue
            base = Path(name).name
            if not base:
                continue
            target = output_dir / base
            with zf.open(name) as src, open(target, "wb") as dst:
                dst.write(src.read())
            extracted.append(base)

    extracted.sort(key=str.casefold)
    respond({"ok": True, "images": extracted, "count": len(extracted)})


def command_task_cap_nhat_cls_tieuphau(args):
    """Chuyển các CLS vừa bổ sung; lưu ngay và để Bác Sĩ trống cho EMR xử lý."""
    import chuyen_thu_thuat_sang_tieuphau_t5 as mod

    input_file = Path(args.file)
    output_file = output_path_from_args(input_file, args.output, "_NHAP_LIEU")
    report_file = report_path_for(output_file)
    mod.FILE_NGUON = input_file
    mod.FILE_DAU_RA = output_file
    mod.FILE_BAO_CAO = report_file
    mod.DIEN_BS_TU_SO_SAU_CHUYEN = False

    result = mod.chuyen_thu_thuat_sang_tieuphau()
    moved = int(result.get("so_dong_chuyen") or 0)
    result.update({
        "ok": True,
        "file": str(output_file),
        "report": str(report_file),
        "message": (
            f"Đã chuyển thêm {moved} dòng CLS còn sót và lưu vào file làm việc."
            if moved
            else "Không còn dòng CLS mới cần chuyển; tiếp tục kiểm tra Bác sĩ trên EMR."
        ),
    })
    respond(result)


def command_task_dien_bs(args):
    import dien_bs_tieuphau_t5_tu_so as mod
    input_file = Path(args.file)
    output_file = output_path_from_args(input_file, args.output, "_da_dien_BS")
    report_file = report_path_for(output_file)
    mod.FILE_NGUON = input_file
    mod.FILE_DAU_RA = output_file
    mod.FILE_BAO_CAO_BS = report_file
    result = mod.dien_bs_tieuphau()
    result.update({"ok": True, "file": str(output_file), "report": str(report_file)})
    respond(result)


def command_task_tat_ca(args):
    import bo_sung_phu_mo_t5_tu_so_phau_thuat as mod_phu_mo
    import chuyen_thu_thuat_sang_tieuphau_t5 as mod_tieu_phau

    input_file = Path(args.file)
    output_file = output_path_from_args(input_file, args.output, "_NHAP_LIEU")
    report_file = report_path_for(output_file)

    # Bước 1: bổ sung PTV chính/phụ mổ vào chính file làm việc.
    mod_phu_mo.FILE_PM_T5 = input_file
    mod_phu_mo.FILE_DAU_RA = output_file
    mod_phu_mo.FILE_BAO_CAO = report_file
    mod_phu_mo.GHI_DE_FILE_GOC = False
    mod_phu_mo.bo_sung_phu_mo()

    # Bước 2: chuyển thủ thuật sang tiểu phẫu + điền BS, tiếp tục ghi vào cùng file làm việc.
    mod_tieu_phau.FILE_NGUON = output_file
    mod_tieu_phau.FILE_DAU_RA = output_file
    mod_tieu_phau.FILE_BAO_CAO = report_file
    result = mod_tieu_phau.chuyen_thu_thuat_sang_tieuphau()

    result.update({
        "ok": True,
        "file": str(output_file),
        "report": str(report_file),
        "message": "Đã chạy toàn bộ vào cùng một file làm việc."
    })
    respond(result)


def build_parser():
    parser = argparse.ArgumentParser(description="Node Excel API")
    sub = parser.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("staff-list")

    p = sub.add_parser("save-staff-list")
    p.add_argument("--staff-json", required=True)

    p = sub.add_parser("cls-list")

    p = sub.add_parser("save-cls-list")
    p.add_argument("--cls-json", required=True)

    p = sub.add_parser("finalize-workbook")
    p.add_argument("--file", required=True)
    p.add_argument("--output")

    p = sub.add_parser("list-files")
    p.add_argument("--root")

    p = sub.add_parser("workbook-info")
    p.add_argument("--file", required=True)

    p = sub.add_parser("normalize-dates")
    p.add_argument("--file", required=True)
    p.add_argument("--output")

    p = sub.add_parser("read-sheet")
    p.add_argument("--file", required=True)
    p.add_argument("--sheet")
    p.add_argument("--query", default="")
    p.add_argument("--missing-only", action="store_true")
    p.add_argument("--limit", default="500")

    p = sub.add_parser("update-row")
    p.add_argument("--file", required=True)
    p.add_argument("--sheet", required=True)
    p.add_argument("--row", required=True)
    p.add_argument("--data-json", required=True)
    p.add_argument("--output")
    p.add_argument("--suffix")

    p = sub.add_parser("update-rows")
    p.add_argument("--file", required=True)
    p.add_argument("--sheet", required=True)
    p.add_argument("--rows-json", required=True)
    p.add_argument("--output")
    p.add_argument("--suffix")

    p = sub.add_parser("clear-columns")
    p.add_argument("--file", required=True)
    p.add_argument("--sheet", required=True)
    p.add_argument("--columns-json", required=True)
    p.add_argument("--preview", action="store_true")
    p.add_argument("--output")
    p.add_argument("--suffix")

    p = sub.add_parser("insert-row")
    p.add_argument("--file", required=True)
    p.add_argument("--sheet", required=True)
    p.add_argument("--data-json", required=True)
    p.add_argument("--insert-at")
    p.add_argument("--output")
    p.add_argument("--suffix")

    p = sub.add_parser("delete-row")
    p.add_argument("--file", required=True)
    p.add_argument("--sheet", required=True)
    p.add_argument("--row", required=True)
    p.add_argument("--output")
    p.add_argument("--suffix")

    p = sub.add_parser("task-tat-ca")
    p.add_argument("--file", required=True)
    p.add_argument("--output")

    p = sub.add_parser("task-bo-sung-phu-mo")
    p.add_argument("--file", required=True)
    p.add_argument("--output")

    p = sub.add_parser("task-chuyen-tieu-phau")
    p.add_argument("--file", required=True)
    p.add_argument("--output")

    p = sub.add_parser("task-cap-nhat-cls-tieuphau")
    p.add_argument("--file", required=True)
    p.add_argument("--output")

    p = sub.add_parser("task-dien-bs")
    p.add_argument("--file", required=True)
    p.add_argument("--output")

    p = sub.add_parser("task-nhap-phau-thuat-so-bo")
    p.add_argument("--file", required=True)
    p.add_argument("--so-bo", required=True, dest="so_bo")
    p.add_argument("--zip", required=True)
    p.add_argument("--sheet-so-bo", dest="sheet_so_bo")
    p.add_argument("--output")

    p = sub.add_parser("extract-zip-images")
    p.add_argument("--zip", required=True)
    p.add_argument("--output-dir", required=True, dest="output_dir")

    return parser


def main():
    parser = build_parser()
    args = parser.parse_args()
    try:
        commands = {
            "staff-list": command_staff_list,
            "save-staff-list": command_save_staff_list,
            "cls-list": command_cls_list,
            "save-cls-list": command_save_cls_list,
            "finalize-workbook": command_finalize_workbook,
            "list-files": command_list_files,
            "workbook-info": command_workbook_info,
            "normalize-dates": command_normalize_dates,
            "read-sheet": command_read_sheet,
            "update-row": command_update_row,
            "update-rows": command_update_rows,
            "clear-columns": command_clear_columns,
            "insert-row": command_insert_row,
            "delete-row": command_delete_row,
            "task-tat-ca": command_task_tat_ca,
            "task-bo-sung-phu-mo": command_task_bo_sung_phu_mo,
            "task-chuyen-tieu-phau": command_task_chuyen_tieu_phau,
            "task-cap-nhat-cls-tieuphau": command_task_cap_nhat_cls_tieuphau,
            "task-dien-bs": command_task_dien_bs,
            "task-nhap-phau-thuat-so-bo": command_task_nhap_phau_thuat_so_bo,
            "extract-zip-images": command_extract_zip_images,
        }
        commands[args.cmd](args)
    except SystemExit:
        raise
    except Exception as exc:
        fail(str(exc))


if __name__ == "__main__":
    main()
