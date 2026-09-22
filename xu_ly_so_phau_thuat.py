#!/usr/bin/env python3
"""
Chuẩn hóa danh sách phẫu thuật từ file Excel + ZIP ảnh.

Chạy:
    python xu_ly_so_phau_thuat.py \
        --excel Danh_sach_phau_thuat.xlsx \
        --zip img.zip \
        --output Danh_sach_phau_thuat_da_chuan_hoa.xlsx

Lưu ý: script này xử lý phần chuẩn hóa Excel. Việc đọc chữ viết tay
(tên bác sĩ, bệnh nhân, chẩn đoán...) vẫn cần OCR/AI hoặc nhập sơ bộ
để script có dữ liệu đầu vào.
"""

from __future__ import annotations

import argparse
import re
import zipfile
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment


DOCTOR_HEADERS = [
    "BS phẫu thuật 1",
    "BS phẫu thuật 2",
    "BS phẫu thuật 3",
    "BS phẫu thuật 4",
]


def header_map(ws):
    return {
        str(cell.value).strip(): cell.column
        for cell in ws[5]
        if cell.value not in (None, "")
    }


def find_header(headers, *names):
    for name in names:
        if name in headers:
            return headers[name]
    return None


def zip_image_names(zip_path: Path):
    with zipfile.ZipFile(zip_path) as zf:
        names = [
            Path(name).name
            for name in zf.namelist()
            if name.lower().endswith(('.jpg', '.jpeg', '.png'))
        ]
    return sorted(names, key=str.casefold)


def original_zip_name(source, image_names):
    text = str(source or "").strip()
    match = re.fullmatch(r"0*(\d+)\.(?:jpg|jpeg|png)", text, re.I)
    if match:
        index = int(match.group(1)) - 1
        if 0 <= index < len(image_names):
            return image_names[index]
    return text


def split_doctors(raw):
    """Tách chuỗi PTV; BS4 chỉ giữ một người, phần còn lại bị bỏ."""
    text = str(raw or "").replace("\r", "\n").strip()
    if not text:
        return ["-", "-", "-", "-"], ""

    # Chuỗi gốc thường dùng dấu gạch ngang để nối tên.
    parts = [p.strip() for p in re.split(r"\s+-\s+", text) if p.strip()]
    # Nếu dữ liệu đã có bốn cột, tên thừa đôi khi nằm trong BS4 dạng A / B.
    if len(parts) >= 4 and " / " in parts[3]:
        parts[3] = parts[3].split(" / ", 1)[0].strip()

    warning = ""
    if len(parts) >= 4 or "\n" in text:
        warning = "CẢNH BÁO: PTV có thể được ghi tiếp trên 2 dòng."
        if len(parts) > 4 or " / " in text:
            warning += " Đã bỏ tên phía sau BS4."
        warning += " Kiểm tra ảnh nguồn."

    doctors = parts[:4]
    while len(doctors) < 4:
        doctors.append("-")
    return doctors, warning


def copy_basic_style(ws, source_col, target_col):
    """Giữ kiểu chữ/căn lề cơ bản khi tạo cột mới."""
    for row in range(5, ws.max_row + 1):
        src = ws.cell(row, source_col)
        dst = ws.cell(row, target_col)
        if src.has_style:
            # _style là mảng nội bộ của openpyxl, không có .copy() ở một số
            # phiên bản — copy từng thuộc tính qua module copy chuẩn để chắc
            # ăn với mọi phiên bản openpyxl.
            dst.font = copy(src.font)
            dst.fill = copy(src.fill)
            dst.border = copy(src.border)
            dst.protection = copy(src.protection)
        if src.number_format:
            dst.number_format = src.number_format
        dst.alignment = Alignment(
            horizontal=src.alignment.horizontal,
            vertical="center",
            wrap_text=True,
        )


def normalize_workbook(excel_path: Path, zip_path: Path, output_path: Path, sheet_name: str):
    wb = load_workbook(excel_path)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Không tìm thấy sheet: {sheet_name}")
    ws = wb[sheet_name]
    image_names = zip_image_names(zip_path)
    if not image_names:
        raise ValueError("ZIP không có file ảnh")

    headers = header_map(ws)
    source_col = find_header(headers, "Ảnh nguồn trong ZIP", "Ảnh nguồn")
    if source_col is None:
        raise ValueError("Không tìm thấy cột Ảnh nguồn")

    doctor_cols = [headers.get(name) for name in DOCTOR_HEADERS]
    has_four_columns = all(doctor_cols)
    raw_team_col = find_header(headers, "Kíp phẫu thuật (đọc theo chuỗi)", "Bác sĩ phẫu thuật")

    # Xóa các cột không còn dùng. Xóa từ phải sang trái để không lệch chỉ số.
    # "Bác sĩ phẫu thuật" chỉ xóa khi đã có sẵn 4 cột tách riêng (lúc đó cột
    # chuỗi gốc mới thành thừa) — nếu chưa có 4 cột thì đây chính là cột
    # nguồn cần đọc để tách bên dưới, xóa sớm sẽ mất dữ liệu.
    remove_cols = []
    for title in ("Bác sĩ gây mê", "Ghi chú"):
        col = headers.get(title)
        if col:
            remove_cols.append(col)
    if has_four_columns:
        col = headers.get("Bác sĩ phẫu thuật")
        if col:
            remove_cols.append(col)
    for col in sorted(set(remove_cols), reverse=True):
        ws.delete_cols(col, 1)

    headers = header_map(ws)
    source_col = find_header(headers, "Ảnh nguồn trong ZIP", "Ảnh nguồn")
    doctor_cols = [headers.get(name) for name in DOCTOR_HEADERS]
    has_four_columns = all(doctor_cols)

    if not has_four_columns:
        raw_team_col = find_header(headers, "Kíp phẫu thuật (đọc theo chuỗi)", "Bác sĩ phẫu thuật")
        if raw_team_col is None:
            raise ValueError("Không tìm thấy chuỗi bác sĩ để tách thành 4 cột")
        raw_values = [ws.cell(row, raw_team_col).value for row in range(6, ws.max_row + 1)]
        warning_values = []
        doctor_values = []
        for raw in raw_values:
            doctors, warning = split_doctors(raw)
            doctor_values.append(doctors)
            warning_values.append(warning)

        ws.delete_cols(raw_team_col, 1)
        ws.insert_cols(raw_team_col, 4)
        for offset, title in enumerate(DOCTOR_HEADERS):
            target = raw_team_col + offset
            ws.cell(5, target).value = title
            copy_basic_style(ws, raw_team_col + 4, target)
        for idx, doctors in enumerate(doctor_values, start=6):
            for offset, value in enumerate(doctors):
                ws.cell(idx, raw_team_col + offset).value = value
        headers = header_map(ws)
        source_col = find_header(headers, "Ảnh nguồn trong ZIP", "Ảnh nguồn")
        doctor_cols = [headers[name] for name in DOCTOR_HEADERS]
    else:
        warning_values = []
        for row in range(6, ws.max_row + 1):
            current = [ws.cell(row, col).value for col in doctor_cols]
            cleaned = []
            warning = ""
            for value in current:
                text = str(value or "").strip()
                if " / " in text:
                    text = text.split(" / ", 1)[0].strip()
                    warning = "CẢNH BÁO: Đã bỏ tên phía sau BS4. Kiểm tra ảnh nguồn."
                cleaned.append(text or "-")
            while len(cleaned) < 4:
                cleaned.append("-")
            for col, value in zip(doctor_cols, cleaned):
                ws.cell(row, col).value = value
            if sum(value != "-" for value in cleaned) >= 4 and not warning:
                warning = "CẢNH BÁO: PTV có thể được ghi tiếp trên 2 dòng. Kiểm tra ảnh nguồn."
            warning_values.append(warning)

    # Cập nhật tên ảnh gốc trong ZIP.
    for row in range(6, ws.max_row + 1):
        cell = ws.cell(row, source_col)
        cell.value = original_zip_name(cell.value, image_names)
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    headers = header_map(ws)
    warning_col = headers.get("Cảnh báo")
    if warning_col is None:
        warning_col = ws.max_column + 1
        ws.cell(5, warning_col).value = "Cảnh báo"
        copy_basic_style(ws, source_col, warning_col)
    for row, warning in enumerate(warning_values, start=6):
        ws.cell(row, warning_col).value = warning
        ws.cell(row, warning_col).alignment = Alignment(vertical="center", wrap_text=True)

    # Kích thước dễ kiểm tra.
    for col in doctor_cols:
        ws.column_dimensions[ws.cell(5, col).column_letter].width = 20
    ws.column_dimensions[ws.cell(5, source_col).column_letter].width = 42
    ws.column_dimensions[ws.cell(5, warning_col).column_letter].width = 48

    # Cập nhật tab đối chiếu ảnh nếu có.
    if "Đối chiếu ảnh" in wb.sheetnames:
        audit = wb["Đối chiếu ảnh"]
        audit.cell(4, 1).value = "Tên file ảnh trong ZIP"
        for index, name in enumerate(image_names, start=5):
            audit.cell(index, 1).value = name
            audit.cell(index, 1).alignment = Alignment(vertical="center", wrap_text=True)
        audit.column_dimensions["A"].width = 58

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--excel", required=True, type=Path)
    parser.add_argument("--zip", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--sheet", default="Danh sách sơ bộ")
    args = parser.parse_args()
    normalize_workbook(args.excel, args.zip, args.output, args.sheet)
    print(f"Đã tạo: {args.output}")


if __name__ == "__main__":
    main()
