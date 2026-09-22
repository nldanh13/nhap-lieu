# -*- coding: utf-8 -*-
"""Ghép kết quả OCR thô (ocr_so_phau_thuat_google.run_ocr — từng dòng chữ
kèm tọa độ tương đối) thành bảng dữ liệu dạng "Danh sách sơ bộ", cùng cấu
trúc cột với file bạn vẫn dùng (STT, Ngày ghi sổ, Họ tên người bệnh, Tuổi,
Giới, Chẩn đoán đọc từ sổ, Phương pháp phẫu thuật, Bác sĩ phẫu thuật...).

QUAN TRỌNG — đây là bước suy đoán theo VỊ TRÍ chữ trên ảnh, không phải đọc
hiểu nội dung: các dòng chữ được gom theo hàng (toạ độ Y gần nhau) rồi cắt
theo các mốc toạ độ X trong RANH_GIOI_COT_MAC_DINH bên dưới để đoán dòng
chữ đó thuộc cột nào. Vì mỗi cuốn sổ có thể kẻ cột ở vị trí khác nhau, các
mốc mặc định CHỈ LÀ PHỎNG ĐOÁN ban đầu — cần xem file kết quả thật rồi
chỉnh lại RANH_GIOI_COT_MAC_DINH (hoặc truyền --cot-json) cho khớp đúng
cuốn sổ đang dùng. Mọi dòng đều bị đánh dấu "Cần kiểm tra" vì độ tin cậy
của bước này thấp hơn nhiều so với người đọc trực tiếp ảnh.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


# Mốc ranh giới bên phải của từng cột, tính theo tỉ lệ chiều rộng trang
# (0 = mép trái, 1 = mép phải). Một dòng chữ được xếp vào cột đầu tiên có
# x_max lớn hơn tâm X của dòng chữ đó. ĐÂY LÀ SUY ĐOÁN BAN ĐẦU, chưa hiệu
# chỉnh theo cuốn sổ thật — xem hướng dẫn ở đầu file.
RANH_GIOI_COT_MAC_DINH: list[dict[str, Any]] = [
    {"field": "STT", "x_max": 0.05},
    {"field": "Ngày ghi sổ", "x_max": 0.16},
    {"field": "Họ tên người bệnh", "x_max": 0.34},
    {"field": "Tuổi", "x_max": 0.38},
    {"field": "Giới", "x_max": 0.42},
    {"field": "Chẩn đoán đọc từ sổ", "x_max": 0.62},
    {"field": "Phương pháp phẫu thuật", "x_max": 0.84},
    {"field": "Bác sĩ phẫu thuật", "x_max": 1.01},
]

# Hai dòng chữ được coi là cùng một hàng bảng nếu tâm Y lệch nhau không
# quá ngưỡng này (tỉ lệ chiều cao trang).
NGUONG_CUNG_HANG_MAC_DINH = 0.012

CAC_TRUONG_XUAT = [
    "STT",
    "Ngày ghi sổ",
    "Họ tên người bệnh",
    "Tuổi",
    "Giới",
    "Chẩn đoán đọc từ sổ",
    "Phương pháp phẫu thuật",
    "Bác sĩ phẫu thuật",
    "Trạng thái",
    "Ảnh nguồn",
    "Ghi chú",
]


def _tam_o(box: list[dict[str, float]]) -> tuple[float, float]:
    xs = [p.get("x", 0.0) for p in box] or [0.0]
    ys = [p.get("y", 0.0) for p in box] or [0.0]
    return sum(xs) / len(xs), sum(ys) / len(ys)


def _gom_hang(lines: list[dict[str, Any]], nguong_y: float) -> list[list[dict[str, Any]]]:
    """Gom các dòng chữ có tâm Y gần nhau thành một hàng bảng."""
    items = []
    for line in lines:
        cx, cy = _tam_o(line.get("box") or [])
        items.append({**line, "_cx": cx, "_cy": cy})
    items.sort(key=lambda x: x["_cy"])

    hangs: list[list[dict[str, Any]]] = []
    for item in items:
        if hangs and abs(item["_cy"] - hangs[-1][-1]["_cy"]) <= nguong_y:
            hangs[-1].append(item)
        else:
            hangs.append([item])

    for hang in hangs:
        hang.sort(key=lambda x: x["_cx"])
    return hangs


def _gan_cot(hang: list[dict[str, Any]], ranh_gioi: list[dict[str, Any]]) -> dict[str, str]:
    ket_qua: dict[str, list[str]] = {b["field"]: [] for b in ranh_gioi}
    for item in hang:
        cx = item["_cx"]
        for bound in ranh_gioi:
            if cx <= bound["x_max"]:
                ket_qua[bound["field"]].append(item.get("text", ""))
                break
        else:
            ket_qua[ranh_gioi[-1]["field"]].append(item.get("text", ""))
    return {field: " ".join(v).strip() for field, v in ket_qua.items()}


def phan_tich(
    ocr_payload: dict[str, Any],
    ranh_gioi_cot: list[dict[str, Any]] | None = None,
    nguong_cung_hang: float = NGUONG_CUNG_HANG_MAC_DINH,
    bo_qua_hang_dau_tren_moi_anh: int = 0,
) -> list[dict[str, Any]]:
    """Trả về danh sách dòng đã đoán cột, mỗi dòng gắn kèm tên ảnh nguồn."""
    ranh_gioi = ranh_gioi_cot or RANH_GIOI_COT_MAC_DINH
    stt = 0
    dong_ra: list[dict[str, Any]] = []

    for anh in ocr_payload.get("results", []):
        ten_anh = anh.get("file", "")
        for page in anh.get("pages", []):
            hangs = _gom_hang(page.get("lines", []), nguong_cung_hang)
            for hang in hangs[bo_qua_hang_dau_tren_moi_anh:]:
                gia_tri = _gan_cot(hang, ranh_gioi)
                # Bỏ qua hàng gần như trống (nhiễu OCR/đường kẻ bảng).
                if not any(v for k, v in gia_tri.items() if k != "STT"):
                    continue
                stt += 1
                dong = {
                    "STT": stt,
                    "Ngày ghi sổ": gia_tri.get("Ngày ghi sổ", ""),
                    "Họ tên người bệnh": gia_tri.get("Họ tên người bệnh", ""),
                    "Tuổi": gia_tri.get("Tuổi", ""),
                    "Giới": gia_tri.get("Giới", ""),
                    "Chẩn đoán đọc từ sổ": gia_tri.get("Chẩn đoán đọc từ sổ", ""),
                    "Phương pháp phẫu thuật": gia_tri.get("Phương pháp phẫu thuật", ""),
                    "Bác sĩ phẫu thuật": gia_tri.get("Bác sĩ phẫu thuật", ""),
                    "Trạng thái": "Cần kiểm tra",
                    "Ảnh nguồn": ten_anh,
                    "Ghi chú": "Tự tách cột theo tọa độ — chưa được người/AI đọc hiểu nội dung, luôn xem ảnh gốc trước khi dùng.",
                }
                dong_ra.append(dong)

    return dong_ra


def xuat_excel(dong_ra: list[dict[str, Any]], output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Danh sách sơ bộ"

    ws.cell(2, 1).value = "DANH SÁCH PHẪU THUẬT — TỰ TÁCH CỘT TỪ OCR (CẦN KIỂM TRA LẠI TOÀN BỘ)"
    ws.cell(3, 1).value = "Tổng dòng"
    ws.cell(3, 2).value = len(dong_ra)
    ws.cell(3, 4).value = "Cần kiểm"
    ws.cell(3, 5).value = len(dong_ra)
    ws.cell(3, 7).value = (
        "Bản tự tách cột theo tọa độ chữ trên ảnh (chưa ai đọc hiểu nội dung). "
        "Độ tin cậy thấp hơn nhiều so với đọc trực tiếp — đối chiếu ảnh nguồn trước khi dùng."
    )

    header_row = 5
    for col_idx, field in enumerate(CAC_TRUONG_XUAT, start=1):
        ws.cell(header_row, col_idx).value = field

    for row_offset, dong in enumerate(dong_ra, start=1):
        for col_idx, field in enumerate(CAC_TRUONG_XUAT, start=1):
            ws.cell(header_row + row_offset, col_idx).value = dong.get(field, "")

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[header_row]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{ws.cell(header_row, len(CAC_TRUONG_XUAT)).column_letter}{header_row + len(dong_ra)}"

    widths = [6, 14, 26, 8, 8, 32, 32, 20, 14, 16, 55]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(header_row, idx).column_letter].width = width

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Tự tách cột từ OCR thô theo tọa độ (cần hiệu chỉnh theo cuốn sổ thật).")
    parser.add_argument("--ocr-json", required=True, type=Path, dest="ocr_json")
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--cot-json", type=Path, dest="cot_json", help="File JSON tùy chỉnh ranh giới cột, ghi đè mặc định.")
    return parser


def main() -> None:
    args = build_parser().parse_args()
    ocr_payload = json.loads(args.ocr_json.read_text(encoding="utf-8"))
    ranh_gioi = None
    if args.cot_json:
        ranh_gioi = json.loads(args.cot_json.read_text(encoding="utf-8"))
    dong_ra = phan_tich(ocr_payload, ranh_gioi)
    xuat_excel(dong_ra, args.output)
    print(f"Đã tách được {len(dong_ra)} dòng (CẦN KIỂM TRA LẠI TOÀN BỘ). File: {args.output}")


if __name__ == "__main__":
    main()
