from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from cau_hinh_cot import COT_NHAN_SU_SO_PHAU_THUAT, COT_NHAN_SU_SO_THU_THUAT
from cau_hinh_file import NAM_DU_LIEU, THANG_DA_CHON_THEO_SO, OUTPUT_FILE_THEO_SO
from cau_hinh_phuong_phap import NHOM_PHAN_LOAI_THEO_THU_TU
from utils_text import normalize_text
from xu_ly_du_lieu import make_summary, process_file


# =========================
# CẤU HÌNH CHẠY THƯ MỤC THEOSO
# =========================

BASE_DIR = Path(__file__).resolve().parent
THU_MUC_THEO_SO = BASE_DIR / "TheoSo"

# Các file Excel bắt đầu bằng các tiền tố này sẽ bị bỏ qua.
BO_QUA_FILE_BAT_DAU_BANG = ["~$", "thong_ke_"]

# Cột dùng để chia tháng trong sổ HIS.
# Hai file Sổ Thủ Thuật và Sổ Phẫu Thuật hiện đều có cột "Thời gian thực hiện".
COT_THOI_GIAN_THUC_HIEN_CO_THE = [
    "Thời gian thực hiện",
    "Thoi gian thuc hien",
    "Thời gian bắt đầu",
    "Thoi gian bat dau",
    "Ngày thực hiện",
    "Ngay thuc hien",
]

# Quy tắc tách nhóm Phục hồi chức năng/VLTL trong thư mục TheoSo.
# Nếu khoa chỉ định là Khoa Khám Bệnh => VLTL Ngoại trú.
# Các khoa còn lại => VLTL Nội trú.
KHOA_KHAM_BENH_DE_XEP_VLTL_NGOAI_TRU = "Khoa Khám Bệnh"
NHOM_VLTL_NGOAI_TRU = "VLTL Ngoại trú"
NHOM_VLTL_NOI_TRU = "VLTL Nội trú"
NHOM_VLTL_THEO_THU_TU = [NHOM_VLTL_NGOAI_TRU, NHOM_VLTL_NOI_TRU]
NHOM_VLTL_GOC_CAN_TACH = {"vltl", "phuc hoi chuc nang"}


def _is_vltl_goc(value):
    """Nhận diện nhóm gốc cần tách tiếp thành VLTL Ngoại trú/Nội trú."""
    return normalize_text(value).strip() in NHOM_VLTL_GOC_CAN_TACH


def _configured_groups_for_theoso():
    """
    Thứ tự nhóm khi xuất TheoSo.
    Nếu cấu hình có nhóm Phục hồi chức năng/VLTL thì thay bằng:
    - VLTL Ngoại trú
    - VLTL Nội trú
    """
    result = []
    inserted_vltl_split = False

    for group_name in NHOM_PHAN_LOAI_THEO_THU_TU:
        group_text = str(group_name).strip()
        if not group_text:
            continue

        if _is_vltl_goc(group_text):
            if not inserted_vltl_split:
                for vltl_group in NHOM_VLTL_THEO_THU_TU:
                    if vltl_group not in result:
                        result.append(vltl_group)
                inserted_vltl_split = True
            continue

        if group_text not in result:
            result.append(group_text)

    for vltl_group in NHOM_VLTL_THEO_THU_TU:
        if vltl_group not in result:
            result.append(vltl_group)

    return result


# =========================
# TÌM FILE TRONG THƯ MỤC THEOSO
# =========================

def _is_excel_file(path: Path) -> bool:
    return path.is_file() and path.suffix.lower() in {".xlsx", ".xls"}


def _is_skip_file(path: Path) -> bool:
    return any(path.name.startswith(prefix) for prefix in BO_QUA_FILE_BAT_DAU_BANG)


def _detect_source_type(file_path: Path):
    """Nhận diện file trong thư mục TheoSo."""
    name = normalize_text(file_path.stem).replace(" ", "")

    # Sổ Thủ Thuật: SoThuThuat_...
    if "sothuthuat" in name or "thuthuat" in name:
        return "Sổ Thủ Thuật"

    # Sổ Phẫu Thuật: PTTT_SoPhauThuat_... hoặc SoPhauThuat_...
    if "sophauthuat" in name or "phauthuat" in name or name.startswith("pttt"):
        return "Sổ Phẫu Thuật"

    return None


def get_file_theo_so_list():
    """Trả về danh sách file HIS trong thư mục TheoSo."""
    if not THU_MUC_THEO_SO.exists():
        raise FileNotFoundError(f"Không tìm thấy thư mục TheoSo: {THU_MUC_THEO_SO}")

    result = []
    for file_path in sorted(THU_MUC_THEO_SO.iterdir()):
        if not _is_excel_file(file_path):
            continue
        if _is_skip_file(file_path):
            continue

        source_type = _detect_source_type(file_path)
        if source_type is None:
            print(f"Bỏ qua file chưa nhận diện được loại sổ: {file_path.name}")
            continue

        result.append((file_path, source_type))

    return result


# =========================
# CHIA THÁNG THEO THỜI GIAN THỰC HIỆN
# =========================

def _find_time_column(df):
    """Tìm cột thời gian thực hiện trong dữ liệu đã lọc."""
    col_map = {normalize_text(c): c for c in df.columns}

    for col_name in COT_THOI_GIAN_THUC_HIEN_CO_THE:
        key = normalize_text(col_name)
        if key in col_map:
            return col_map[key]

    # Dự phòng: nếu có cột nào chứa đủ cụm "thời gian" và "thực hiện" thì dùng.
    for col in df.columns:
        col_norm = normalize_text(col)
        if "thoi gian" in col_norm and "thuc hien" in col_norm:
            return col

    return None


def _parse_datetime_series(series):
    """
    Đọc được các dạng thời gian thường gặp trong HIS:
    - 00:00 05/06/2026
    - 02:15 01/01/2026
    - 2026-01-01 02:45:00
    """
    try:
        return pd.to_datetime(series, format="mixed", dayfirst=True, errors="coerce")
    except TypeError:
        return pd.to_datetime(series, dayfirst=True, errors="coerce")


def _months_to_export():
    """
    Xác định tháng cần xuất theo cấu hình THANG_CAN_CHAY_THEO_SO trong cau_hinh_file.py.
    - THANG_CAN_CHAY_THEO_SO = [] hoặc "tat_ca" => tháng 1 đến tháng 12.
    - THANG_CAN_CHAY_THEO_SO = [6] => chỉ tháng 6.
    """
    if THANG_DA_CHON_THEO_SO is None:
        return list(range(1, 13))
    return sorted(int(x) for x in THANG_DA_CHON_THEO_SO)


def _add_execution_month_columns(df_filtered):
    df = df_filtered.copy()
    time_col = _find_time_column(df)

    if not time_col:
        available_cols = ", ".join(map(str, df.columns.tolist()))
        raise ValueError(
            "Không tìm thấy cột thời gian thực hiện để chia tháng. "
            "Hãy kiểm tra file sổ hoặc thêm tên cột vào COT_THOI_GIAN_THUC_HIEN_CO_THE trong chay_theo_so.py.\n"
            f"Các cột hiện có: {available_cols}"
        )

    time_values = _parse_datetime_series(df[time_col])
    df["Cột thời gian dùng chia tháng"] = time_col
    df["_thoi_gian_thuc_hien_parsed"] = time_values
    df["_nam_so"] = time_values.dt.year
    df["_thang_so"] = time_values.dt.month

    return df, time_col


# =========================
# TÁCH VLTL NGOẠI TRÚ / NỘI TRÚ TRONG THEOSO
# =========================

def _is_khoa_kham_benh(value):
    """Kiểm tra khoa chỉ định có phải Khoa Khám Bệnh hay không."""
    return normalize_text(KHOA_KHAM_BENH_DE_XEP_VLTL_NGOAI_TRU) in normalize_text(value)


def apply_vltl_classification_by_khoa(df_filtered):
    """
    Với các dòng đã được phân loại là Phục hồi chức năng hoặc VLTL:
    - Khoa Khám Bệnh => VLTL Ngoại trú
    - Khoa khác => VLTL Nội trú

    Lưu ý: không tự chuyển toàn bộ dòng Chưa phân loại sang VLTL.
    Các dòng chưa phân loại vẫn giữ nguyên để xuất sheet Chưa phân loại cho bạn rà soát.
    """
    if df_filtered.empty or "Nhóm phân loại" not in df_filtered.columns:
        return df_filtered, {
            "vltl_ngoai_tru": 0,
            "vltl_noi_tru": 0,
            "tong_vltl": 0,
        }

    df = df_filtered.copy()
    nhom_series = df["Nhóm phân loại"].fillna("").astype(str)
    mask_vltl = nhom_series.map(_is_vltl_goc)

    if not mask_vltl.any():
        return df, {
            "vltl_ngoai_tru": 0,
            "vltl_noi_tru": 0,
            "tong_vltl": 0,
        }

    if "Khoa chỉ định dùng lọc" in df.columns:
        khoa_series = df["Khoa chỉ định dùng lọc"]
    else:
        khoa_series = pd.Series("", index=df.index)

    mask_ngoai_tru = mask_vltl & khoa_series.map(_is_khoa_kham_benh)
    mask_noi_tru = mask_vltl & ~mask_ngoai_tru

    df.loc[mask_ngoai_tru, "Nhóm phân loại"] = NHOM_VLTL_NGOAI_TRU
    df.loc[mask_noi_tru, "Nhóm phân loại"] = NHOM_VLTL_NOI_TRU

    # Ghi lại lý do để khi cần mở chi tiết vẫn biết dòng này được tách theo khoa.
    if "Kiểu phân loại" in df.columns:
        current_value = df.loc[mask_vltl, "Kiểu phân loại"].fillna("").astype(str)
        df.loc[mask_vltl, "Kiểu phân loại"] = current_value.where(
            current_value.eq(""),
            current_value + " | "
        ) + "tach_vltl_theo_khoa_chi_dinh"

    return df, {
        "vltl_ngoai_tru": int(mask_ngoai_tru.sum()),
        "vltl_noi_tru": int(mask_noi_tru.sum()),
        "tong_vltl": int(mask_vltl.sum()),
    }


# =========================
# XUẤT 1 SHEET THEO NHÓM / THEO THÁNG
# =========================

def _prepare_df_by_selected_months(df_filtered):
    """Chuẩn bị dữ liệu có cột năm/tháng và lọc đúng năm, đúng tháng đang chọn."""
    if df_filtered.empty:
        return df_filtered.copy(), "", 0, 0

    df, time_col = _add_execution_month_columns(df_filtered)
    months = _months_to_export()

    invalid_time_count = int(df["_thoi_gian_thuc_hien_parsed"].isna().sum())

    # Chỉ lấy đúng năm đang cấu hình và đúng tháng được chọn.
    df_in_year = df[df["_nam_so"] == NAM_DU_LIEU].copy()
    df_in_months = df_in_year[df_in_year["_thang_so"].isin(months)].copy()
    outside_selected_count = len(df) - len(df_in_months)

    return df_in_months, time_col, invalid_time_count, outside_selected_count


def make_group_month_table(df_filtered):
    """Tạo bảng: Theo nhóm phân loại | Tháng 1 | Tháng 2 | ... | Tổng số."""
    months = _months_to_export()
    month_columns = [f"Tháng {m}" for m in months]

    if df_filtered.empty:
        table = pd.DataFrame({"Theo nhóm phân loại": _configured_groups_for_theoso()})
        for col in month_columns:
            table[col] = 0
        table["Tổng số"] = 0
        return table, "", 0, 0, 0

    df_in_months, time_col, invalid_time_count, outside_selected_count = _prepare_df_by_selected_months(df_filtered)

    if df_in_months.empty:
        table = pd.DataFrame({"Theo nhóm phân loại": _configured_groups_for_theoso()})
        for col in month_columns:
            table[col] = 0
        table["Tổng số"] = 0
        return table, time_col, invalid_time_count, outside_selected_count, 0

    pivot = (
        df_in_months
        .groupby(["Nhóm phân loại", "_thang_so"], dropna=False)
        .size()
        .unstack(fill_value=0)
    )

    for month in months:
        if month not in pivot.columns:
            pivot[month] = 0

    pivot = pivot[months]

    configured_groups = _configured_groups_for_theoso()

    extra_groups = [g for g in pivot.index.tolist() if g not in configured_groups]
    if "Chưa phân loại" in extra_groups:
        extra_groups = [g for g in extra_groups if g != "Chưa phân loại"] + ["Chưa phân loại"]

    ordered_groups = configured_groups + extra_groups
    pivot = pivot.reindex(ordered_groups, fill_value=0)

    result = pivot.reset_index().rename(columns={"Nhóm phân loại": "Theo nhóm phân loại"})
    result.columns = ["Theo nhóm phân loại"] + month_columns
    result["Tổng số"] = result[month_columns].sum(axis=1)

    total_row = {"Theo nhóm phân loại": "Tổng cộng"}
    for col in month_columns:
        total_row[col] = int(result[col].sum())
    total_row["Tổng số"] = int(result["Tổng số"].sum())

    chua_phan_loai_count = int((df_in_months["Nhóm phân loại"] == "Chưa phân loại").sum())

    return (
        pd.concat([result, pd.DataFrame([total_row])], ignore_index=True),
        time_col,
        invalid_time_count,
        outside_selected_count,
        chua_phan_loai_count,
    )


def make_unclassified_detail_table(df_filtered):
    """
    Tạo sheet chi tiết các dòng chưa phân loại trong đúng năm/tháng đang chạy.
    Sheet này giúp biết dòng chưa phân loại nằm ở file nào, sổ nào, tháng nào,
    phương pháp nào để bổ sung vào cau_hinh_phuong_phap.py.
    """
    output_columns = [
        "File nguồn",
        "Nguồn",
        "TT",
        "Thời gian thực hiện",
        "Năm",
        "Tháng",
        "Cột phương pháp",
        "Phương pháp dùng phân loại",
        "Khoa chỉ định dùng lọc",
        "Bác sĩ khớp",
        "Điều dưỡng khớp",
        "KTV khớp",
        "Kiểu dò nhân sự",
        "Cột thời gian dùng chia tháng",
        "Gợi ý cần thêm vào cau_hinh_phuong_phap.py",
    ]

    if df_filtered.empty:
        return pd.DataFrame(columns=output_columns)

    df_in_months, time_col, _, _ = _prepare_df_by_selected_months(df_filtered)
    if df_in_months.empty:
        return pd.DataFrame(columns=output_columns)

    df_unclassified = df_in_months[df_in_months["Nhóm phân loại"] == "Chưa phân loại"].copy()
    if df_unclassified.empty:
        return pd.DataFrame(columns=output_columns)

    df_unclassified["Năm"] = df_unclassified["_nam_so"].astype("Int64")
    df_unclassified["Tháng"] = df_unclassified["_thang_so"].astype("Int64")

    if time_col and time_col in df_unclassified.columns:
        df_unclassified["Thời gian thực hiện"] = df_unclassified[time_col]
    elif "Thời gian thực hiện" not in df_unclassified.columns:
        df_unclassified["Thời gian thực hiện"] = ""

    if "TT" not in df_unclassified.columns:
        df_unclassified["TT"] = ""

    # Cột gợi ý: copy nhanh tên phương pháp để dán vào cau_hinh_phuong_phap.py.
    df_unclassified["Gợi ý cần thêm vào cau_hinh_phuong_phap.py"] = df_unclassified["Phương pháp dùng phân loại"].map(
        lambda x: f'{{"phuong_phap": "{str(x).strip()}", "nhom": "", "kieu_khop": "chinh_xac"}},'
    )

    # Đảm bảo các cột chính luôn tồn tại.
    for col in output_columns:
        if col not in df_unclassified.columns:
            df_unclassified[col] = ""

    # Thêm một số cột định danh thường gặp nếu có trong sổ để dễ dò lại.
    optional_cols = []
    possible_optional_names = [
        "Mã y tế",
        "Mã Y tế",
        "Mã người bệnh",
        "Mã BN",
        "Họ tên người bệnh",
        "Họ tên NB",
        "Người bệnh",
        "Tên người bệnh",
        "Năm sinh",
        "Tuổi",
        "Giới tính",
        "Chẩn đoán",
        "Chẩn đoán trước phẫu thuật",
        "Chẩn đoán trước thủ thuật",
        "Phương pháp phẫu thuật",
        "Phương pháp thủ thuật",
    ]
    for col in possible_optional_names:
        if col in df_unclassified.columns and col not in output_columns and col not in optional_cols:
            optional_cols.append(col)

    final_columns = output_columns[:8] + optional_cols + output_columns[8:]
    final_columns = [col for i, col in enumerate(final_columns) if col not in final_columns[:i]]

    return df_unclassified[final_columns].sort_values(
        by=["Tháng", "File nguồn", "Nguồn", "TT"],
        kind="stable",
    ).reset_index(drop=True)


def _format_sheet(writer, sheet_name):
    ws = writer.book[sheet_name]

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    total_fill = PatternFill("solid", fgColor="FFF2CC")
    header_font = Font(bold=True)
    total_font = Font(bold=True)

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=2, max_col=ws.max_column):
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.number_format = "#,##0"

    if ws.max_row >= 2:
        for cell in ws[ws.max_row]:
            cell.font = total_font
            cell.fill = total_fill

    ws.column_dimensions["A"].width = 45
    for col_cells in ws.iter_cols(min_col=2, max_col=ws.max_column):
        ws.column_dimensions[col_cells[0].column_letter].width = 13


def _format_detail_sheet(writer, sheet_name):
    ws = writer.book[sheet_name]

    header_fill = PatternFill("solid", fgColor="FCE4D6")
    header_font = Font(bold=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    # Độ rộng cột dễ đọc, không quá rộng.
    width_by_header = {
        "File nguồn": 36,
        "Nguồn": 18,
        "TT": 10,
        "Thời gian thực hiện": 20,
        "Năm": 10,
        "Tháng": 10,
        "Cột phương pháp": 24,
        "Phương pháp dùng phân loại": 45,
        "Khoa chỉ định dùng lọc": 28,
        "Bác sĩ khớp": 24,
        "Điều dưỡng khớp": 24,
        "KTV khớp": 24,
        "Kiểu dò nhân sự": 22,
        "Cột thời gian dùng chia tháng": 26,
        "Gợi ý cần thêm vào cau_hinh_phuong_phap.py": 70,
    }

    for col_cells in ws.iter_cols(min_row=1, max_row=1):
        header = str(col_cells[0].value or "")
        width = width_by_header.get(header, 22)
        ws.column_dimensions[col_cells[0].column_letter].width = width


def export_theo_so_excel(output_file, df_filtered):
    table, time_col, invalid_time_count, outside_selected_count, chua_phan_loai_count = make_group_month_table(df_filtered)
    detail_table = make_unclassified_detail_table(df_filtered)

    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        table.to_excel(writer, sheet_name="Theo nhóm tháng", index=False)
        _format_sheet(writer, "Theo nhóm tháng")

        if not detail_table.empty:
            detail_table.to_excel(writer, sheet_name="Chưa phân loại", index=False)
            _format_detail_sheet(writer, "Chưa phân loại")

    return {
        "time_col": time_col,
        "invalid_time_count": invalid_time_count,
        "outside_selected_count": outside_selected_count,
        "chua_phan_loai_count": chua_phan_loai_count,
        "created_unclassified_sheet": not detail_table.empty,
        "total_exported": int(table.loc[table["Theo nhóm phân loại"] == "Tổng cộng", "Tổng số"].iloc[0]) if not table.empty else 0,
    }


# =========================
# MAIN
# =========================

def main():
    file_items = get_file_theo_so_list()

    if not file_items:
        raise ValueError(
            "Không tìm thấy file Sổ Thủ Thuật hoặc Sổ Phẫu Thuật trong thư mục TheoSo. "
            "Hãy kiểm tra lại tên file Excel."
        )

    print("Đang chạy chế độ TheoSo")
    print("Thư mục dữ liệu:", THU_MUC_THEO_SO)
    print("Năm dữ liệu:", NAM_DU_LIEU)
    print("Tháng cần chạy TheoSo:", "Tất cả" if THANG_DA_CHON_THEO_SO is None else THANG_DA_CHON_THEO_SO)
    print("Số file đầu vào:", len(file_items))
    print("Phân loại dựa vào: cột Phương pháp thủ thuật / Phương pháp phẫu thuật")
    print("Danh sách phân loại nằm trong: cau_hinh_phuong_phap.py -> PHAN_LOAI_PHUONG_PHAP_UU_TIEN")
    print("Dò nhân sự TheoSo: chỉ theo họ tên đầy đủ, không dùng bí danh viết tắt")
    print("Có hỗ trợ nhóm KTV nếu bạn khai báo KTV trong danh_sach_nhan_su.py")
    print("Tách nhóm Phục hồi chức năng/VLTL: Khoa Khám Bệnh => VLTL Ngoại trú, khoa khác => VLTL Nội trú")

    all_filtered = []
    all_roles = []

    for file_path, source_type in file_items:
        print(f"- {file_path.name} => {source_type}")

        if source_type == "Sổ Thủ Thuật":
            staff_columns = COT_NHAN_SU_SO_THU_THUAT
        elif source_type == "Sổ Phẫu Thuật":
            staff_columns = COT_NHAN_SU_SO_PHAU_THUAT
        else:
            continue

        # File trong thư mục TheoSo là sổ HIS có họ tên đầy đủ.
        # Không dùng bí danh viết tắt để tránh nhầm, ví dụ alias "DUY" khớp sai với "Nguyễn Duy Linh".
        df_filtered, df_roles = process_file(file_path, source_type, staff_columns, use_alias=False)

        if not df_filtered.empty:
            df_filtered.insert(0, "File nguồn", file_path.name)
        if not df_roles.empty:
            df_roles.insert(0, "File nguồn", file_path.name)

        all_filtered.append(df_filtered)
        all_roles.append(df_roles)

    df_filtered_all = pd.concat(all_filtered, ignore_index=True) if all_filtered else pd.DataFrame()
    df_roles_all = pd.concat(all_roles, ignore_index=True) if all_roles else pd.DataFrame()

    # Sau khi phân loại theo danh sách trong cau_hinh_phuong_phap.py,
    # các dòng thuộc nhóm Phục hồi chức năng/VLTL sẽ được tách thành
    # VLTL Ngoại trú hoặc VLTL Nội trú theo khoa chỉ định.
    # Các dòng Chưa phân loại vẫn giữ nguyên để xuất sheet rà soát.
    df_filtered_all, vltl_info = apply_vltl_classification_by_khoa(df_filtered_all)

    # Gọi make_summary để giữ kiểm tra logic cũ, dù file xuất chỉ dùng bảng Theo nhóm tháng.
    _ = make_summary(df_filtered_all, df_roles_all)

    export_info = export_theo_so_excel(OUTPUT_FILE_THEO_SO, df_filtered_all)

    print("Chia tháng dựa vào cột:", export_info["time_col"] or "Không xác định")
    print("Đã xuất file:", OUTPUT_FILE_THEO_SO)
    print("Tổng dòng lọc được trước khi chia tháng:", len(df_filtered_all))
    print("Tổng dòng đã đưa vào bảng tháng:", export_info["total_exported"])
    print("Số dòng không đọc được thời gian thực hiện:", export_info["invalid_time_count"])
    print("Số dòng ngoài năm/tháng đang chọn:", export_info["outside_selected_count"])
    print("Đã tách Phục hồi chức năng/VLTL:", vltl_info["tong_vltl"])
    print("  - VLTL Ngoại trú:", vltl_info["vltl_ngoai_tru"])
    print("  - VLTL Nội trú:", vltl_info["vltl_noi_tru"])
    print("Số dòng chưa phân loại trong năm/tháng đang chạy:", export_info["chua_phan_loai_count"])
    print("Đã tạo sheet Chưa phân loại:", "Có" if export_info["created_unclassified_sheet"] else "Không")
    if not df_filtered_all.empty:
        print("Số dòng chưa phân loại trước khi chia tháng:", (df_filtered_all["Nhóm phân loại"] == "Chưa phân loại").sum())
        print("Tổng lượt nhân sự xuất hiện:", len(df_roles_all))


if __name__ == "__main__":
    main()
