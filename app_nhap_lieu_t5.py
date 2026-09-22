# -*- coding: utf-8 -*-
"""
Giao diện nhập liệu T5 cho file PM khoa CTCH.

Chạy:
    streamlit run app_nhap_lieu_t5.py

Mục tiêu:
- Xem/sửa nhanh dữ liệu trong các sheet phauthuat, thu thuat, tieuphau.
- Thêm dòng mới trước dòng Tổng Cộng để không phá phần chữ ký/ô merge bên dưới.
- Chạy các chức năng tự động đã có: bổ sung phụ mổ, chuyển thủ thuật sang tiểu phẫu,
  điền BS bằng bí danh từ Sổ Phẫu Thuật/Sổ Thủ Thuật.
- Xuất file Excel mới, giữ file gốc.
"""

from __future__ import annotations

import io
import os
import re
import shutil
import sys
import traceback
import unicodedata
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.worksheet.cell_range import CellRange

# Đảm bảo import được các module trong cùng thư mục khi chạy bằng Streamlit.
THIS_DIR = Path(__file__).resolve().parent
if str(THIS_DIR) not in sys.path:
    sys.path.insert(0, str(THIS_DIR))

try:
    from cau_hinh_file import BASE_DIR
except Exception:
    BASE_DIR = THIS_DIR

DEFAULT_PM_FILE = BASE_DIR / "2026" / "PM khoa CTCH T5.xlsx"
DEFAULT_PM_PHU_MO_FILE = BASE_DIR / "2026" / "PM khoa CTCH T5_da_bo_sung_phu_mo.xlsx"
DEFAULT_PM_CHUYEN_FILE = BASE_DIR / "2026" / "PM khoa CTCH T5_da_chuyen_tieu_phau.xlsx"
DEFAULT_PM_CHUYEN_BS_FILE = BASE_DIR / "2026" / "PM khoa CTCH T5_da_chuyen_tieu_phau_da_dien_BS.xlsx"
DEFAULT_OUTPUT_DIR = BASE_DIR / "2026"
DEFAULT_REPORT_DIR = BASE_DIR

SHEET_PRESETS = {
    "phauthuat": {
        "label": "Phẫu thuật",
        "required": ["STT", "NGÀY", "Họ và tên bệnh nhân", "Chẩn đoán và phương pháp phẫu thuật"],
        "quick_fields": [
            "NGÀY",
            "Họ và tên bệnh nhân",
            "Tuổi",
            "Chẩn đoán và phương pháp phẫu thuật",
            "PTV chính",
            "Phụ mổ 1",
            "Phụ mổ 2",
            "Phụ mổ 3",
        ],
        "missing_cols": ["PTV chính", "Phụ mổ 1", "Phụ mổ 2", "Phụ mổ 3"],
    },
    "thu thuat": {
        "label": "Thủ thuật",
        "required": ["STT", "Ngày", "Họ và tên", "Tuổi", "Tên CLS", "Số Lượng", "Thành tiền"],
        "quick_fields": ["Ngày", "Họ và tên", "Tuổi", "Tên CLS", "Số Lượng", "Thành tiền"],
        "missing_cols": [],
    },
    "tieuphau": {
        "label": "Tiểu phẫu",
        "required": ["STT", "Ngày", "Họ và tên", "Tuổi", "Tên CLS", "Số Lượng", "Thành tiền"],
        "quick_fields": ["Ngày", "Họ và tên", "Tuổi", "Tên CLS", "Số Lượng", "Thành tiền", "Bác Sĩ", "Điều Dưỡng"],
        "missing_cols": ["Bác Sĩ"],
    },
}

# Các tên cột phụ có thể gặp trong file.
COLUMN_ALIASES = {
    "Bác Sĩ": ["Bác Sĩ", "BS", "Bác sĩ", "Bac Si", "Bac si"],
    "Điều Dưỡng": ["Điều Dưỡng", "Điều dưỡng", "DD", "ĐD", "Dieu Duong", "Dieu duong"],
    "Số Lượng": ["Số Lượng", "Số lượng", "SL", "So Luong", "So luong"],
}


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    if not text or text in {"nan", "none"}:
        return ""
    text = unicodedata.normalize("NFD", text)
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    text = text.replace("đ", "d").replace("Đ", "D")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def normalize_header(value: Any) -> str:
    return normalize_text(value).replace(" ", "")


def value_is_blank(value: Any) -> bool:
    if value is None:
        return True
    text = str(value).strip()
    return text == "" or text.lower() in {"nan", "none", "nat"}


def sheet_key(name: str) -> str:
    return normalize_header(name)


def find_sheet_name(wb, wanted: str) -> str:
    wanted_key = sheet_key(wanted)
    for name in wb.sheetnames:
        if sheet_key(name) == wanted_key:
            return name
    raise ValueError(f"Không tìm thấy sheet '{wanted}'. Các sheet hiện có: {', '.join(wb.sheetnames)}")


def expand_required_columns(required: Iterable[str]) -> list[str]:
    result: list[str] = []
    for col in required:
        result.extend(COLUMN_ALIASES.get(col, [col]))
    return result


def find_header_row_and_map(ws, required_columns: list[str] | None = None) -> tuple[int, dict[str, int], dict[int, str]]:
    """Tìm dòng tiêu đề và map header -> cột.

    Nếu required_columns có alias, chỉ cần match theo tên chuẩn trong required gốc.
    """
    required_columns = required_columns or ["STT"]
    required_keys: set[str] = set()
    for col in required_columns:
        aliases = COLUMN_ALIASES.get(col, [col])
        required_keys.add(normalize_header(aliases[0]))

    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 50)):
        by_key: dict[str, int] = {}
        by_col: dict[int, str] = {}
        alias_seen: set[str] = set()

        for cell in row:
            if value_is_blank(cell.value):
                continue
            raw = str(cell.value).strip()
            key = normalize_header(raw)
            if not key:
                continue
            by_key.setdefault(key, cell.column)
            by_col[cell.column] = raw

            for canonical, aliases in COLUMN_ALIASES.items():
                if key in {normalize_header(x) for x in aliases}:
                    by_key.setdefault(normalize_header(canonical), cell.column)
                    alias_seen.add(normalize_header(canonical))

        # Kiểm tra required theo canonical/alias.
        ok = True
        for col in required_columns:
            keys = [normalize_header(x) for x in COLUMN_ALIASES.get(col, [col])]
            if not any(k in by_key for k in keys):
                ok = False
                break
        if ok and by_col:
            return row[0].row, by_key, by_col

    raise ValueError(f"Không tìm thấy dòng tiêu đề có đủ các cột: {required_columns}")


def col_index(header_map: dict[str, int], col_name: str) -> int | None:
    for alias in COLUMN_ALIASES.get(col_name, [col_name]):
        key = normalize_header(alias)
        if key in header_map:
            return header_map[key]
    return None


def find_total_row(ws, header_row: int) -> int:
    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row):
        for cell in row:
            if normalize_text(cell.value) == "tong cong":
                return cell.row
    return ws.max_row + 1


def row_has_data(ws, row_idx: int, columns: Iterable[int]) -> bool:
    return any(not value_is_blank(ws.cell(row_idx, col).value) for col in columns)


def get_data_range(ws, header_row: int, header_by_col: dict[int, str], total_row: int) -> tuple[int, int]:
    data_cols = sorted(header_by_col.keys())
    last = header_row
    end_limit = total_row - 1 if total_row <= ws.max_row else ws.max_row
    for row_idx in range(end_limit, header_row, -1):
        if row_has_data(ws, row_idx, data_cols):
            last = row_idx
            break
    return header_row + 1, last


def choose_table_columns(header_by_col: dict[int, str]) -> list[tuple[int, str]]:
    cols: list[tuple[int, str]] = []
    used: set[str] = set()
    for col_idx in sorted(header_by_col):
        name = str(header_by_col[col_idx]).strip()
        if not name:
            continue
        base = name
        suffix = 2
        while name in used:
            name = f"{base}_{suffix}"
            suffix += 1
        used.add(name)
        cols.append((col_idx, name))
    return cols


def workbook_to_dataframe(file_path: Path, sheet_wanted: str) -> tuple[pd.DataFrame, dict[str, Any]]:
    wb = load_workbook(file_path, data_only=False)
    sheet_name = find_sheet_name(wb, sheet_wanted)
    ws = wb[sheet_name]
    preset = SHEET_PRESETS.get(sheet_wanted, {})
    header_row, header_map, header_by_col = find_header_row_and_map(ws, preset.get("required"))
    total_row = find_total_row(ws, header_row)
    start_row, last_row = get_data_range(ws, header_row, header_by_col, total_row)
    table_cols = choose_table_columns(header_by_col)

    rows: list[dict[str, Any]] = []
    for row_idx in range(start_row, last_row + 1):
        item: dict[str, Any] = {"__dong_excel": row_idx}
        for col_idx, name in table_cols:
            item[name] = ws.cell(row_idx, col_idx).value
        rows.append(item)

    df = pd.DataFrame(rows)
    if df.empty:
        df = pd.DataFrame(columns=["__dong_excel", *[name for _, name in table_cols]])

    meta = {
        "sheet_name": sheet_name,
        "header_row": header_row,
        "total_row": total_row,
        "start_row": start_row,
        "last_row": last_row,
        "table_cols": table_cols,
        "headers": [name for _, name in table_cols],
        "header_map": header_map,
    }
    return df, meta


def copy_cell_style(src_cell, dst_cell) -> None:
    if src_cell.has_style:
        dst_cell.font = copy(src_cell.font)
        dst_cell.fill = copy(src_cell.fill)
        dst_cell.border = copy(src_cell.border)
        dst_cell.alignment = copy(src_cell.alignment)
        dst_cell.number_format = src_cell.number_format
        dst_cell.protection = copy(src_cell.protection)


def copy_row_style(ws, src_row: int, dst_row: int, max_col: int) -> None:
    ws.row_dimensions[dst_row].height = ws.row_dimensions[src_row].height
    for col_idx in range(1, max_col + 1):
        copy_cell_style(ws.cell(src_row, col_idx), ws.cell(dst_row, col_idx))
        src = ws.cell(src_row, col_idx)
        dst = ws.cell(dst_row, col_idx)
        if isinstance(src.value, str) and src.value.startswith("="):
            try:
                dst.value = Translator(src.value, origin=src.coordinate).translate_formula(dst.coordinate)
            except Exception:
                dst.value = src.value
        else:
            dst.value = None


def clone_merged_ranges(ws) -> list[CellRange]:
    return [CellRange(str(rng)) for rng in ws.merged_cells.ranges]


def shift_merged_ranges_for_insert(ranges: list[CellRange], insert_at: int, amount: int) -> list[CellRange]:
    shifted: list[CellRange] = []
    for rng in ranges:
        new = CellRange(str(rng))
        if new.min_row >= insert_at:
            new.shift(row_shift=amount)
        elif new.max_row >= insert_at:
            new.max_row += amount
        shifted.append(new)
    return shifted


def safe_insert_rows(ws, insert_at: int, amount: int) -> None:
    ranges = clone_merged_ranges(ws)
    for rng in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(rng))
    ws.insert_rows(insert_at, amount)
    for rng in shift_merged_ranges_for_insert(ranges, insert_at, amount):
        ws.merge_cells(str(rng))


def dataframe_to_workbook(
    source_file: Path,
    sheet_wanted: str,
    edited_df: pd.DataFrame,
    output_file: Path,
    allow_add_rows: bool = True,
) -> dict[str, Any]:
    wb = load_workbook(source_file)
    sheet_name = find_sheet_name(wb, sheet_wanted)
    ws = wb[sheet_name]
    preset = SHEET_PRESETS.get(sheet_wanted, {})
    header_row, header_map, header_by_col = find_header_row_and_map(ws, preset.get("required"))
    total_row_before = find_total_row(ws, header_row)
    table_cols = choose_table_columns(header_by_col)
    col_by_name = {name: col_idx for col_idx, name in table_cols}

    updated = 0
    inserted = 0
    skipped_new = 0

    # Chuẩn hóa dataframe từ st.data_editor.
    df = edited_df.copy()
    if "__dong_excel" not in df.columns:
        raise ValueError("Thiếu cột kỹ thuật __dong_excel, không thể ghi file an toàn.")

    existing_rows: list[tuple[int, pd.Series]] = []
    new_rows: list[pd.Series] = []
    for _, row in df.iterrows():
        row_no = row.get("__dong_excel")
        try:
            row_no_int = int(row_no) if not pd.isna(row_no) and str(row_no).strip() else 0
        except Exception:
            row_no_int = 0
        if row_no_int > 0:
            existing_rows.append((row_no_int, row))
        else:
            new_rows.append(row)

    for row_idx, row in existing_rows:
        if row_idx <= header_row or row_idx >= total_row_before:
            continue
        for name, col_idx in col_by_name.items():
            if name in row:
                value = row[name]
                if pd.isna(value):
                    value = None
                ws.cell(row_idx, col_idx).value = value
        updated += 1

    if allow_add_rows and new_rows:
        insert_at = total_row_before
        safe_insert_rows(ws, insert_at, len(new_rows))
        style_row = max(header_row + 1, insert_at - 1)
        for offset, row in enumerate(new_rows):
            dst_row = insert_at + offset
            copy_row_style(ws, style_row, dst_row, ws.max_column)
            for name, col_idx in col_by_name.items():
                if name == "STT":
                    continue
                if name in row:
                    value = row[name]
                    if pd.isna(value):
                        value = None
                    ws.cell(dst_row, col_idx).value = value
            inserted += 1
    else:
        skipped_new = len(new_rows)

    # Cập nhật STT nếu có.
    stt_col = col_index(header_map, "STT")
    if stt_col:
        total_row_after = find_total_row(ws, header_row)
        stt = 1
        for r in range(header_row + 1, total_row_after):
            # Chỉ đánh STT cho dòng có dữ liệu ở các cột khác.
            data_cols = [c for c in col_by_name.values() if c != stt_col]
            if row_has_data(ws, r, data_cols):
                ws.cell(r, stt_col).value = stt
                stt += 1

    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_file)
    return {
        "file_dau_ra": str(output_file),
        "sheet": sheet_name,
        "so_dong_cap_nhat": updated,
        "so_dong_chen_moi": inserted,
        "so_dong_moi_bo_qua": skipped_new,
    }


def save_upload(uploaded_file, save_dir: Path) -> Path:
    save_dir.mkdir(parents=True, exist_ok=True)
    file_path = save_dir / uploaded_file.name
    with open(file_path, "wb") as f:
        f.write(uploaded_file.getbuffer())
    return file_path


def path_input(label: str, default_path: Path, key: str) -> Path:
    return Path(st.text_input(label, value=str(default_path), key=key).strip().strip('"'))


def file_download_button(file_path: Path, label: str) -> None:
    if file_path.exists():
        with open(file_path, "rb") as f:
            st.download_button(
                label=label,
                data=f.read(),
                file_name=file_path.name,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )


def summarize_missing(file_path: Path, sheet_wanted: str) -> pd.DataFrame:
    df, meta = workbook_to_dataframe(file_path, sheet_wanted)
    preset = SHEET_PRESETS.get(sheet_wanted, {})
    missing_cols = preset.get("missing_cols", [])
    if not missing_cols or df.empty:
        return pd.DataFrame()

    cols_existing = [c for c in missing_cols if c in df.columns]
    # Alias Bác Sĩ/BS.
    if sheet_wanted == "tieuphau" and not cols_existing:
        for c in df.columns:
            if normalize_header(c) in {normalize_header(x) for x in COLUMN_ALIASES["Bác Sĩ"]}:
                cols_existing.append(c)
                break

    if not cols_existing:
        return pd.DataFrame()

    mask = False
    for c in cols_existing:
        mask = mask | df[c].apply(value_is_blank)

    view_cols = ["__dong_excel"]
    for wanted in ["NGÀY", "Ngày", "Họ và tên bệnh nhân", "Họ và tên", "Tên CLS", "Chẩn đoán và phương pháp phẫu thuật", *cols_existing]:
        if wanted in df.columns and wanted not in view_cols:
            view_cols.append(wanted)
    return df.loc[mask, view_cols]


def get_workbook_sheets(file_path: Path) -> list[str]:
    wb = load_workbook(file_path, read_only=True, data_only=False)
    return wb.sheetnames


def run_bo_sung_phu_mo(pm_file: Path, output_file: Path) -> dict[str, Any]:
    import bo_sung_phu_mo_t5_tu_so_phau_thuat as mod

    mod.FILE_PM_T5 = pm_file
    mod.FILE_DAU_RA = output_file
    mod.FILE_BAO_CAO = DEFAULT_REPORT_DIR / "bao_cao_bo_sung_phu_mo_T5.xlsx"
    mod.GHI_DE_FILE_GOC = False
    mod.bo_sung_phu_mo()
    return {
        "file_dau_ra": str(output_file),
        "file_bao_cao": str(mod.FILE_BAO_CAO),
    }


def run_chuyen_tieu_phau(pm_file: Path, output_file: Path, xoa_khoi_thu_thuat: bool, chi_dien_bs_trong: bool, danh_sach_cls_text: str) -> dict[str, Any]:
    import chuyen_thu_thuat_sang_tieuphau_t5 as mod

    mod.FILE_NGUON = pm_file
    mod.FILE_DAU_RA = output_file
    mod.FILE_BAO_CAO = DEFAULT_REPORT_DIR / "bao_cao_chuyen_thu_thuat_sang_tieuphau_T5.xlsx"
    mod.XOA_DONG_DA_CHUYEN_KHOI_THU_THUAT = xoa_khoi_thu_thuat
    mod.CHI_DIEN_BS_KHI_DANG_TRONG = chi_dien_bs_trong
    danh_sach = [line.strip() for line in danh_sach_cls_text.splitlines() if line.strip()]
    if danh_sach:
        mod.DANH_SACH_TEN_CLS_CAN_CHUYEN = danh_sach
    return mod.chuyen_thu_thuat_sang_tieuphau()


def run_dien_bs_tieuphau(pm_file: Path, output_file: Path) -> dict[str, Any]:
    import dien_bs_tieuphau_t5_tu_so as mod

    mod.FILE_NGUON = pm_file
    mod.FILE_DAU_RA = output_file
    mod.FILE_BAO_CAO_BS = DEFAULT_REPORT_DIR / "bao_cao_dien_bs_tieuphau_T5.xlsx"
    return mod.dien_bs_tieuphau()


def make_output_name(source_file: Path, suffix: str) -> Path:
    safe_suffix = suffix.strip()
    if not safe_suffix:
        safe_suffix = "da_sua"
    return source_file.with_name(f"{source_file.stem}_{safe_suffix}{source_file.suffix}")


def show_exception(err: Exception) -> None:
    st.error(str(err))
    with st.expander("Chi tiết lỗi"):
        st.code(traceback.format_exc())


st.set_page_config(page_title="Nhập liệu T5 - Khoa CTCH", layout="wide")
st.title("Giao diện nhập liệu T5")
st.caption("Hỗ trợ sửa dữ liệu, thêm dòng, chuyển thủ thuật sang tiểu phẫu, bổ sung phụ mổ và điền BS bằng bí danh.")

with st.sidebar:
    st.header("File làm việc")
    source_mode = st.radio("Nguồn file", ["Chọn đường dẫn có sẵn", "Tải file Excel lên"], horizontal=False)
    upload_dir = BASE_DIR / "_file_tai_len"

    if source_mode == "Tải file Excel lên":
        uploaded = st.file_uploader("Tải file PM khoa CTCH", type=["xlsx", "xlsm"])
        if uploaded is not None:
            current_file = save_upload(uploaded, upload_dir)
            st.success(f"Đã lưu file: {current_file}")
        else:
            current_file = DEFAULT_PM_FILE
    else:
        current_file = path_input("Đường dẫn file PM", DEFAULT_PM_FILE, "pm_file_path")

    output_suffix = st.text_input("Hậu tố file xuất khi lưu thủ công", value="da_nhap_lieu")
    manual_output_file = make_output_name(current_file, output_suffix)

    st.divider()
    st.header("Tùy chọn tự động")
    xoa_sau_khi_chuyen = st.checkbox("Xóa dòng khỏi sheet thu thuat sau khi chuyển", value=True)
    chi_dien_bs_trong = st.checkbox("Chỉ điền BS nếu ô đang trống", value=True)

    default_cls_text = "\n".join([
        "Hút ổ viêm/áp xe phần mềm",
        "Tiêm khớp gối",
        "Khâu vết thương phần mềm dài dưới 10 cm [tổn thương nông]",
        "Khâu vết thương phần mềm nông dài < 5cm",
        "Tiêm gân gấp ngón tay",
        "Chích rạch nhọt, Apxe nhỏ dẫn lưu",
        "Tiêm gân duỗi ngón tay",
        "Tiêm gân lồi cầu cánh tay trong/ngoài",
        "Mắt cá 1-2 thương tổn",
        "Khâu vết thương phần mềm sâu dài > 5 cm",
        "Tiêm gân Dequervein",
        "Khâu vết thương phần mềm nông dài > 5 cm",
        "Hút dịch khớp gối",
    ])
    danh_sach_cls_text = st.text_area("Danh sách Tên CLS cần chuyển sang tieuphau", value=default_cls_text, height=230)

if not current_file.exists():
    st.warning(f"Chưa tìm thấy file: {current_file}")
    st.stop()

try:
    sheet_names = get_workbook_sheets(current_file)
except Exception as e:
    show_exception(e)
    st.stop()

st.info(f"File đang mở: `{current_file}`")

main_tabs = st.tabs([
    "1. Nhập liệu / sửa nhanh",
    "2. Thêm dòng nhanh",
    "3. Chức năng tự động",
    "4. Kiểm tra thiếu dữ liệu",
    "5. Tải file",
])

with main_tabs[0]:
    st.subheader("Nhập liệu / sửa nhanh")
    available_presets = [s for s in SHEET_PRESETS if any(sheet_key(x) == sheet_key(s) for x in sheet_names)]
    if not available_presets:
        st.warning("Không tìm thấy các sheet chuẩn: phauthuat, thu thuat, tieuphau.")
    else:
        selected_sheet = st.selectbox(
            "Chọn sheet",
            available_presets,
            format_func=lambda x: f"{SHEET_PRESETS[x]['label']} ({x})",
            key="edit_sheet_select",
        )
        col_a, col_b, col_c = st.columns([1.2, 1.2, 2])
        with col_a:
            max_rows_show = st.number_input("Số dòng hiển thị", min_value=20, max_value=1000, value=200, step=20)
        with col_b:
            allow_add_rows = st.checkbox("Cho phép thêm dòng mới", value=True)
        with col_c:
            search_text = st.text_input("Tìm nhanh theo họ tên / tên CLS / phương pháp", value="")

        try:
            df, meta = workbook_to_dataframe(current_file, selected_sheet)
            df_show = df.copy()
            if search_text.strip():
                needle = normalize_text(search_text)
                mask = df_show.apply(lambda row: needle in normalize_text(" ".join(str(x) for x in row.values)), axis=1)
                df_show = df_show[mask]
                st.caption("Đang lọc để xem. Khi lưu, giao diện chỉ lưu các dòng đang hiển thị/sửa. Muốn sửa toàn bộ, hãy xóa ô tìm kiếm trước khi lưu.")
            df_show = df_show.head(int(max_rows_show))

            disabled_cols = ["__dong_excel"]
            edited = st.data_editor(
                df_show,
                use_container_width=True,
                hide_index=True,
                num_rows="dynamic" if allow_add_rows else "fixed",
                disabled=disabled_cols,
                height=620,
                key=f"editor_{selected_sheet}_{str(current_file)}",
            )

            st.caption(
                f"Header dòng {meta['header_row']} | dữ liệu dòng {meta['start_row']}–{meta['last_row']} | "
                f"Tổng Cộng dòng {meta['total_row']}"
            )

            save_col, reload_col = st.columns([1, 5])
            with save_col:
                if st.button("Lưu ra file mới", type="primary", key="save_manual_edit"):
                    result = dataframe_to_workbook(
                        source_file=current_file,
                        sheet_wanted=selected_sheet,
                        edited_df=edited,
                        output_file=manual_output_file,
                        allow_add_rows=allow_add_rows,
                    )
                    st.success(
                        f"Đã lưu: {result['file_dau_ra']} | Cập nhật {result['so_dong_cap_nhat']} dòng | "
                        f"Chèn mới {result['so_dong_chen_moi']} dòng."
                    )
                    file_download_button(Path(result["file_dau_ra"]), "Tải file vừa lưu")
            with reload_col:
                st.caption("Giao diện không ghi đè file gốc. Sau khi lưu, có thể lấy file mới làm file làm việc ở sidebar.")
        except Exception as e:
            show_exception(e)

with main_tabs[1]:
    st.subheader("Thêm dòng nhanh")
    selected_sheet_add = st.selectbox(
        "Chọn sheet cần thêm dòng",
        [s for s in SHEET_PRESETS if any(sheet_key(x) == sheet_key(s) for x in sheet_names)],
        format_func=lambda x: f"{SHEET_PRESETS[x]['label']} ({x})",
        key="add_sheet_select",
    )
    preset = SHEET_PRESETS[selected_sheet_add]
    st.caption("Dòng mới sẽ được insert ngay trước dòng Tổng Cộng để phần chữ ký/ô merge bên dưới tự đẩy xuống.")

    form_values: dict[str, Any] = {}
    with st.form("quick_add_form"):
        fields = preset["quick_fields"]
        cols = st.columns(3)
        for i, field in enumerate(fields):
            with cols[i % 3]:
                if normalize_header(field) in {normalize_header("NGÀY"), normalize_header("Ngày")}:
                    form_values[field] = st.date_input(field, value=date.today())
                elif field in {"Số Lượng", "Thành tiền", "Tuổi"}:
                    form_values[field] = st.text_input(field, value="")
                else:
                    form_values[field] = st.text_input(field, value="")
        submitted = st.form_submit_button("Thêm dòng và lưu ra file mới", type="primary")

    if submitted:
        try:
            df, meta = workbook_to_dataframe(current_file, selected_sheet_add)
            new_row = {"__dong_excel": ""}
            for col in df.columns:
                if col != "__dong_excel":
                    new_row[col] = None
            for field, value in form_values.items():
                # Ghi theo tên thật hoặc alias trong dataframe.
                target_col = None
                if field in df.columns:
                    target_col = field
                else:
                    for alias in COLUMN_ALIASES.get(field, []):
                        if alias in df.columns:
                            target_col = alias
                            break
                if target_col:
                    if isinstance(value, date):
                        value = datetime(value.year, value.month, value.day)
                    new_row[target_col] = value
            df_new = pd.concat([df, pd.DataFrame([new_row])], ignore_index=True)
            result = dataframe_to_workbook(
                source_file=current_file,
                sheet_wanted=selected_sheet_add,
                edited_df=df_new,
                output_file=manual_output_file,
                allow_add_rows=True,
            )
            st.success(
                f"Đã thêm {result['so_dong_chen_moi']} dòng vào sheet {result['sheet']} và lưu file mới."
            )
            file_download_button(Path(result["file_dau_ra"]), "Tải file vừa thêm dòng")
        except Exception as e:
            show_exception(e)

with main_tabs[2]:
    st.subheader("Chức năng tự động")
    st.write("Các nút này gọi lại những script riêng đã bổ sung trước đó, nhưng chạy qua giao diện.")

    auto_cols = st.columns(3)
    with auto_cols[0]:
        st.markdown("**Bổ sung phụ mổ từ Sổ Phẫu Thuật**")
        out_phu_mo = make_output_name(current_file, "da_bo_sung_phu_mo")
        st.caption(f"File xuất: `{out_phu_mo.name}`")
        if st.button("Chạy bổ sung phụ mổ", type="primary"):
            try:
                result = run_bo_sung_phu_mo(current_file, out_phu_mo)
                st.success("Đã chạy xong bổ sung phụ mổ.")
                st.json(result)
                file_download_button(Path(result["file_dau_ra"]), "Tải file phụ mổ")
                file_download_button(Path(result["file_bao_cao"]), "Tải báo cáo phụ mổ")
            except Exception as e:
                show_exception(e)

    with auto_cols[1]:
        st.markdown("**Chuyển thủ thuật sang tiểu phẫu + điền BS**")
        out_chuyen = make_output_name(current_file, "da_chuyen_tieu_phau")
        st.caption(f"File xuất: `{out_chuyen.name}`")
        if st.button("Chạy chuyển tiểu phẫu", type="primary"):
            try:
                result = run_chuyen_tieu_phau(
                    current_file,
                    out_chuyen,
                    xoa_khoi_thu_thuat=xoa_sau_khi_chuyen,
                    chi_dien_bs_trong=chi_dien_bs_trong,
                    danh_sach_cls_text=danh_sach_cls_text,
                )
                st.success("Đã chạy xong chuyển thủ thuật sang tiểu phẫu và điền BS.")
                st.json(result)
                file_download_button(Path(result["file_dau_ra"]), "Tải file chuyển tiểu phẫu")
                file_download_button(Path(result["file_bao_cao"]), "Tải báo cáo chuyển")
            except Exception as e:
                show_exception(e)

    with auto_cols[2]:
        st.markdown("**Chỉ điền BS cho sheet tieuphau**")
        out_bs = make_output_name(current_file, "da_dien_BS")
        st.caption(f"File xuất: `{out_bs.name}`")
        if st.button("Chạy điền BS", type="primary"):
            try:
                result = run_dien_bs_tieuphau(current_file, out_bs)
                st.success("Đã chạy xong điền BS.")
                st.json(result)
                file_download_button(Path(result["file_dau_ra"]), "Tải file đã điền BS")
                file_download_button(Path(result["file_bao_cao"]), "Tải báo cáo BS")
            except Exception as e:
                show_exception(e)

with main_tabs[3]:
    st.subheader("Kiểm tra thiếu dữ liệu")
    check_cols = st.columns(3)
    for i, sheet in enumerate(["phauthuat", "thu thuat", "tieuphau"]):
        with check_cols[i]:
            st.markdown(f"**{SHEET_PRESETS[sheet]['label']}**")
            if any(sheet_key(x) == sheet_key(sheet) for x in sheet_names):
                try:
                    missing = summarize_missing(current_file, sheet)
                    if missing.empty:
                        st.success("Không thấy thiếu dữ liệu ở cột kiểm tra chính.")
                    else:
                        st.warning(f"Có {len(missing)} dòng cần kiểm tra.")
                        st.dataframe(missing.head(200), use_container_width=True, hide_index=True)
                except Exception as e:
                    st.error(str(e))
            else:
                st.info("Không có sheet này trong file.")

with main_tabs[4]:
    st.subheader("Tải file / báo cáo")
    candidate_files = [
        current_file,
        manual_output_file,
        make_output_name(current_file, "da_bo_sung_phu_mo"),
        make_output_name(current_file, "da_chuyen_tieu_phau"),
        make_output_name(current_file, "da_dien_BS"),
        DEFAULT_REPORT_DIR / "bao_cao_bo_sung_phu_mo_T5.xlsx",
        DEFAULT_REPORT_DIR / "bao_cao_chuyen_thu_thuat_sang_tieuphau_T5.xlsx",
        DEFAULT_REPORT_DIR / "bao_cao_dien_bs_tieuphau_T5.xlsx",
    ]
    seen: set[Path] = set()
    for file in candidate_files:
        file = Path(file)
        if file in seen:
            continue
        seen.add(file)
        if file.exists():
            st.write(f"`{file}`")
            file_download_button(file, f"Tải {file.name}")
