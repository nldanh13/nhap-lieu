# -*- coding: utf-8 -*-
"""
Điền cột BS/Bác Sĩ trong sheet tieuphau của file T5 bằng bí danh nhân sự,
dựa vào Sổ Phẫu Thuật và Sổ Thủ Thuật.

Chạy độc lập sau khi đã chuyển thủ thuật sang tiểu phẫu:
    python dien_bs_tieuphau_t5_tu_so.py

Ghi chú:
- Chỉ điền bí danh có trong danh_sach_nhan_su.py.
- Nếu người trong sổ chưa có bí danh, cột BS để trống và ghi rõ trong báo cáo.
- Mặc định chỉ điền các ô BS đang trống, không ghi đè ô đã có dữ liệu.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from chuyen_thu_thuat_sang_tieuphau_t5 import (
    BASE_DIR,
    COT_CAN_CHUYEN,
    SHEET_TIEU_PHAU,
    dien_bs_tieuphau_tu_so,
    find_header_row_and_columns,
    find_sheet_case_insensitive,
)

FILE_NGUON = BASE_DIR / "2026" / "PM khoa CTCH T5_da_chuyen_tieu_phau.xlsx"
FILE_DAU_RA = BASE_DIR / "2026" / "PM khoa CTCH T5_da_chuyen_tieu_phau_da_dien_BS.xlsx"
FILE_BAO_CAO_BS = BASE_DIR / "bao_cao_dien_bs_tieuphau_T5.xlsx"


def tao_bao_cao_dien_bs(report_rows: list[dict[str, Any]], output_file: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Bao_cao_dien_BS"

    headers = [
        "Dòng Excel tieuphau",
        "Ngày T5",
        "Họ tên T5",
        "Tên CLS T5",
        "BS đã điền",
        "Trạng thái",
        "Điểm khớp",
        "Nguồn",
        "Dòng nguồn",
        "Thời gian trong sổ",
        "Tên CLS trong sổ",
        "Nhân sự trong sổ",
        "Khoa chỉ định trong sổ",
    ]
    ws.append(headers)

    for item in report_rows:
        ws.append([item.get(h, "") for h in headers])

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    widths = {
        "A": 18,
        "B": 20,
        "C": 28,
        "D": 55,
        "E": 14,
        "F": 42,
        "G": 12,
        "H": 18,
        "I": 12,
        "J": 24,
        "K": 55,
        "L": 36,
        "M": 42,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    ws.freeze_panes = "A2"
    wb.save(output_file)


def dien_bs_tieuphau() -> dict[str, Any]:
    if not FILE_NGUON.exists():
        raise FileNotFoundError(
            f"Không tìm thấy file nguồn: {FILE_NGUON}. "
            "Hãy chạy chuyen_thu_thuat_sang_tieuphau_t5.py trước."
        )

    wb = load_workbook(FILE_NGUON)
    sheet_name = find_sheet_case_insensitive(wb, SHEET_TIEU_PHAU)
    ws_tieu = wb[sheet_name]

    header_row, header_map = find_header_row_and_columns(ws_tieu, ["STT", *COT_CAN_CHUYEN])
    report_rows, filled, skipped_has_bs, not_filled = dien_bs_tieuphau_tu_so(ws_tieu, header_row, header_map)

    wb.save(FILE_DAU_RA)
    tao_bao_cao_dien_bs(report_rows, FILE_BAO_CAO_BS)

    return {
        "file_nguon": str(FILE_NGUON),
        "file_dau_ra": str(FILE_DAU_RA),
        "file_bao_cao": str(FILE_BAO_CAO_BS),
        "so_dong_dien_bs": filled,
        "so_dong_bo_qua_da_co_bs": skipped_has_bs,
        "so_dong_chua_dien_bs": not_filled,
    }


if __name__ == "__main__":
    result = dien_bs_tieuphau()
    print("ĐÃ ĐIỀN BS CHO SHEET TIEUPHAU")
    print(f"- File nguồn: {result['file_nguon']}")
    print(f"- Số dòng đã điền BS theo bí danh: {result['so_dong_dien_bs']}")
    print(f"- Số dòng bỏ qua vì đã có BS: {result['so_dong_bo_qua_da_co_bs']}")
    print(f"- Số dòng chưa điền được BS: {result['so_dong_chua_dien_bs']}")
    print(f"- File đầu ra: {result['file_dau_ra']}")
    print(f"- File báo cáo: {result['file_bao_cao']}")
