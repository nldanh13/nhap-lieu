# -*- coding: utf-8 -*-
"""
Chuyển một số kỹ thuật đang nằm ở sheet "thu thuat" sang sheet "tieuphau"
trong file PM khoa CTCH T5.

Chạy độc lập, không ảnh hưởng các file thống kê khác:
    python chuyen_thu_thuat_sang_tieuphau_t5.py

Luồng xử lý:
1. Đọc file T5.
2. Quét sheet "thu thuat".
3. Nếu cột "Tên CLS" thuộc danh sách cần chuyển, dòng đó sẽ được đưa sang sheet "tieuphau".
4. Sheet "tieuphau" được insert thêm dòng ngay sau dòng dữ liệu cuối cùng, để các dòng tổng cộng
   và các ô merge phía dưới bị đẩy xuống đúng vị trí.
5. Mặc định xóa các dòng đã chuyển khỏi sheet "thu thuat" để không bị tính trùng.
6. Cập nhật lại STT và công thức tổng cộng ở cả hai sheet.
7. Xuất ra file Excel mới, giữ nguyên file gốc.
"""

from __future__ import annotations

import json
import re
import unicodedata
from collections import defaultdict
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange

try:
    from cau_hinh_file import BASE_DIR
except Exception:
    BASE_DIR = Path(__file__).resolve().parent

try:
    from cau_hinh_khoa import get_khoa_chi_dinh_can_lay
except Exception:
    get_khoa_chi_dinh_can_lay = None

try:
    from bo_sung_phu_mo_t5_tu_so_phau_thuat import staff_to_alias
except Exception:
    staff_to_alias = None

# =========================
# CẤU HÌNH FILE
# =========================

FILE_PM_T5_GOC = BASE_DIR / "2026" / "PM khoa CTCH T5.xlsx"
FILE_PM_T5_DA_BO_SUNG_PHU_MO = BASE_DIR / "2026" / "PM khoa CTCH T5_da_bo_sung_phu_mo.xlsx"

# Nếu muốn chỉ định file nguồn thủ công, điền Path(...) tại đây.
# Nếu để None: chương trình ưu tiên file đã bổ sung phụ mổ nếu có,
# sau đó mới dùng file T5 gốc.
FILE_NGUON = None

SHEET_THU_THUAT = "thu thuat"
SHEET_TIEU_PHAU = "tieuphau"
SHEET_BAO_CAO = "Bao_cao_chuyen_tieu_phau"

FILE_DAU_RA = BASE_DIR / "2026" / "PM khoa CTCH T5_da_chuyen_tieu_phau.xlsx"
FILE_BAO_CAO = BASE_DIR / "bao_cao_chuyen_thu_thuat_sang_tieuphau_T5.xlsx"

THU_MUC_THEO_SO = BASE_DIR / "TheoSo"
FILE_SO_PHAU_THUAT = None  # Nếu muốn chỉ định thủ công, điền Path(...) tại đây.
FILE_SO_THU_THUAT = None   # Nếu muốn chỉ định thủ công, điền Path(...) tại đây.

# Sau khi chuyển sang sheet tieuphau, tự điền cột Bác Sĩ/BS bằng bí danh
# dựa vào Sổ Phẫu Thuật và Sổ Thủ Thuật.
TU_DONG_DIEN_BS_TIEU_PHAU_TU_SO = True
CHI_DIEN_BS_KHI_DANG_TRONG = True

# Ngưỡng khớp phương pháp. Nếu cùng ngày + cùng họ tên nhưng tên CLS khác quá nhiều,
# chương trình sẽ không tự điền BS để tránh điền nhầm.
DIEM_KHOP_TEN_CLS_TOI_THIEU = 0.20

# Nếu True: sau khi chèn sang tieuphau, xóa dòng đó khỏi thu thuat.
# Nếu False: chỉ copy sang tieuphau, không xóa dòng cũ.
XOA_DONG_DA_CHUYEN_KHOI_THU_THUAT = True

# Nếu True: tạo công thức tiền như các dòng mẫu đang có trong sheet tieuphau.
# Các cột đưa sang từ thu thuat vẫn chỉ là: Ngày, Họ và tên, Tuổi, Tên CLS, Số Lượng, Thành tiền.
SAO_CHEP_CONG_THUC_TIEN_TIEU_PHAU = True

# Nút cập nhật CLS bổ sung sẽ tắt bước điền Bác Sĩ từ sổ cũ,
# để các dòng mới được lấy trực tiếp từ EMR.
DIEN_BS_TU_SO_SAU_CHUYEN = True

# Giá trị mặc định ở cột Điều Dưỡng nếu dòng mẫu có sẵn "ĐD".
# Để None nếu muốn bỏ trống hoàn toàn.
GIA_TRI_DIEU_DUONG_MAC_DINH = "ĐD"

CLS_CONFIG_FILE = Path(__file__).resolve().with_name("ten_cls_tieuphau.json")

DANH_SACH_TEN_CLS_MAC_DINH = [
    "Hút ổ viêm/áp xe phần mềm",
    "Tiêm khớp gối",
    "Khâu vết thương phần mềm dài dưới 10 cm [tổn thương nông]",
    "Khâu vết thương phần mềm nông dài < 5cm",
    "Tiêm gân gấp ngón tay",
    "Chích rạch nhọt, Apxe nhỏ dẫn lưu",
    "Tiêm gân duỗi ngón tay",
    "Tiêm gân lồi cầu cánh tay trong/ngoài",
    "Mắt cá 1-2 thương tổn",
    "Khâu vết thương phần mềm sâu dài > 5 cm",
    "Tiêm gân Dequervein",
    "Khâu vết thương phần mềm nông dài > 5 cm",
    "Hút dịch khớp gối",
    "Điều trị thoái hóa khớp bằng huyết tương giàu tiểu cầu",
    "Khâu vết thương phần mềm sâu dài < 5cm",
    "Tiêm khớp vai",
    "Mắt cá 3-4 thương tổn",
    "Khâu vết thương phần mềm dài dưới 10 cm [tổn thương sâu]",
]


def load_ten_cls_can_chuyen() -> list[str]:
    """Đọc danh mục CLS đang bật từ tab Quản lý tên CLS.

    Nếu file cấu hình chưa tồn tại hoặc bị lỗi, chương trình vẫn dùng danh sách
    mặc định để không làm gián đoạn tác vụ chuyển tiểu phẫu.
    """
    if not CLS_CONFIG_FILE.exists():
        return list(DANH_SACH_TEN_CLS_MAC_DINH)
    try:
        payload = json.loads(CLS_CONFIG_FILE.read_text(encoding="utf-8"))
        rows = payload.get("cls", []) if isinstance(payload, dict) else payload
        result: list[str] = []
        seen: set[str] = set()
        for item in rows if isinstance(rows, list) else []:
            if isinstance(item, str):
                name = item.strip()
                active = True
            elif isinstance(item, dict):
                name = str(item.get("tenCls") or item.get("ten_cls") or item.get("name") or "").strip()
                active = bool(item.get("active", True))
            else:
                continue
            key_text = unicodedata.normalize("NFD", name.lower())
            key_text = "".join(ch for ch in key_text if unicodedata.category(ch) != "Mn")
            key = re.sub(r"[^a-z0-9]+", " ", key_text.replace("đ", "d")).strip()
            if name and active and key and key not in seen:
                seen.add(key)
                result.append(name)
        return result or list(DANH_SACH_TEN_CLS_MAC_DINH)
    except Exception:
        return list(DANH_SACH_TEN_CLS_MAC_DINH)


DANH_SACH_TEN_CLS_CAN_CHUYEN = load_ten_cls_can_chuyen()

# Các cột dữ liệu chuyển từ thu thuat sang tieuphau.
COT_CAN_CHUYEN = ["Ngày", "Họ và tên", "Tuổi", "Tên CLS", "Số Lượng", "Thành tiền"]

# =========================
# TIỆN ÍCH
# =========================


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    if not text or text in {"nan", "none"}:
        return ""

    text = unicodedata.normalize("NFD", text)
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    text = text.replace("đ", "d").replace("Đ", "D")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def normalize_header(value: Any) -> str:
    return normalize_text(value).replace(" ", "")


def value_is_blank(value: Any) -> bool:
    if value is None:
        return True
    text = str(value).strip()
    return text == "" or text.lower() in {"nan", "none"}


def find_sheet_case_insensitive(wb, wanted_name: str) -> str:
    wanted_key = normalize_header(wanted_name)
    for name in wb.sheetnames:
        if normalize_header(name) == wanted_key:
            return name
    raise ValueError(f"Không tìm thấy sheet '{wanted_name}'. Các sheet hiện có: {wb.sheetnames}")


def find_header_row_and_columns(ws, required_columns: list[str]) -> tuple[int, dict[str, int]]:
    required_keys = {normalize_header(x) for x in required_columns}

    best_row = None
    best_mapping: dict[str, int] = {}
    best_matched = -1

    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 40)):
        mapping: dict[str, int] = {}
        for cell in row:
            if value_is_blank(cell.value):
                continue
            key = normalize_header(cell.value)
            if key and key not in mapping:
                mapping[key] = cell.column

        if required_keys.issubset(set(mapping.keys())):
            return row[0].row, mapping

        matched = len(required_keys & set(mapping.keys()))
        if matched > best_matched:
            best_matched = matched
            best_row = row[0].row
            best_mapping = mapping

    # Dòng khớp nhiều cột nhất trong 40 dòng đầu, dùng để báo cho người dùng
    # biết chính xác đang thiếu cột nào thay vì chỉ báo chung chung.
    missing_keys = required_keys - set(best_mapping.keys())
    missing_labels = [c for c in required_columns if normalize_header(c) in missing_keys]
    detail = ""
    if best_row is not None:
        found_headers = [ws.cell(best_row, col).value for col in sorted(best_mapping.values())]
        detail = (
            f" Dòng gần khớp nhất là dòng {best_row} với các tiêu đề đọc được: {found_headers}."
            f" Còn thiếu cột: {missing_labels}."
        )

    raise ValueError(
        f"Không tìm thấy dòng tiêu đề có đủ các cột {required_columns} trong sheet '{ws.title}'.{detail}"
    )


def require_col(header_map: dict[str, int], col_name: str) -> int:
    key = normalize_header(col_name)
    if key not in header_map:
        raise ValueError(f"Không tìm thấy cột '{col_name}'.")
    return header_map[key]


def find_total_row(ws) -> int:
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for cell in row:
            if normalize_text(cell.value) == "tong cong":
                return cell.row
    raise ValueError(f"Không tìm thấy dòng Tổng Cộng trong sheet '{ws.title}'.")


def row_has_any_data(ws, row_idx: int, cols: list[int]) -> bool:
    return any(not value_is_blank(ws.cell(row_idx, col_idx).value) for col_idx in cols)


def find_last_data_row(ws, header_row: int, total_row: int, data_cols: list[int]) -> int:
    for row_idx in range(total_row - 1, header_row, -1):
        if row_has_any_data(ws, row_idx, data_cols):
            return row_idx
    return header_row


def copy_cell_style(src_cell, dst_cell) -> None:
    if src_cell.has_style:
        dst_cell.font = copy(src_cell.font)
        dst_cell.fill = copy(src_cell.fill)
        dst_cell.border = copy(src_cell.border)
        dst_cell.alignment = copy(src_cell.alignment)
        dst_cell.number_format = src_cell.number_format
        dst_cell.protection = copy(src_cell.protection)


def copy_row_style(ws, src_row: int, dst_row: int, max_col: int) -> None:
    ws.row_dimensions[dst_row].height = ws.row_dimensions[src_row].height
    for col_idx in range(1, max_col + 1):
        copy_cell_style(ws.cell(src_row, col_idx), ws.cell(dst_row, col_idx))


def translate_formula(formula: Any, src_cell, dst_cell) -> Any:
    if not isinstance(formula, str) or not formula.startswith("="):
        return formula
    try:
        return Translator(formula, origin=src_cell.coordinate).translate_formula(dst_cell.coordinate)
    except Exception:
        return formula


# =========================
# GIỮ MERGE KHI INSERT / DELETE DÒNG
# =========================


def clone_merged_ranges(ws) -> list[CellRange]:
    return [CellRange(str(rng)) for rng in ws.merged_cells.ranges]


def unmerge_all(ws) -> None:
    for rng in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(rng))


def remerge_ranges(ws, ranges: list[CellRange]) -> None:
    for rng in ranges:
        try:
            ws.merge_cells(str(rng))
        except Exception:
            pass


def shift_merged_ranges_for_insert(ranges: list[CellRange], insert_at: int, amount: int) -> list[CellRange]:
    shifted = []
    for rng in ranges:
        new_rng = CellRange(str(rng))
        if new_rng.min_row >= insert_at:
            new_rng.shift(row_shift=amount, col_shift=0)
        elif new_rng.max_row >= insert_at:
            # Nếu insert nằm giữa vùng merge, mở rộng vùng merge xuống dưới.
            new_rng.max_row += amount
        shifted.append(new_rng)
    return shifted


def insert_rows_preserve_merges(ws, insert_at: int, amount: int) -> None:
    if amount <= 0:
        return
    ranges = clone_merged_ranges(ws)
    unmerge_all(ws)
    ws.insert_rows(insert_at, amount)
    remerge_ranges(ws, shift_merged_ranges_for_insert(ranges, insert_at, amount))


def shift_merged_ranges_for_delete(ranges: list[CellRange], deleted_rows: list[int]) -> list[CellRange]:
    deleted_rows = sorted(set(deleted_rows))
    shifted = []

    for rng in ranges:
        new_rng = CellRange(str(rng))

        # Dữ liệu cần xóa nằm ở vùng bảng, còn merge thường nằm ở phần tiêu đề/footer.
        # Nếu merge nằm dưới các dòng đã xóa thì dịch lên tương ứng.
        rows_before_min = sum(1 for row_idx in deleted_rows if row_idx < new_rng.min_row)
        rows_inside = sum(1 for row_idx in deleted_rows if new_rng.min_row <= row_idx <= new_rng.max_row)

        if rows_inside >= (new_rng.max_row - new_rng.min_row + 1):
            continue

        new_rng.min_row -= rows_before_min
        new_rng.max_row -= rows_before_min

        if rows_inside:
            new_rng.max_row -= rows_inside
            if new_rng.max_row < new_rng.min_row:
                continue

        shifted.append(new_rng)

    return shifted


def delete_rows_preserve_merges(ws, rows_to_delete: list[int]) -> None:
    rows_to_delete = sorted(set(rows_to_delete), reverse=True)
    if not rows_to_delete:
        return

    original_ranges = clone_merged_ranges(ws)
    unmerge_all(ws)

    for row_idx in rows_to_delete:
        ws.delete_rows(row_idx, 1)

    remerge_ranges(ws, shift_merged_ranges_for_delete(original_ranges, rows_to_delete))


# =========================
# XỬ LÝ DỮ LIỆU
# =========================


@dataclass
class DongCanChuyen:
    source_row: int
    stt: Any
    ngay: Any
    ho_ten: Any
    tuoi: Any
    ten_cls: Any
    so_luong: Any
    thanh_tien: Any


TARGET_CLS_NORMALIZED = {normalize_text(x) for x in DANH_SACH_TEN_CLS_CAN_CHUYEN}


def resolve_input_file() -> Path:
    if FILE_NGUON:
        path = Path(FILE_NGUON)
        if not path.exists():
            raise FileNotFoundError(f"Không tìm thấy FILE_NGUON: {path}")
        return path

    for path in [FILE_PM_T5_DA_BO_SUNG_PHU_MO, FILE_PM_T5_GOC]:
        if path.exists():
            return path

    raise FileNotFoundError(
        f"Không tìm thấy file T5. Đã kiểm tra: {FILE_PM_T5_DA_BO_SUNG_PHU_MO} và {FILE_PM_T5_GOC}"
    )


def lay_dong_can_chuyen(ws_thu, header_map_thu: dict[str, int], header_row_thu: int, total_row_thu: int) -> list[DongCanChuyen]:
    col_stt = require_col(header_map_thu, "STT")
    col_ngay = require_col(header_map_thu, "Ngày")
    col_ho_ten = require_col(header_map_thu, "Họ và tên")
    col_tuoi = require_col(header_map_thu, "Tuổi")
    col_ten_cls = require_col(header_map_thu, "Tên CLS")
    col_so_luong = require_col(header_map_thu, "Số Lượng")
    col_thanh_tien = require_col(header_map_thu, "Thành tiền")

    result: list[DongCanChuyen] = []
    for row_idx in range(header_row_thu + 1, total_row_thu):
        ten_cls = ws_thu.cell(row_idx, col_ten_cls).value
        if normalize_text(ten_cls) not in TARGET_CLS_NORMALIZED:
            continue

        result.append(
            DongCanChuyen(
                source_row=row_idx,
                stt=ws_thu.cell(row_idx, col_stt).value,
                ngay=ws_thu.cell(row_idx, col_ngay).value,
                ho_ten=ws_thu.cell(row_idx, col_ho_ten).value,
                tuoi=ws_thu.cell(row_idx, col_tuoi).value,
                ten_cls=ten_cls,
                so_luong=ws_thu.cell(row_idx, col_so_luong).value,
                thanh_tien=ws_thu.cell(row_idx, col_thanh_tien).value,
            )
        )

    return result


def find_formula_template_row(ws_tieu, header_row_tieu: int, total_row_tieu: int, formula_cols: list[int]) -> int:
    for row_idx in range(header_row_tieu + 1, total_row_tieu):
        for col_idx in formula_cols:
            value = ws_tieu.cell(row_idx, col_idx).value
            if isinstance(value, str) and value.startswith("="):
                return row_idx
    return header_row_tieu + 1


def chen_sang_tieuphau(ws_tieu, header_map_tieu: dict[str, int], rows_to_move: list[DongCanChuyen]) -> tuple[int, int, int]:
    if not rows_to_move:
        return 0, 0, 0

    header_row = find_header_row_and_columns(ws_tieu, ["STT", *COT_CAN_CHUYEN])[0]
    total_row_before = find_total_row(ws_tieu)

    col_stt = require_col(header_map_tieu, "STT")
    col_ngay = require_col(header_map_tieu, "Ngày")
    col_ho_ten = require_col(header_map_tieu, "Họ và tên")
    col_tuoi = require_col(header_map_tieu, "Tuổi")
    col_ten_cls = require_col(header_map_tieu, "Tên CLS")
    col_so_luong = require_col(header_map_tieu, "Số Lượng")
    col_thanh_tien = require_col(header_map_tieu, "Thành tiền")

    # Các cột tiền/phân công trong sheet tieuphau.
    col_bac_si = header_map_tieu.get(normalize_header("Bác Sĩ"))
    col_so_tien_bs = header_map_tieu.get(normalize_header("Số Tiền"))  # cột I nếu header trùng có thể lấy cột đầu tiên
    col_dieu_duong = header_map_tieu.get(normalize_header("Điều Dưỡng"))

    # Header có 2 cột "Số Tiền". Mapping lấy cột đầu tiên, nên xác định thêm bằng vị trí cố định nếu có.
    # Với mẫu hiện tại: I = số tiền bác sĩ, K = số tiền điều dưỡng, L = thực lãnh.
    col_i = 9 if ws_tieu.max_column >= 9 else None
    col_k = 11 if ws_tieu.max_column >= 11 else None
    col_l = 12 if ws_tieu.max_column >= 12 else None
    formula_cols = [c for c in [col_i, col_k, col_l] if c]

    data_cols = [col_ngay, col_ho_ten, col_tuoi, col_ten_cls, col_so_luong, col_thanh_tien]
    last_data_row_before = find_last_data_row(ws_tieu, header_row, total_row_before, data_cols)
    insert_at = last_data_row_before + 1
    amount = len(rows_to_move)

    # Tìm STT hiện tại lớn nhất trong bảng tieuphau.
    current_stt_values = []
    for row_idx in range(header_row + 1, total_row_before):
        value = ws_tieu.cell(row_idx, col_stt).value
        if isinstance(value, (int, float)):
            current_stt_values.append(int(value))
    next_stt = (max(current_stt_values) if current_stt_values else 0) + 1

    style_template_row = header_row + 1
    formula_template_row = find_formula_template_row(ws_tieu, header_row, total_row_before, formula_cols)

    insert_rows_preserve_merges(ws_tieu, insert_at, amount)

    for offset, item in enumerate(rows_to_move):
        row_idx = insert_at + offset
        copy_row_style(ws_tieu, style_template_row, row_idx, ws_tieu.max_column)

        ws_tieu.cell(row_idx, col_stt).value = next_stt + offset
        ws_tieu.cell(row_idx, col_ngay).value = item.ngay
        ws_tieu.cell(row_idx, col_ho_ten).value = item.ho_ten
        ws_tieu.cell(row_idx, col_tuoi).value = item.tuoi
        ws_tieu.cell(row_idx, col_ten_cls).value = item.ten_cls
        ws_tieu.cell(row_idx, col_so_luong).value = item.so_luong
        ws_tieu.cell(row_idx, col_thanh_tien).value = item.thanh_tien

        # Bác sĩ sẽ được điền ở bước tra Sổ Phẫu Thuật/Sổ Thủ Thuật phía sau.
        if col_bac_si:
            ws_tieu.cell(row_idx, col_bac_si).value = None

        if SAO_CHEP_CONG_THUC_TIEN_TIEU_PHAU:
            if col_dieu_duong:
                ws_tieu.cell(row_idx, col_dieu_duong).value = GIA_TRI_DIEU_DUONG_MAC_DINH

            for col_idx in formula_cols:
                src = ws_tieu.cell(formula_template_row, col_idx)
                dst = ws_tieu.cell(row_idx, col_idx)
                if isinstance(src.value, str) and src.value.startswith("="):
                    dst.value = translate_formula(src.value, src, dst)

    total_row_after = find_total_row(ws_tieu)
    refresh_formula_references_tieuphau(ws_tieu, header_row, total_row_after)
    update_total_formulas_tieuphau(ws_tieu, header_row, total_row_after)

    return insert_at, insert_at + amount - 1, total_row_after


def xoa_khoi_thu_thuat_va_cap_nhat(ws_thu, header_map_thu: dict[str, int], header_row_thu: int, rows_to_move: list[DongCanChuyen]) -> int:
    if not rows_to_move or not XOA_DONG_DA_CHUYEN_KHOI_THU_THUAT:
        total_row = find_total_row(ws_thu)
        update_total_formulas_thu_thuat(ws_thu, header_row_thu, total_row)
        return total_row

    delete_rows_preserve_merges(ws_thu, [item.source_row for item in rows_to_move])

    # Sau khi xóa, tìm lại header/total vì vị trí đã thay đổi.
    header_row, header_map = find_header_row_and_columns(ws_thu, ["STT", *COT_CAN_CHUYEN])
    total_row = find_total_row(ws_thu)

    col_stt = require_col(header_map, "STT")
    data_cols = [require_col(header_map, name) for name in COT_CAN_CHUYEN]

    stt = 1
    for row_idx in range(header_row + 1, total_row):
        if row_has_any_data(ws_thu, row_idx, data_cols):
            ws_thu.cell(row_idx, col_stt).value = stt
            stt += 1
        else:
            ws_thu.cell(row_idx, col_stt).value = None

    update_total_formulas_thu_thuat(ws_thu, header_row, total_row)
    return total_row


def refresh_formula_references_tieuphau(ws, header_row: int, total_row: int) -> None:
    """Sửa lại công thức theo đúng số dòng sau khi insert.

    openpyxl không tự dịch công thức khi insert dòng. Trong mẫu hiện tại có một
    dòng trống kèm công thức nằm ngay trước Tổng Cộng; sau khi insert, công thức
    đó có thể vẫn tham chiếu dòng cũ. Hàm này rà lại các ô công thức I/K/L và
    dịch theo đúng dòng hiện tại để tránh cộng nhầm.
    """
    formula_cols = [9, 11, 12]
    template_row = find_formula_template_row(ws, header_row, total_row, formula_cols)

    for row_idx in range(header_row + 1, total_row):
        for col_idx in formula_cols:
            cell = ws.cell(row_idx, col_idx)
            if isinstance(cell.value, str) and cell.value.startswith("="):
                src = ws.cell(template_row, col_idx)
                cell.value = translate_formula(src.value, src, cell)


def update_total_formulas_thu_thuat(ws, header_row: int, total_row: int) -> None:
    start = header_row + 1
    end = total_row - 1
    for col_letter in ["F", "G", "H"]:
        ws[f"{col_letter}{total_row}"] = f"=SUM({col_letter}{start}:{col_letter}{end})"


def update_total_formulas_tieuphau(ws, header_row: int, total_row: int) -> None:
    start = header_row + 1
    end = total_row - 1
    for col_letter in ["F", "G", "I", "K", "L"]:
        ws[f"{col_letter}{total_row}"] = f"=SUM({col_letter}{start}:{col_letter}{end})"


def update_formula_lien_ket_tong_cong(ws_tieu, total_row_tieu: int, total_row_thu: int) -> None:
    # Công thức mẫu hiện tại nằm ở G42: =+G39+'thu thuat'!G1258
    # Sau khi insert/delete dòng, tự tìm công thức chứa 'thu thuat' rồi cập nhật lại.
    for row in ws_tieu.iter_rows(min_row=total_row_tieu + 1, max_row=ws_tieu.max_row):
        for cell in row:
            if isinstance(cell.value, str) and "thu thuat" in cell.value.lower():
                cell.value = f"=+G{total_row_tieu}+'thu thuat'!G{total_row_thu}"
                return



# =========================
# ĐIỀN CỘT BÁC SĨ TỪ SỔ PHẪU THUẬT / SỔ THỦ THUẬT
# =========================


def excel_serial_to_datetime(value: int | float) -> datetime | None:
    try:
        return datetime(1899, 12, 30) + timedelta(days=float(value))
    except Exception:
        return None


def parse_datetime_any(value: Any) -> datetime | None:
    if value is None:
        return None

    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    if isinstance(value, (int, float)):
        return excel_serial_to_datetime(value)

    text = str(value).strip()
    if not text or text.lower() in {"nan", "none"}:
        return None

    # Sửa lỗi nhập nhầm năm kiểu 07/5/5026 -> 07/5/2026.
    text = re.sub(r"/(5\d{3})(\s|$)", r"/2026\2", text)

    formats = [
        "%Y/%m/%d %H:%M",
        "%Y-%m-%d %H:%M",
        "%H:%M %d/%m/%Y",
        "%d/%m/%Y %H:%M",
        "%d/%m/%Y",
        "%d/%m/%y",
        "%Y/%m/%d",
        "%Y-%m-%d",
    ]
    for fmt in formats:
        try:
            return datetime.strptime(text, fmt)
        except Exception:
            pass

    # Dạng Sổ: 11:40 07/05/2026
    m = re.search(r"(\d{1,2}:\d{2})\s+(\d{1,2}/\d{1,2}/\d{4})", text)
    if m:
        try:
            return datetime.strptime(f"{m.group(1)} {m.group(2)}", "%H:%M %d/%m/%Y")
        except Exception:
            pass

    # Dạng file tháng: 2026/05/04 09:03
    m = re.search(r"(\d{4})/(\d{1,2})/(\d{1,2})(?:\s+(\d{1,2}:\d{2}))?", text)
    if m:
        y, month, day, hm = m.groups()
        hm = hm or "00:00"
        try:
            return datetime.strptime(f"{y}/{month}/{day} {hm}", "%Y/%m/%d %H:%M")
        except Exception:
            pass

    # Bắt lại ngày/tháng/năm trong chuỗi nếu có thêm ký tự khác.
    m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", text)
    if m:
        day, month, y = map(int, m.groups())
        if y >= 3000:
            y = 2026
        try:
            return datetime(y, month, day)
        except Exception:
            return None

    return None


def simplify_method(value: Any) -> str:
    text = normalize_text(value)
    if not text:
        return ""

    text = re.sub(r"\b\d+\s*%\b", " ", text)
    replacements = {
        "ptns": "phau thuat noi soi",
        "pttt": "phau thuat thu thuat",
        "pt": "phau thuat",
        "khx": "ket hop xuong",
        "dc khx": "dung cu ket hop xuong",
        "vpm": "vet phan mem",
        "apxe": "ap xe",
        "dequervein": "de quervain",
        "de quervain": "de quervain",
    }
    for old, new in sorted(replacements.items(), key=lambda x: -len(x[0])):
        text = re.sub(rf"\b{re.escape(old)}\b", new, text)

    text = re.sub(r"\s+", " ", text).strip()
    return text


def token_overlap_score(a: str, b: str) -> float:
    tokens_a = {x for x in a.split() if len(x) >= 2}
    tokens_b = {x for x in b.split() if len(x) >= 2}
    if not tokens_a or not tokens_b:
        return 0.0
    return len(tokens_a & tokens_b) / max(len(tokens_a), len(tokens_b))


def method_score(target_method: Any, source_method: Any) -> float:
    a = simplify_method(target_method)
    b = simplify_method(source_method)
    if not a or not b:
        return 0.0

    if a == b:
        return 1.0
    if a in b or b in a:
        return 0.90

    ratio = SequenceMatcher(None, a, b).ratio()
    overlap = token_overlap_score(a, b)
    return max(ratio, overlap)


def tim_file_theo_mau(patterns: list[str], file_chi_dinh: Path | None, ten_loai: str) -> Path:
    if file_chi_dinh:
        path = Path(file_chi_dinh)
        if not path.exists():
            raise FileNotFoundError(f"Không tìm thấy {ten_loai}: {path}")
        return path

    candidates = []
    if THU_MUC_THEO_SO.exists():
        for pattern in patterns:
            candidates.extend(THU_MUC_THEO_SO.glob(pattern))

    candidates = [p for p in candidates if not p.name.startswith("~$")]
    candidates = sorted(set(candidates), key=lambda p: p.stat().st_mtime, reverse=True)

    if not candidates:
        raise FileNotFoundError(
            f"Không tìm thấy {ten_loai} trong thư mục {THU_MUC_THEO_SO}. "
            "Hãy điền đường dẫn file thủ công trong phần cấu hình."
        )

    return candidates[0]


def tim_file_so_phau_thuat() -> Path:
    return tim_file_theo_mau(
        ["*SoPhauThuat*.xlsx", "*SoPhauThuat*.xls", "*PhauThuat*.xlsx", "*PhauThuat*.xls"],
        FILE_SO_PHAU_THUAT,
        "Sổ Phẫu Thuật",
    )


def tim_file_so_thu_thuat() -> Path:
    return tim_file_theo_mau(
        ["*SoThuThuat*.xlsx", "*SoThuThuat*.xls", "*ThuThuat*.xlsx", "*ThuThuat*.xls"],
        FILE_SO_THU_THUAT,
        "Sổ Thủ Thuật",
    )


def read_book_rows(file_path: Path) -> list[dict[str, Any]]:
    """Đọc sổ HIS có header 2 dòng."""
    wb = load_workbook(file_path, data_only=True, read_only=True)
    ws = wb.active
    max_col = min(ws.max_column or 60, 80)

    header_row_1 = None
    for row_idx, row_values in enumerate(
        ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20), max_col=max_col, values_only=True),
        start=1,
    ):
        normalized_values = {normalize_header(v) for v in row_values if v is not None}
        if "tt" in normalized_values and "hovaten" in normalized_values:
            header_row_1 = row_idx
            break

    if header_row_1 is None:
        header_row_1 = 6

    header_row_2 = header_row_1 + 1
    header_rows = list(
        ws.iter_rows(min_row=header_row_1, max_row=header_row_2, max_col=max_col, values_only=True)
    )
    header_1_values = header_rows[0]
    header_2_values = header_rows[1]

    headers = []
    used = defaultdict(int)

    for col_idx in range(max_col):
        h1 = header_1_values[col_idx] if col_idx < len(header_1_values) else None
        h2 = header_2_values[col_idx] if col_idx < len(header_2_values) else None
        parts = [str(x).strip() for x in (h1, h2) if x is not None and str(x).strip()]

        if not parts:
            name = f"Cot_{col_idx + 1}"
        elif parts[0].lower().startswith("nhân viên") or parts[0].lower().startswith("nhan vien"):
            name = parts[-1]
        elif len(parts) == 1:
            name = parts[0]
        else:
            name = " ".join(parts)

        used[name] += 1
        if used[name] > 1:
            name = f"{name}_{used[name]}"
        headers.append(name)

    rows = []
    for values in ws.iter_rows(min_row=header_row_2 + 1, max_row=ws.max_row, max_col=max_col, values_only=True):
        record = dict(zip(headers, values))

        try:
            float(str(record.get("TT", "")).strip())
        except Exception:
            continue

        rows.append(record)

    wb.close()
    return rows


def khoa_duoc_lay(khoa_value: Any, danh_sach_khoa: list[str]) -> bool:
    if not danh_sach_khoa:
        return True
    khoa_norm = normalize_text(khoa_value)
    return any(normalize_text(k) in khoa_norm for k in danh_sach_khoa)


def get_khoa_filter(ten_so: str) -> list[str]:
    if get_khoa_chi_dinh_can_lay is not None:
        try:
            return list(get_khoa_chi_dinh_can_lay(ten_so) or [])
        except Exception:
            pass

    if "Phẫu" in ten_so or "Phau" in ten_so:
        return ["Khoa Cấp Cứu", "Khoa Ngoại Chấn Thương"]
    return ["Khoa Khám Bệnh", "Khoa Ngoại Chấn Thương"]


@dataclass
class SourceBSRecord:
    source_name: str
    source_index: int
    ngay: date
    ho_ten_key: str
    ho_ten: str
    ten_cls: str
    thoi_gian: datetime | None
    khoa: str
    raw_staff: str
    bs_alias: str


def build_source_bs_records() -> list[SourceBSRecord]:
    if staff_to_alias is None:
        raise RuntimeError(
            "Không import được hàm staff_to_alias từ bo_sung_phu_mo_t5_tu_so_phau_thuat.py."
        )

    configs = [
        {
            "source_name": "Sổ Phẫu Thuật",
            "file": tim_file_so_phau_thuat(),
            "method_col": "Phương pháp phẫu thuật",
            "staff_col": "Phẫu thuật viên",
        },
        {
            "source_name": "Sổ Thủ Thuật",
            "file": tim_file_so_thu_thuat(),
            "method_col": "Phương pháp thủ thuật",
            "staff_col": "TT chính",
        },
    ]

    records: list[SourceBSRecord] = []
    for cfg in configs:
        khoa_filter = get_khoa_filter(cfg["source_name"])
        for i, row in enumerate(read_book_rows(cfg["file"]), start=1):
            dt = parse_datetime_any(row.get("Thời gian thực hiện"))
            if not dt:
                continue

            ho_ten = row.get("Họ và tên")
            ho_ten_key = normalize_text(ho_ten)
            if not ho_ten_key:
                continue

            khoa = row.get("Khoa phòng chỉ định")
            if not khoa_duoc_lay(khoa, khoa_filter):
                continue

            raw_staff = str(row.get(cfg["staff_col"]) or "").strip()
            bs_alias = staff_to_alias(raw_staff)

            records.append(
                SourceBSRecord(
                    source_name=cfg["source_name"],
                    source_index=i,
                    ngay=dt.date(),
                    ho_ten_key=ho_ten_key,
                    ho_ten=str(ho_ten).strip() if ho_ten is not None else "",
                    ten_cls=str(row.get(cfg["method_col"]) or "").strip(),
                    thoi_gian=dt,
                    khoa=str(khoa or "").strip(),
                    raw_staff=raw_staff,
                    bs_alias=bs_alias,
                )
            )

    return records


def choose_best_bs_match(
    target_date: date | None,
    target_name: Any,
    target_cls: Any,
    index_by_date_name: dict[tuple[date, str], list[SourceBSRecord]],
) -> tuple[SourceBSRecord | None, float, str]:
    if not target_date:
        return None, 0.0, "Không đọc được ngày"

    name_key = normalize_text(target_name)
    if not name_key:
        return None, 0.0, "Không có họ tên"

    candidates = index_by_date_name.get((target_date, name_key), [])
    if not candidates:
        return None, 0.0, "Không tìm thấy theo ngày + họ tên trong sổ"

    scored = [(method_score(target_cls, candidate.ten_cls), candidate) for candidate in candidates]
    scored.sort(key=lambda x: x[0], reverse=True)

    best_score, best_candidate = scored[0]
    if best_score < DIEM_KHOP_TEN_CLS_TOI_THIEU:
        return best_candidate, best_score, "Có ngày + họ tên nhưng Tên CLS khác nhiều, chưa tự điền"

    if len(scored) >= 2 and abs(scored[0][0] - scored[1][0]) < 0.001:
        return best_candidate, best_score, "Có nhiều dòng cùng ngày + họ tên gần giống nhau, đã lấy dòng đầu tiên"

    return best_candidate, best_score, "Khớp"


def find_bs_column(header_map_tieu: dict[str, int]) -> int:
    for col_name in ["Bác Sĩ", "Bác sĩ", "BS"]:
        key = normalize_header(col_name)
        if key in header_map_tieu:
            return header_map_tieu[key]
    raise ValueError("Không tìm thấy cột Bác Sĩ/BS trong sheet tieuphau.")


def dien_bs_tieuphau_tu_so(ws_tieu, header_row_tieu: int, header_map_tieu: dict[str, int]) -> tuple[list[dict[str, Any]], int, int, int]:
    if not TU_DONG_DIEN_BS_TIEU_PHAU_TU_SO:
        return [], 0, 0, 0

    total_row = find_total_row(ws_tieu)
    col_ngay = require_col(header_map_tieu, "Ngày")
    col_ho_ten = require_col(header_map_tieu, "Họ và tên")
    col_ten_cls = require_col(header_map_tieu, "Tên CLS")
    col_bs = find_bs_column(header_map_tieu)

    source_records = build_source_bs_records()
    index_by_date_name: dict[tuple[date, str], list[SourceBSRecord]] = defaultdict(list)
    for record in source_records:
        index_by_date_name[(record.ngay, record.ho_ten_key)].append(record)

    report_rows: list[dict[str, Any]] = []
    filled = 0
    skipped_has_bs = 0
    not_filled = 0

    for row_idx in range(header_row_tieu + 1, total_row):
        current_bs = ws_tieu.cell(row_idx, col_bs).value
        if CHI_DIEN_BS_KHI_DANG_TRONG and not value_is_blank(current_bs):
            skipped_has_bs += 1
            continue

        target_name = ws_tieu.cell(row_idx, col_ho_ten).value
        target_cls = ws_tieu.cell(row_idx, col_ten_cls).value

        # Bỏ qua các dòng trống/template nằm trước Tổng Cộng.
        if value_is_blank(ws_tieu.cell(row_idx, col_ngay).value) and value_is_blank(target_name) and value_is_blank(target_cls):
            continue

        target_dt = parse_datetime_any(ws_tieu.cell(row_idx, col_ngay).value)
        target_date = target_dt.date() if target_dt else None

        match, score, status = choose_best_bs_match(target_date, target_name, target_cls, index_by_date_name)

        if match and score >= DIEM_KHOP_TEN_CLS_TOI_THIEU and match.bs_alias:
            ws_tieu.cell(row_idx, col_bs).value = match.bs_alias
            filled += 1
            bs_value = match.bs_alias
        elif match and score >= DIEM_KHOP_TEN_CLS_TOI_THIEU and not match.bs_alias:
            not_filled += 1
            bs_value = ""
            status = "Khớp trong sổ nhưng nhân sự chưa có bí danh"
        else:
            not_filled += 1
            bs_value = ""

        report_rows.append(
            {
                "Dòng Excel tieuphau": row_idx,
                "Ngày T5": ws_tieu.cell(row_idx, col_ngay).value,
                "Họ tên T5": target_name,
                "Tên CLS T5": target_cls,
                "BS đã điền": bs_value,
                "Trạng thái": status,
                "Điểm khớp": round(float(score or 0), 3),
                "Nguồn": match.source_name if match else "",
                "Dòng nguồn": match.source_index if match else "",
                "Thời gian trong sổ": match.thoi_gian if match else "",
                "Tên CLS trong sổ": match.ten_cls if match else "",
                "Nhân sự trong sổ": match.raw_staff if match else "",
                "Khoa chỉ định trong sổ": match.khoa if match else "",
            }
        )

    return report_rows, filled, skipped_has_bs, not_filled


def tao_bao_cao(
    rows_to_move: list[DongCanChuyen],
    input_file: Path,
    output_file: Path,
    bs_report_rows: list[dict[str, Any]] | None = None,
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_BAO_CAO

    headers = [
        "STT",
        "Dòng gốc sheet thu thuat",
        "Ngày",
        "Họ và tên",
        "Tuổi",
        "Tên CLS",
        "Số Lượng",
        "Thành tiền",
        "Trạng thái",
        "File nguồn",
        "File đầu ra",
    ]
    ws.append(headers)

    for idx, item in enumerate(rows_to_move, start=1):
        ws.append(
            [
                idx,
                item.source_row,
                item.ngay,
                item.ho_ten,
                item.tuoi,
                item.ten_cls,
                item.so_luong,
                item.thanh_tien,
                "Đã chuyển sang tieuphau" if XOA_DONG_DA_CHUYEN_KHOI_THU_THUAT else "Đã copy sang tieuphau",
                str(input_file),
                str(output_file),
            ]
        )

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    widths = {
        "A": 8,
        "B": 22,
        "C": 20,
        "D": 28,
        "E": 16,
        "F": 55,
        "G": 12,
        "H": 16,
        "I": 26,
        "J": 65,
        "K": 65,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    ws.freeze_panes = "A2"

    if bs_report_rows is not None:
        ws_bs = wb.create_sheet("Bao_cao_dien_BS")
        headers_bs = [
            "Dòng Excel tieuphau",
            "Ngày T5",
            "Họ tên T5",
            "Tên CLS T5",
            "BS đã điền",
            "Trạng thái",
            "Điểm khớp",
            "Nguồn",
            "Dòng nguồn",
            "Thời gian trong sổ",
            "Tên CLS trong sổ",
            "Nhân sự trong sổ",
            "Khoa chỉ định trong sổ",
        ]
        ws_bs.append(headers_bs)
        for item in bs_report_rows:
            ws_bs.append([item.get(h, "") for h in headers_bs])

        for cell in ws_bs[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        widths_bs = {
            "A": 18,
            "B": 20,
            "C": 28,
            "D": 55,
            "E": 14,
            "F": 42,
            "G": 12,
            "H": 18,
            "I": 12,
            "J": 24,
            "K": 55,
            "L": 36,
            "M": 42,
        }
        for col, width in widths_bs.items():
            ws_bs.column_dimensions[col].width = width
        ws_bs.freeze_panes = "A2"

    wb.save(FILE_BAO_CAO)


def chuyen_thu_thuat_sang_tieuphau() -> dict[str, Any]:
    input_file = resolve_input_file()
    wb = load_workbook(input_file)

    sheet_thu_name = find_sheet_case_insensitive(wb, SHEET_THU_THUAT)
    sheet_tieu_name = find_sheet_case_insensitive(wb, SHEET_TIEU_PHAU)
    ws_thu = wb[sheet_thu_name]
    ws_tieu = wb[sheet_tieu_name]

    header_row_thu, header_map_thu = find_header_row_and_columns(ws_thu, ["STT", *COT_CAN_CHUYEN])
    header_row_tieu, header_map_tieu = find_header_row_and_columns(ws_tieu, ["STT", *COT_CAN_CHUYEN])

    total_row_thu_before = find_total_row(ws_thu)
    total_row_tieu_before = find_total_row(ws_tieu)

    rows_to_move = lay_dong_can_chuyen(ws_thu, header_map_thu, header_row_thu, total_row_thu_before)

    if rows_to_move:
        insert_start, insert_end, total_row_tieu_after = chen_sang_tieuphau(ws_tieu, header_map_tieu, rows_to_move)
        total_row_thu_after = xoa_khoi_thu_thuat_va_cap_nhat(ws_thu, header_map_thu, header_row_thu, rows_to_move)
        update_formula_lien_ket_tong_cong(ws_tieu, total_row_tieu_after, total_row_thu_after)
    else:
        insert_start = insert_end = 0
        total_row_tieu_after = total_row_tieu_before
        total_row_thu_after = total_row_thu_before
        refresh_formula_references_tieuphau(ws_tieu, header_row_tieu, total_row_tieu_after)
        update_total_formulas_tieuphau(ws_tieu, header_row_tieu, total_row_tieu_after)
        update_total_formulas_thu_thuat(ws_thu, header_row_thu, total_row_thu_after)

    # Sau khi đã insert/xóa dòng, có thể điền Bác Sĩ từ hai sổ cũ.
    # Chế độ cập nhật CLS bổ sung sẽ để trống Bác Sĩ để giao diện lấy trực tiếp từ EMR.
    header_row_tieu_after, header_map_tieu_after = find_header_row_and_columns(ws_tieu, ["STT", *COT_CAN_CHUYEN])
    if DIEN_BS_TU_SO_SAU_CHUYEN:
        bs_report_rows, so_dong_dien_bs, so_dong_bo_qua_da_co_bs, so_dong_chua_dien_bs = dien_bs_tieuphau_tu_so(
            ws_tieu,
            header_row_tieu_after,
            header_map_tieu_after,
        )
    else:
        bs_report_rows = []
        so_dong_dien_bs = 0
        so_dong_bo_qua_da_co_bs = 0
        so_dong_chua_dien_bs = len(rows_to_move)

    wb.save(FILE_DAU_RA)
    tao_bao_cao(rows_to_move, input_file, FILE_DAU_RA, bs_report_rows)

    return {
        "file_nguon": str(input_file),
        "file_dau_ra": str(FILE_DAU_RA),
        "file_bao_cao": str(FILE_BAO_CAO),
        "so_dong_chuyen": len(rows_to_move),
        "xoa_khoi_thu_thuat": XOA_DONG_DA_CHUYEN_KHOI_THU_THUAT,
        "dong_chen_tieuphau_tu": insert_start,
        "dong_chen_tieuphau_den": insert_end,
        "tong_cong_thu_thuat_truoc": total_row_thu_before,
        "tong_cong_thu_thuat_sau": total_row_thu_after,
        "tong_cong_tieuphau_truoc": total_row_tieu_before,
        "tong_cong_tieuphau_sau": total_row_tieu_after,
        "so_dong_dien_bs": so_dong_dien_bs,
        "so_dong_bo_qua_da_co_bs": so_dong_bo_qua_da_co_bs,
        "so_dong_chua_dien_bs": so_dong_chua_dien_bs,
        "dien_bs_tu_so_sau_chuyen": DIEN_BS_TU_SO_SAU_CHUYEN,
    }


if __name__ == "__main__":
    result = chuyen_thu_thuat_sang_tieuphau()
    print("ĐÃ CHUYỂN THỦ THUẬT SANG TIỂU PHẪU")
    print(f"- File nguồn: {result['file_nguon']}")
    print(f"- Số dòng đã chuyển: {result['so_dong_chuyen']}")
    print(f"- Xóa khỏi sheet thu thuat: {result['xoa_khoi_thu_thuat']}")
    print(
        f"- Dòng chèn ở sheet tieuphau: {result['dong_chen_tieuphau_tu']}"
        f" đến {result['dong_chen_tieuphau_den']}"
    )
    print(
        f"- Dòng Tổng Cộng thu thuat: {result['tong_cong_thu_thuat_truoc']}"
        f" -> {result['tong_cong_thu_thuat_sau']}"
    )
    print(
        f"- Dòng Tổng Cộng tieuphau: {result['tong_cong_tieuphau_truoc']}"
        f" -> {result['tong_cong_tieuphau_sau']}"
    )
    print(f"- Số dòng đã điền BS theo bí danh: {result['so_dong_dien_bs']}")
    print(f"- Số dòng bỏ qua vì đã có BS: {result['so_dong_bo_qua_da_co_bs']}")
    print(f"- Số dòng chưa điền được BS: {result['so_dong_chua_dien_bs']}")
    print(f"- File đầu ra: {result['file_dau_ra']}")
    print(f"- File báo cáo: {result['file_bao_cao']}")
