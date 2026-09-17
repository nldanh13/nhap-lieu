import re
import unicodedata

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from cau_hinh_file import (
    NAM_DU_LIEU,
    THANG_DA_CHON_TONG_HOP,
    XUAT_SHEET_TRUNG_LAP_PHAUTHUAT_BHYT,
)
from cau_hinh_phuong_phap import NHOM_PHAN_LOAI_THEO_THU_TU


SHEET_THONG_KE = "Theo nhóm tháng"
SHEET_CHUA_PHAN_LOAI = "Chưa phân loại"
SHEET_TRUNG_LAP = "Trùng PT-BHYT"


def _remove_accents(text):
    text = str(text)
    text = unicodedata.normalize("NFD", text)
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    text = text.replace("đ", "d").replace("Đ", "D")
    return text


def _extract_month_from_text(value):
    """Lấy số tháng từ tên file/nhãn tháng như TH1.2025, T02, Tháng 6..."""
    if value is None or pd.isna(value):
        return None

    name = _remove_accents(str(value)).lower()

    patterns = [
        r"(?:^|[^a-z0-9])th\s*0?([1-9]|1[0-2])(?:[^0-9]|$)",
        r"(?:^|[^a-z0-9])t\s*0?([1-9]|1[0-2])(?:[^0-9]|$)",
        r"thang\s*0?([1-9]|1[0-2])(?:[^0-9]|$)",
    ]

    for pattern in patterns:
        match = re.search(pattern, name)
        if match:
            return int(match.group(1))

    return None


def _months_to_export(df_filtered):
    """
    Xác định các tháng cần hiển thị:
    - Nếu cấu hình chọn tháng cụ thể: hiện đúng các tháng đã chọn.
    - Nếu chạy cả năm: hiện đủ tháng 1 đến tháng 12.
    - Nếu không xác định được cấu hình: lấy các tháng có trong dữ liệu.
    """
    if THANG_DA_CHON_TONG_HOP is None:
        return list(range(1, 13))

    if isinstance(THANG_DA_CHON_TONG_HOP, (list, tuple, set)) and THANG_DA_CHON_TONG_HOP:
        return sorted(int(x) for x in THANG_DA_CHON_TONG_HOP)

    month_values = []
    if not df_filtered.empty:
        for col in ["Tháng/File", "File nguồn", "Nguồn"]:
            if col in df_filtered.columns:
                month_values.extend(df_filtered[col].map(_extract_month_from_text).dropna().astype(int).tolist())

    return sorted(set(month_values)) or list(range(1, 13))


def _add_month_column(df_filtered):
    df = df_filtered.copy()

    month = pd.Series([None] * len(df), index=df.index, dtype="object")
    for col in ["Tháng/File", "File nguồn", "Nguồn"]:
        if col not in df.columns:
            continue
        parsed = df[col].map(_extract_month_from_text)
        month = month.where(month.notna(), parsed)

    df["_thang_so"] = month
    return df


def make_group_month_table(df_filtered):
    """Tạo bảng duy nhất: Theo nhóm phân loại | các tháng | Tổng số."""
    months = _months_to_export(df_filtered)
    month_columns = [f"Tháng {m}" for m in months]

    if df_filtered.empty or "Nhóm phân loại" not in df_filtered.columns:
        columns = ["Theo nhóm phân loại"] + month_columns + ["Tổng số"]
        return pd.DataFrame(columns=columns)

    df = _add_month_column(df_filtered)
    df = df[df["_thang_so"].isin(months)].copy()

    pivot = (
        df.groupby(["Nhóm phân loại", "_thang_so"], dropna=False)
        .size()
        .unstack(fill_value=0)
    )

    for month in months:
        if month not in pivot.columns:
            pivot[month] = 0

    pivot = pivot[months]

    # Sắp xếp theo thứ tự nhóm đã cấu hình, các nhóm phát sinh thêm để cuối.
    configured_groups = [g for g in NHOM_PHAN_LOAI_THEO_THU_TU if g in pivot.index or True]
    extra_groups = [g for g in pivot.index.tolist() if g not in configured_groups]
    if "Chưa phân loại" in pivot.index and "Chưa phân loại" not in configured_groups:
        extra_groups = [g for g in extra_groups if g != "Chưa phân loại"] + ["Chưa phân loại"]

    ordered_groups = configured_groups + extra_groups
    pivot = pivot.reindex(ordered_groups, fill_value=0)

    # Nếu không muốn hiện nhóm hoàn toàn bằng 0, có thể bỏ comment dòng bên dưới.
    # pivot = pivot[pivot.sum(axis=1) > 0]

    result = pivot.reset_index().rename(columns={"Nhóm phân loại": "Theo nhóm phân loại"})
    result.columns = ["Theo nhóm phân loại"] + month_columns
    result["Tổng số"] = result[month_columns].sum(axis=1)

    total_row = {"Theo nhóm phân loại": "Tổng cộng"}
    for col in month_columns:
        total_row[col] = int(result[col].sum())
    total_row["Tổng số"] = int(result["Tổng số"].sum())
    result = pd.concat([result, pd.DataFrame([total_row])], ignore_index=True)

    return result


def _format_single_sheet(writer, sheet_name):
    workbook = writer.book
    ws = workbook[sheet_name]

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    total_fill = PatternFill("solid", fgColor="FFF2CC")
    header_font = Font(bold=True)
    total_font = Font(bold=True)

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Căn giữa và định dạng số cho các cột tháng/tổng.
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=2, max_col=ws.max_column):
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.number_format = "#,##0"

    # Dòng tổng cộng.
    if ws.max_row >= 2:
        for cell in ws[ws.max_row]:
            cell.font = total_font
            cell.fill = total_fill

    ws.column_dimensions["A"].width = 45
    for col_cells in ws.iter_cols(min_col=2, max_col=ws.max_column):
        col_letter = col_cells[0].column_letter
        ws.column_dimensions[col_letter].width = 13


def _format_detail_sheet(writer, sheet_name):
    ws = writer.book[sheet_name]
    header_fill = PatternFill("solid", fgColor="FCE4D6")
    header_font = Font(bold=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for col_cells in ws.iter_cols(min_row=1, max_row=ws.max_row):
        header = str(col_cells[0].value or "")
        max_len = max((len(str(cell.value or "")) for cell in col_cells[:200]), default=len(header))
        width = min(max(max_len + 2, 12), 55)
        if "phuong phap" in _remove_accents(header).lower():
            width = min(max(width, 35), 60)
        ws.column_dimensions[col_cells[0].column_letter].width = width


def export_excel(output_file, summaries, df_filtered, df_roles):
    """
    Xuất bảng thống kê chính và sheet rà soát Chưa phân loại.
    """
    table = make_group_month_table(df_filtered)
    unclassified = summaries.get(SHEET_CHUA_PHAN_LOAI, pd.DataFrame()).copy()

    duplicate_report = pd.DataFrame()
    if hasattr(df_filtered, "attrs"):
        report = df_filtered.attrs.get("trung_lap_phauthuat_bhyt")
        if isinstance(report, pd.DataFrame):
            duplicate_report = report.copy()

    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        table.to_excel(writer, sheet_name=SHEET_THONG_KE, index=False)
        _format_single_sheet(writer, SHEET_THONG_KE)

        # Luôn tạo sheet này, kể cả khi không có dòng, để đầu ra nhất quán.
        if unclassified.empty and len(unclassified.columns) == 0:
            unclassified = pd.DataFrame(
                columns=["Phương pháp chưa phân loại", "Số lần", "Gợi ý thêm vào file"]
            )
        unclassified.to_excel(writer, sheet_name=SHEET_CHUA_PHAN_LOAI, index=False)
        _format_detail_sheet(writer, SHEET_CHUA_PHAN_LOAI)

        if XUAT_SHEET_TRUNG_LAP_PHAUTHUAT_BHYT and not duplicate_report.empty:
            duplicate_report.to_excel(writer, sheet_name=SHEET_TRUNG_LAP, index=False)
            _format_detail_sheet(writer, SHEET_TRUNG_LAP)
