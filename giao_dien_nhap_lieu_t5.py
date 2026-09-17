# -*- coding: utf-8 -*-
"""
Giao diện nhập liệu T5 dùng Tkinter, chạy offline, không cần cài Streamlit.

Chạy:
    python giao_dien_nhap_lieu_t5.py

Chức năng:
- Mở file PM khoa CTCH T5.
- Xem/tìm/sửa nhanh các sheet: phauthuat, thu thuat, tieuphau.
- Thêm dòng mới bằng cách insert trước dòng Tổng Cộng để đẩy phần chữ ký/ô merge xuống.
- Chạy các chức năng tự động: bổ sung phụ mổ, chuyển thủ thuật sang tiểu phẫu, điền BS.
- Luôn xuất ra file mới, không ghi đè file gốc.
"""

from __future__ import annotations

import os
import re
import sys
import threading
import traceback
import unicodedata
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.worksheet.cell_range import CellRange

THIS_DIR = Path(__file__).resolve().parent
if str(THIS_DIR) not in sys.path:
    sys.path.insert(0, str(THIS_DIR))

try:
    from cau_hinh_file import BASE_DIR
except Exception:
    BASE_DIR = THIS_DIR

DEFAULT_FILE = BASE_DIR / "2026" / "PM khoa CTCH T5.xlsx"
DEFAULT_REPORT_DIR = BASE_DIR

SHEET_PRESETS = {
    "phauthuat": {
        "label": "Phẫu thuật",
        "required": ["STT", "NGÀY", "Họ và tên bệnh nhân", "Chẩn đoán và phương pháp phẫu thuật"],
        "fields": [
            "NGÀY",
            "Họ và tên bệnh nhân",
            "Tuổi",
            "Chẩn đoán và phương pháp phẫu thuật",
            "PTV chính",
            "Phụ mổ 1",
            "Phụ mổ 2",
            "Phụ mổ 3",
        ],
        "missing_cols": ["PTV chính", "Phụ mổ 1", "Phụ mổ 2", "Phụ mổ 3"],
    },
    "thu thuat": {
        "label": "Thủ thuật",
        "required": ["STT", "Ngày", "Họ và tên", "Tuổi", "Tên CLS", "Số Lượng", "Thành tiền"],
        "fields": ["Ngày", "Họ và tên", "Tuổi", "Tên CLS", "Số Lượng", "Thành tiền"],
        "missing_cols": [],
    },
    "tieuphau": {
        "label": "Tiểu phẫu",
        "required": ["STT", "Ngày", "Họ và tên", "Tuổi", "Tên CLS", "Số Lượng", "Thành tiền"],
        "fields": ["Ngày", "Họ và tên", "Tuổi", "Tên CLS", "Số Lượng", "Thành tiền", "Bác Sĩ", "Điều Dưỡng"],
        "missing_cols": ["Bác Sĩ"],
    },
}

COLUMN_ALIASES = {
    "Bác Sĩ": ["Bác Sĩ", "BS", "Bác sĩ", "Bac Si", "Bac si"],
    "Điều Dưỡng": ["Điều Dưỡng", "Điều dưỡng", "DD", "ĐD", "Dieu Duong", "Dieu duong"],
    "Số Lượng": ["Số Lượng", "Số lượng", "SL", "So Luong", "So luong"],
}

DEFAULT_CLS_LIST = [
    "Hút ổ viêm/áp xe phần mềm",
    "Tiêm khớp gối",
    "Khâu vết thương phần mềm dài dưới 10 cm [tổn thương nông]",
    "Khâu vết thương phần mềm nông dài < 5cm",
    "Tiêm gân gấp ngón tay",
    "Chích rạch nhọt, Apxe nhỏ dẫn lưu",
    "Tiêm gân duỗi ngón tay",
    "Tiêm gân lồi cầu cánh tay trong/ngoài",
    "Mắt cá 1-2 thương tổn",
    "Khâu vết thương phần mềm sâu dài > 5 cm",
    "Tiêm gân Dequervein",
    "Khâu vết thương phần mềm nông dài > 5 cm",
    "Hút dịch khớp gối",
]


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    if not text or text in {"nan", "none", "nat"}:
        return ""
    text = unicodedata.normalize("NFD", text)
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    text = text.replace("đ", "d").replace("Đ", "D")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def normalize_header(value: Any) -> str:
    return normalize_text(value).replace(" ", "")


def value_is_blank(value: Any) -> bool:
    if value is None:
        return True
    text = str(value).strip()
    return text == "" or text.lower() in {"nan", "none", "nat"}


def sheet_key(name: str) -> str:
    return normalize_header(name)


def find_sheet_name(wb, wanted: str) -> str:
    wanted_key = sheet_key(wanted)
    for name in wb.sheetnames:
        if sheet_key(name) == wanted_key:
            return name
    raise ValueError(f"Không tìm thấy sheet '{wanted}'.")


def find_header_row_and_map(ws, required_columns: list[str]) -> tuple[int, dict[str, int], dict[int, str]]:
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 50)):
        by_key: dict[str, int] = {}
        by_col: dict[int, str] = {}
        for cell in row:
            if value_is_blank(cell.value):
                continue
            raw = str(cell.value).strip()
            key = normalize_header(raw)
            by_key.setdefault(key, cell.column)
            by_col[cell.column] = raw
            for canonical, aliases in COLUMN_ALIASES.items():
                alias_keys = {normalize_header(x) for x in aliases}
                if key in alias_keys:
                    by_key.setdefault(normalize_header(canonical), cell.column)

        ok = True
        for col in required_columns:
            keys = [normalize_header(x) for x in COLUMN_ALIASES.get(col, [col])]
            if not any(k in by_key for k in keys):
                ok = False
                break
        if ok:
            return row[0].row, by_key, by_col

    raise ValueError(f"Không tìm thấy dòng tiêu đề có đủ cột: {required_columns}")


def col_index(header_map: dict[str, int], col_name: str) -> int | None:
    keys = [normalize_header(x) for x in COLUMN_ALIASES.get(col_name, [col_name])]
    for key in keys:
        if key in header_map:
            return header_map[key]
    return None


def find_total_row(ws, header_row: int) -> int:
    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row):
        for cell in row:
            if normalize_text(cell.value) == "tong cong":
                return cell.row
    return ws.max_row + 1


def row_has_data(ws, row_idx: int, columns: Iterable[int]) -> bool:
    return any(not value_is_blank(ws.cell(row_idx, col).value) for col in columns)


def get_data_rows(ws, header_row: int, header_by_col: dict[int, str], total_row: int) -> list[int]:
    cols = sorted(header_by_col.keys())
    end = total_row - 1 if total_row <= ws.max_row else ws.max_row
    rows = []
    for row_idx in range(header_row + 1, end + 1):
        if row_has_data(ws, row_idx, cols):
            rows.append(row_idx)
    return rows


def choose_table_columns(header_by_col: dict[int, str]) -> list[tuple[int, str]]:
    cols: list[tuple[int, str]] = []
    used: set[str] = set()
    for col_idx in sorted(header_by_col):
        name = str(header_by_col[col_idx]).strip()
        if not name:
            continue
        base = name
        suffix = 2
        while name in used:
            name = f"{base}_{suffix}"
            suffix += 1
        used.add(name)
        cols.append((col_idx, name))
    return cols


def display_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    return str(value)


def parse_user_value(text: str) -> Any:
    text = (text or "").strip()
    if not text:
        return None
    for fmt in ["%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d"]:
        try:
            return datetime.strptime(text, fmt)
        except Exception:
            pass
    # Giữ dạng số nếu người dùng nhập số rõ ràng.
    if re.fullmatch(r"-?\d+", text):
        try:
            return int(text)
        except Exception:
            return text
    if re.fullmatch(r"-?\d+[\.,]\d+", text):
        try:
            return float(text.replace(",", "."))
        except Exception:
            return text
    return text


def copy_cell_style(src_cell, dst_cell) -> None:
    if src_cell.has_style:
        dst_cell.font = copy(src_cell.font)
        dst_cell.fill = copy(src_cell.fill)
        dst_cell.border = copy(src_cell.border)
        dst_cell.alignment = copy(src_cell.alignment)
        dst_cell.number_format = src_cell.number_format
        dst_cell.protection = copy(src_cell.protection)


def copy_row_style(ws, src_row: int, dst_row: int, max_col: int) -> None:
    ws.row_dimensions[dst_row].height = ws.row_dimensions[src_row].height
    for col_idx in range(1, max_col + 1):
        src = ws.cell(src_row, col_idx)
        dst = ws.cell(dst_row, col_idx)
        copy_cell_style(src, dst)
        if isinstance(src.value, str) and src.value.startswith("="):
            try:
                dst.value = Translator(src.value, origin=src.coordinate).translate_formula(dst.coordinate)
            except Exception:
                dst.value = src.value
        else:
            dst.value = None


def clone_merged_ranges(ws) -> list[CellRange]:
    return [CellRange(str(rng)) for rng in ws.merged_cells.ranges]


def shift_merged_ranges_for_insert(ranges: list[CellRange], insert_at: int, amount: int) -> list[CellRange]:
    shifted = []
    for rng in ranges:
        new = CellRange(str(rng))
        if new.min_row >= insert_at:
            new.shift(row_shift=amount)
        elif new.max_row >= insert_at:
            new.max_row += amount
        shifted.append(new)
    return shifted


def safe_insert_rows(ws, insert_at: int, amount: int) -> None:
    ranges = clone_merged_ranges(ws)
    for rng in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(rng))
    ws.insert_rows(insert_at, amount)
    for rng in shift_merged_ranges_for_insert(ranges, insert_at, amount):
        ws.merge_cells(str(rng))


def output_path_for(source_file: Path, suffix: str) -> Path:
    suffix = suffix.strip() or "da_nhap_lieu"
    return source_file.with_name(f"{source_file.stem}_{suffix}{source_file.suffix}")


class ExcelTable:
    def __init__(self, file_path: Path, sheet_key_name: str):
        self.file_path = Path(file_path)
        self.sheet_key_name = sheet_key_name
        self.wb = load_workbook(self.file_path, data_only=False)
        self.sheet_name = find_sheet_name(self.wb, sheet_key_name)
        self.ws = self.wb[self.sheet_name]
        self.preset = SHEET_PRESETS[sheet_key_name]
        self.header_row, self.header_map, self.header_by_col = find_header_row_and_map(
            self.ws, self.preset["required"]
        )
        self.total_row = find_total_row(self.ws, self.header_row)
        self.table_cols = choose_table_columns(self.header_by_col)
        self.col_by_name = {name: col_idx for col_idx, name in self.table_cols}
        self.data_rows = get_data_rows(self.ws, self.header_row, self.header_by_col, self.total_row)

    def rows_for_tree(self) -> list[dict[str, Any]]:
        out = []
        fields = self.preset["fields"]
        for row_idx in self.data_rows:
            item = {"__dong_excel": row_idx}
            for field in fields:
                col = self.find_col_by_field(field)
                item[field] = display_value(self.ws.cell(row_idx, col).value) if col else ""
            out.append(item)
        return out

    def find_col_by_field(self, field: str) -> int | None:
        if field in self.col_by_name:
            return self.col_by_name[field]
        col = col_index(self.header_map, field)
        if col:
            return col
        field_key = normalize_header(field)
        for name, col_idx in self.col_by_name.items():
            if normalize_header(name) == field_key:
                return col_idx
        return None

    def values_for_row(self, row_idx: int) -> dict[str, str]:
        values = {}
        for field in self.preset["fields"]:
            col = self.find_col_by_field(field)
            values[field] = display_value(self.ws.cell(row_idx, col).value) if col else ""
        return values

    def save_existing_row(self, row_idx: int, field_values: dict[str, str], output_file: Path) -> None:
        for field, text in field_values.items():
            col = self.find_col_by_field(field)
            if col:
                self.ws.cell(row_idx, col).value = parse_user_value(text)
        self.renumber_stt()
        output_file.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(output_file)

    def insert_new_row(self, field_values: dict[str, str], output_file: Path) -> int:
        insert_at = self.total_row
        safe_insert_rows(self.ws, insert_at, 1)
        style_row = max(self.header_row + 1, insert_at - 1)
        copy_row_style(self.ws, style_row, insert_at, self.ws.max_column)
        # Sau insert cần cập nhật lại mapping/total_row vẫn đủ để ghi vào dòng mới.
        for field, text in field_values.items():
            col = self.find_col_by_field(field)
            if col:
                self.ws.cell(insert_at, col).value = parse_user_value(text)
        self.renumber_stt()
        output_file.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(output_file)
        return insert_at

    def renumber_stt(self) -> None:
        stt_col = self.find_col_by_field("STT")
        if not stt_col:
            return
        data_cols = [c for c in self.col_by_name.values() if c != stt_col]
        total = find_total_row(self.ws, self.header_row)
        stt = 1
        for row_idx in range(self.header_row + 1, total):
            if row_has_data(self.ws, row_idx, data_cols):
                self.ws.cell(row_idx, stt_col).value = stt
                stt += 1

    def missing_rows(self) -> list[dict[str, Any]]:
        miss_cols = self.preset.get("missing_cols", [])
        if not miss_cols:
            return []
        out = []
        for row_idx in self.data_rows:
            missing_any = False
            for field in miss_cols:
                col = self.find_col_by_field(field)
                if col and value_is_blank(self.ws.cell(row_idx, col).value):
                    missing_any = True
            if missing_any:
                item = {"__dong_excel": row_idx}
                for field in self.preset["fields"]:
                    col = self.find_col_by_field(field)
                    item[field] = display_value(self.ws.cell(row_idx, col).value) if col else ""
                out.append(item)
        return out


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Nhập liệu T5 - Khoa CTCH")
        self.geometry("1320x780")
        self.current_file = tk.StringVar(value=str(DEFAULT_FILE))
        self.sheet_var = tk.StringVar(value="phauthuat")
        self.search_var = tk.StringVar()
        self.output_suffix_var = tk.StringVar(value="da_nhap_lieu")
        self.status_var = tk.StringVar(value="Sẵn sàng")
        self.selected_row_idx: int | None = None
        self.table: ExcelTable | None = None
        self.field_vars: dict[str, tk.StringVar] = {}
        self.tree_columns: list[str] = []
        self.rows_cache: list[dict[str, Any]] = []
        self._build_ui()
        self.load_sheet()

    def _build_ui(self) -> None:
        top = ttk.Frame(self, padding=8)
        top.pack(fill=tk.X)

        ttk.Label(top, text="File PM:").grid(row=0, column=0, sticky="w")
        ttk.Entry(top, textvariable=self.current_file, width=95).grid(row=0, column=1, sticky="we", padx=5)
        ttk.Button(top, text="Chọn file", command=self.browse_file).grid(row=0, column=2, padx=4)
        ttk.Button(top, text="Mở thư mục", command=self.open_folder).grid(row=0, column=3, padx=4)

        ttk.Label(top, text="Sheet:").grid(row=1, column=0, sticky="w", pady=5)
        sheet_combo = ttk.Combobox(top, textvariable=self.sheet_var, values=list(SHEET_PRESETS.keys()), state="readonly", width=18)
        sheet_combo.grid(row=1, column=1, sticky="w", padx=5)
        sheet_combo.bind("<<ComboboxSelected>>", lambda _e: self.load_sheet())
        ttk.Button(top, text="Tải lại sheet", command=self.load_sheet).grid(row=1, column=2, padx=4)

        ttk.Label(top, text="Hậu tố file xuất:").grid(row=1, column=3, sticky="e")
        ttk.Entry(top, textvariable=self.output_suffix_var, width=22).grid(row=1, column=4, sticky="w", padx=4)
        top.columnconfigure(1, weight=1)

        body = ttk.PanedWindow(self, orient=tk.HORIZONTAL)
        body.pack(fill=tk.BOTH, expand=True, padx=8, pady=4)

        left = ttk.Frame(body)
        right = ttk.Frame(body)
        body.add(left, weight=3)
        body.add(right, weight=2)

        search_frame = ttk.Frame(left)
        search_frame.pack(fill=tk.X, pady=(0, 4))
        ttk.Label(search_frame, text="Tìm:").pack(side=tk.LEFT)
        search_entry = ttk.Entry(search_frame, textvariable=self.search_var, width=50)
        search_entry.pack(side=tk.LEFT, padx=5)
        search_entry.bind("<KeyRelease>", lambda _e: self.refresh_tree())
        ttk.Button(search_frame, text="Xóa tìm", command=self.clear_search).pack(side=tk.LEFT)
        ttk.Button(search_frame, text="Lọc dòng thiếu dữ liệu", command=self.show_missing).pack(side=tk.LEFT, padx=5)

        tree_frame = ttk.Frame(left)
        tree_frame.pack(fill=tk.BOTH, expand=True)
        self.tree = ttk.Treeview(tree_frame, show="headings")
        yscroll = ttk.Scrollbar(tree_frame, orient=tk.VERTICAL, command=self.tree.yview)
        xscroll = ttk.Scrollbar(tree_frame, orient=tk.HORIZONTAL, command=self.tree.xview)
        self.tree.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)
        self.tree.grid(row=0, column=0, sticky="nsew")
        yscroll.grid(row=0, column=1, sticky="ns")
        xscroll.grid(row=1, column=0, sticky="ew")
        tree_frame.rowconfigure(0, weight=1)
        tree_frame.columnconfigure(0, weight=1)
        self.tree.bind("<<TreeviewSelect>>", self.on_select_row)

        self.form_title = ttk.Label(right, text="Form nhập liệu", font=("Arial", 12, "bold"))
        self.form_title.pack(anchor="w", pady=(0, 6))
        self.form_frame = ttk.Frame(right)
        self.form_frame.pack(fill=tk.X)

        button_frame = ttk.Frame(right)
        button_frame.pack(fill=tk.X, pady=8)
        ttk.Button(button_frame, text="Lưu dòng đang chọn ra file mới", command=self.save_selected_row).pack(fill=tk.X, pady=2)
        ttk.Button(button_frame, text="Thêm thành dòng mới", command=self.add_new_row).pack(fill=tk.X, pady=2)
        ttk.Button(button_frame, text="Xóa nội dung form", command=self.clear_form).pack(fill=tk.X, pady=2)

        auto = ttk.LabelFrame(right, text="Chức năng tự động", padding=8)
        auto.pack(fill=tk.X, pady=8)
        ttk.Button(auto, text="1. Bổ sung phụ mổ từ Sổ Phẫu Thuật", command=self.auto_bo_sung_phu_mo).pack(fill=tk.X, pady=2)
        ttk.Button(auto, text="2. Chuyển thủ thuật sang tiểu phẫu + điền BS", command=self.auto_chuyen_tieu_phau).pack(fill=tk.X, pady=2)
        ttk.Button(auto, text="3. Chỉ điền BS cho sheet tieuphau", command=self.auto_dien_bs).pack(fill=tk.X, pady=2)

        cls_frame = ttk.LabelFrame(right, text="Danh sách Tên CLS chuyển sang tieuphau", padding=6)
        cls_frame.pack(fill=tk.BOTH, expand=True, pady=4)
        self.cls_text = tk.Text(cls_frame, height=10, wrap="word")
        self.cls_text.pack(fill=tk.BOTH, expand=True)
        self.cls_text.insert("1.0", "\n".join(DEFAULT_CLS_LIST))

        bottom = ttk.Frame(self, padding=(8, 2, 8, 8))
        bottom.pack(fill=tk.X)
        ttk.Label(bottom, textvariable=self.status_var).pack(side=tk.LEFT)

    def current_path(self) -> Path:
        return Path(self.current_file.get().strip().strip('"'))

    def output_path(self, suffix: str | None = None) -> Path:
        return output_path_for(self.current_path(), suffix or self.output_suffix_var.get())

    def browse_file(self) -> None:
        path = filedialog.askopenfilename(
            title="Chọn file Excel PM khoa CTCH",
            filetypes=[("Excel files", "*.xlsx *.xlsm"), ("All files", "*.*")],
            initialdir=str(DEFAULT_FILE.parent if DEFAULT_FILE.parent.exists() else THIS_DIR),
        )
        if path:
            self.current_file.set(path)
            self.load_sheet()

    def open_folder(self) -> None:
        path = self.current_path()
        folder = path.parent if path.exists() else THIS_DIR
        try:
            os.startfile(folder)  # type: ignore[attr-defined]
        except Exception:
            messagebox.showinfo("Thư mục", str(folder))

    def load_sheet(self) -> None:
        try:
            path = self.current_path()
            if not path.exists():
                self.status_var.set(f"Không tìm thấy file: {path}")
                return
            self.table = ExcelTable(path, self.sheet_var.get())
            self.rows_cache = self.table.rows_for_tree()
            self.selected_row_idx = None
            self.build_tree_columns()
            self.build_form()
            self.refresh_tree()
            self.status_var.set(
                f"Đã tải {self.table.sheet_name}: {len(self.rows_cache)} dòng | "
                f"Header dòng {self.table.header_row}, Tổng Cộng dòng {self.table.total_row}"
            )
        except Exception as e:
            self.show_error(e)

    def build_tree_columns(self) -> None:
        if not self.table:
            return
        fields = self.table.preset["fields"]
        self.tree_columns = ["Dòng Excel", *fields]
        self.tree.configure(columns=self.tree_columns)
        for col in self.tree_columns:
            self.tree.heading(col, text=col)
            width = 110
            if col in {"Họ và tên", "Họ và tên bệnh nhân"}:
                width = 190
            elif col in {"Tên CLS", "Chẩn đoán và phương pháp phẫu thuật"}:
                width = 360
            elif col == "Dòng Excel":
                width = 80
            self.tree.column(col, width=width, minwidth=60, anchor="w")

    def build_form(self) -> None:
        for child in self.form_frame.winfo_children():
            child.destroy()
        self.field_vars = {}
        if not self.table:
            return
        for idx, field in enumerate(self.table.preset["fields"]):
            ttk.Label(self.form_frame, text=field).grid(row=idx, column=0, sticky="nw", padx=2, pady=2)
            var = tk.StringVar()
            self.field_vars[field] = var
            if field in {"Tên CLS", "Chẩn đoán và phương pháp phẫu thuật"}:
                text = tk.Text(self.form_frame, height=3, width=45, wrap="word")
                text.grid(row=idx, column=1, sticky="we", padx=2, pady=2)
                text.bind("<KeyRelease>", lambda _e, f=field, t=text: self.field_vars[f].set(t.get("1.0", "end").strip()))
                text._field_name = field  # type: ignore[attr-defined]
            else:
                ttk.Entry(self.form_frame, textvariable=var, width=48).grid(row=idx, column=1, sticky="we", padx=2, pady=2)
        self.form_frame.columnconfigure(1, weight=1)

    def refresh_tree(self, custom_rows: list[dict[str, Any]] | None = None) -> None:
        for item in self.tree.get_children():
            self.tree.delete(item)
        rows = custom_rows if custom_rows is not None else self.rows_cache
        needle = normalize_text(self.search_var.get())
        count = 0
        for row in rows:
            haystack = normalize_text(" ".join(str(v) for v in row.values()))
            if needle and needle not in haystack:
                continue
            values = [row.get("__dong_excel", "")]
            if self.table:
                for field in self.table.preset["fields"]:
                    values.append(row.get(field, ""))
            self.tree.insert("", tk.END, values=values, iid=str(row.get("__dong_excel", count)))
            count += 1
        self.status_var.set(f"Đang hiển thị {count} dòng")

    def clear_search(self) -> None:
        self.search_var.set("")
        self.refresh_tree()

    def set_text_widget(self, field: str, value: str) -> None:
        for child in self.form_frame.winfo_children():
            if isinstance(child, tk.Text) and getattr(child, "_field_name", None) == field:
                child.delete("1.0", "end")
                child.insert("1.0", value)

    def on_select_row(self, _event=None) -> None:
        if not self.table:
            return
        selected = self.tree.selection()
        if not selected:
            return
        try:
            row_idx = int(selected[0])
        except Exception:
            return
        self.selected_row_idx = row_idx
        values = self.table.values_for_row(row_idx)
        for field, var in self.field_vars.items():
            value = values.get(field, "")
            var.set(value)
            self.set_text_widget(field, value)
        self.status_var.set(f"Đang chọn dòng Excel {row_idx}")

    def get_form_values(self) -> dict[str, str]:
        values = {}
        for field, var in self.field_vars.items():
            # Text widget đã đồng bộ qua KeyRelease, nhưng nếu người dùng chưa nhấn phím sau paste thì đọc lại.
            for child in self.form_frame.winfo_children():
                if isinstance(child, tk.Text) and getattr(child, "_field_name", None) == field:
                    var.set(child.get("1.0", "end").strip())
            values[field] = var.get()
        return values

    def clear_form(self) -> None:
        self.selected_row_idx = None
        for field, var in self.field_vars.items():
            var.set("")
            self.set_text_widget(field, "")
        self.tree.selection_remove(self.tree.selection())

    def save_selected_row(self) -> None:
        if not self.table or not self.selected_row_idx:
            messagebox.showwarning("Chưa chọn dòng", "Hãy chọn một dòng trong bảng trước.")
            return
        try:
            out = self.output_path()
            self.table.save_existing_row(self.selected_row_idx, self.get_form_values(), out)
            self.status_var.set(f"Đã lưu dòng {self.selected_row_idx} ra file: {out}")
            if messagebox.askyesno("Đã lưu", f"Đã lưu file:\n{out}\n\nBạn có muốn chuyển sang làm việc trên file vừa lưu không?"):
                self.current_file.set(str(out))
            self.load_sheet()
        except Exception as e:
            self.show_error(e)

    def add_new_row(self) -> None:
        if not self.table:
            return
        try:
            out = self.output_path()
            new_row = self.table.insert_new_row(self.get_form_values(), out)
            self.status_var.set(f"Đã thêm dòng mới tại dòng Excel {new_row}, file: {out}")
            if messagebox.askyesno("Đã thêm dòng", f"Đã thêm dòng mới và lưu file:\n{out}\n\nBạn có muốn chuyển sang làm việc trên file vừa lưu không?"):
                self.current_file.set(str(out))
            self.load_sheet()
        except Exception as e:
            self.show_error(e)

    def show_missing(self) -> None:
        try:
            if not self.table:
                return
            rows = self.table.missing_rows()
            self.refresh_tree(custom_rows=rows)
            if rows:
                self.status_var.set(f"Có {len(rows)} dòng thiếu dữ liệu trong sheet {self.table.sheet_name}")
            else:
                self.status_var.set(f"Không thấy dòng thiếu dữ liệu trong sheet {self.table.sheet_name}")
        except Exception as e:
            self.show_error(e)

    def run_background(self, title: str, func) -> None:
        def worker():
            try:
                self.status_var.set(f"Đang chạy: {title}...")
                result = func()
                self.status_var.set(f"Hoàn tất: {title}")
                self.after(0, lambda: messagebox.showinfo("Hoàn tất", f"{title}\n\n{result}"))
                self.after(0, self.load_sheet)
            except Exception as e:
                self.after(0, lambda: self.show_error(e))
        threading.Thread(target=worker, daemon=True).start()

    def auto_bo_sung_phu_mo(self) -> None:
        def task():
            import bo_sung_phu_mo_t5_tu_so_phau_thuat as mod
            source = self.current_path()
            out = output_path_for(source, "da_bo_sung_phu_mo")
            mod.FILE_PM_T5 = source
            mod.FILE_DAU_RA = out
            mod.FILE_BAO_CAO = DEFAULT_REPORT_DIR / "bao_cao_bo_sung_phu_mo_T5.xlsx"
            mod.GHI_DE_FILE_GOC = False
            mod.bo_sung_phu_mo()
            return f"File đầu ra: {out}\nBáo cáo: {mod.FILE_BAO_CAO}"
        self.run_background("Bổ sung phụ mổ", task)

    def auto_chuyen_tieu_phau(self) -> None:
        def task():
            import chuyen_thu_thuat_sang_tieuphau_t5 as mod
            source = self.current_path()
            out = output_path_for(source, "da_chuyen_tieu_phau")
            cls_list = [line.strip() for line in self.cls_text.get("1.0", "end").splitlines() if line.strip()]
            mod.FILE_NGUON = source
            mod.FILE_DAU_RA = out
            mod.FILE_BAO_CAO = DEFAULT_REPORT_DIR / "bao_cao_chuyen_thu_thuat_sang_tieuphau_T5.xlsx"
            mod.XOA_DONG_DA_CHUYEN_KHOI_THU_THUAT = True
            mod.CHI_DIEN_BS_KHI_DANG_TRONG = True
            if cls_list:
                mod.DANH_SACH_TEN_CLS_CAN_CHUYEN = cls_list
            result = mod.chuyen_thu_thuat_sang_tieuphau()
            return "\n".join(f"{k}: {v}" for k, v in result.items())
        self.run_background("Chuyển thủ thuật sang tiểu phẫu + điền BS", task)

    def auto_dien_bs(self) -> None:
        def task():
            import dien_bs_tieuphau_t5_tu_so as mod
            source = self.current_path()
            out = output_path_for(source, "da_dien_BS")
            mod.FILE_NGUON = source
            mod.FILE_DAU_RA = out
            mod.FILE_BAO_CAO_BS = DEFAULT_REPORT_DIR / "bao_cao_dien_bs_tieuphau_T5.xlsx"
            result = mod.dien_bs_tieuphau()
            return "\n".join(f"{k}: {v}" for k, v in result.items())
        self.run_background("Điền BS tieuphau", task)

    def show_error(self, e: Exception) -> None:
        detail = traceback.format_exc()
        self.status_var.set(f"Lỗi: {e}")
        messagebox.showerror("Lỗi", f"{e}\n\nChi tiết đã in trong terminal.")
        print(detail)


if __name__ == "__main__":
    app = App()
    app.mainloop()
