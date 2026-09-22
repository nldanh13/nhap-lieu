# -*- coding: utf-8 -*-
"""Đối chiếu các dòng sheet phauthuat CÒN THIẾU PTV chính với dữ liệu OCR
(Danh sách sơ bộ trích từ ảnh sổ viết tay) — dùng khi tra TheoSo (Sổ Phẫu
Thuật) không tìm được PTV. Chỉ đề xuất, không tự ghi vào file: người dùng
bấm "Xem ảnh" để xác minh chữ viết tay đúng, rồi mới áp dụng.

Thứ tự ưu tiên nguồn PTV/phụ mổ:
1. TheoSo (Sổ Phẫu Thuật) — chạy qua bo_sung_phu_mo_t5_tu_so_phau_thuat.py.
2. Dữ liệu ảnh (module này) — chỉ dùng cho các dòng TheoSo không có kết quả.

Chạy độc lập để kiểm tra nhanh:
    python doi_chieu_phau_thuat_tu_anh.py \
        --file "PM khoa CTCH T8_NHAP_LIEU.xlsx" \
        --so-bo "Danh_sach_phau_thuat_da_chuan_hoa.xlsx"
"""
from __future__ import annotations

import argparse
import json
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from chuyen_thu_thuat_sang_tieuphau_t5 import (
    find_header_row_and_columns,
    find_sheet_case_insensitive,
    find_total_row,
    method_score,
    normalize_header,
    normalize_text,
    value_is_blank,
)
from bo_sung_phu_mo_t5_tu_so_phau_thuat import parse_date_t5, staff_to_alias


SHEET_PHAU_THUAT = "phauthuat"
SHEET_SO_BO_MAC_DINH = "Danh sách sơ bộ"

COT_PHAU_THUAT_CAN = [
    "NGÀY",
    "Họ và tên bệnh nhân",
    "Chẩn đoán và phương pháp phẫu thuật",
    "PTV chính",
    "Phụ mổ 1",
    "Phụ mổ 2",
    "Phụ mổ 3",
]

DOCTOR_HEADERS_SO_BO = [
    "BS phẫu thuật 1",
    "BS phẫu thuật 2",
    "BS phẫu thuật 3",
    "BS phẫu thuật 4",
]


def _header_map_so_bo(ws) -> dict[str, int]:
    header_row = 5
    return {
        str(cell.value).strip(): cell.column
        for cell in ws[header_row]
        if not value_is_blank(cell.value)
    }


def _find_col(headers: dict[str, int], *names: str) -> int | None:
    for name in names:
        if name in headers:
            return headers[name]
    return None


def _parse_date_so_bo(value: Any):
    """Ngày đọc từ ảnh đôi khi không chắc chắn, vd 'Trước 12/07/2026' —
    những dòng này cố ý KHÔNG dùng để đối chiếu vì có thể sai ngày thật."""
    text = str(value or "").strip()
    if not text or normalize_text(text).startswith("truoc"):
        return None
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%y"):
        try:
            return datetime.strptime(text, fmt).date()
        except Exception:
            continue
    return None


def doc_du_lieu_so_bo(so_bo_path: Path, sheet_name: str = SHEET_SO_BO_MAC_DINH) -> list[dict[str, Any]]:
    wb = load_workbook(so_bo_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        sheet_name = wb.sheetnames[0]
    ws = wb[sheet_name]
    headers = _header_map_so_bo(ws)

    col_ngay = _find_col(headers, "Ngày ghi sổ", "Ngày")
    col_ho_ten = _find_col(headers, "Họ tên người bệnh", "Họ và tên")
    col_phuong_phap = _find_col(headers, "Phương pháp phẫu thuật", "Phương pháp")
    col_chan_doan = _find_col(headers, "Chẩn đoán đọc từ sổ", "Chẩn đoán")
    col_anh = _find_col(headers, "Ảnh nguồn", "Ảnh nguồn trong ZIP")
    col_bac_si = [_find_col(headers, name) for name in DOCTOR_HEADERS_SO_BO]
    col_bac_si_chuoi = _find_col(headers, "Bác sĩ phẫu thuật", "Kíp phẫu thuật (đọc theo chuỗi)")

    if col_ho_ten is None:
        raise ValueError("Không tìm thấy cột Họ tên người bệnh trong file danh sách sơ bộ.")

    rows: list[dict[str, Any]] = []
    for row_idx in range(6, ws.max_row + 1):
        ho_ten = ws.cell(row_idx, col_ho_ten).value
        if value_is_blank(ho_ten):
            continue

        if any(col_bac_si):
            bac_si_raw = [ws.cell(row_idx, c).value if c else "-" for c in col_bac_si]
        elif col_bac_si_chuoi:
            chuoi = str(ws.cell(row_idx, col_bac_si_chuoi).value or "")
            parts = [p.strip() for p in chuoi.replace("\r", "\n").split(" - ") if p.strip()]
            bac_si_raw = (parts + ["-", "-", "-", "-"])[:4]
        else:
            bac_si_raw = ["-", "-", "-", "-"]

        rows.append(
            {
                "row": row_idx,
                "ngay": _parse_date_so_bo(ws.cell(row_idx, col_ngay).value if col_ngay else None),
                "ho_ten": str(ho_ten).strip(),
                "ho_ten_key": normalize_text(ho_ten),
                "phuong_phap": str(ws.cell(row_idx, col_phuong_phap).value or "").strip() if col_phuong_phap else "",
                "chan_doan": str(ws.cell(row_idx, col_chan_doan).value or "").strip() if col_chan_doan else "",
                "anh_nguon": str(ws.cell(row_idx, col_anh).value or "").strip() if col_anh else "",
                "bac_si_raw": bac_si_raw,
            }
        )
    return rows


def doi_chieu(file_path: Path, so_bo_path: Path, sheet_so_bo: str = SHEET_SO_BO_MAC_DINH) -> list[dict[str, Any]]:
    wb = load_workbook(file_path, data_only=True)
    sheet_name = find_sheet_case_insensitive(wb, SHEET_PHAU_THUAT)
    ws = wb[sheet_name]

    header_row, header_map = find_header_row_and_columns(ws, ["STT", *COT_PHAU_THUAT_CAN])
    col_ngay = header_map[normalize_header("NGÀY")]
    col_ho_ten = header_map[normalize_header("Họ và tên bệnh nhân")]
    col_phuong_phap = header_map[normalize_header("Chẩn đoán và phương pháp phẫu thuật")]
    col_ptv = header_map[normalize_header("PTV chính")]

    total_row = find_total_row(ws)

    so_bo_rows = doc_du_lieu_so_bo(so_bo_path, sheet_so_bo)
    index: dict[tuple, list[dict[str, Any]]] = defaultdict(list)
    for item in so_bo_rows:
        if item["ngay"] is None:
            continue
        index[(item["ngay"], item["ho_ten_key"])].append(item)

    ket_qua: list[dict[str, Any]] = []
    for row_idx in range(header_row + 1, total_row):
        if not value_is_blank(ws.cell(row_idx, col_ptv).value):
            continue  # đã có PTV chính (từ TheoSo hoặc nhập tay) — ưu tiên giữ, không đề xuất đè.

        ho_ten = ws.cell(row_idx, col_ho_ten).value
        target_date = parse_date_t5(ws.cell(row_idx, col_ngay).value)
        if value_is_blank(ho_ten) or target_date is None:
            continue

        candidates = index.get((target_date, normalize_text(ho_ten)), [])
        if not candidates:
            continue

        best = candidates[0]
        best_score = None
        if len(candidates) > 1:
            target_method = ws.cell(row_idx, col_phuong_phap).value
            scored = [(method_score(target_method, c["phuong_phap"]), c) for c in candidates]
            scored.sort(key=lambda x: x[0], reverse=True)
            best_score, best = scored[0]

        alias_list = []
        for raw in best["bac_si_raw"]:
            text = str(raw or "").strip()
            if not text or text == "-":
                alias_list.append("")
                continue
            alias = staff_to_alias(text)
            alias_list.append(alias or text)
        while len(alias_list) < 4:
            alias_list.append("")

        ket_qua.append(
            {
                "row": row_idx,
                "hoTen": str(ho_ten).strip(),
                "ngay": str(ws.cell(row_idx, col_ngay).value),
                "chanDoanPhuongPhap": ws.cell(row_idx, col_phuong_phap).value,
                "anhNguon": best["anh_nguon"],
                "soBoHoTen": best["ho_ten"],
                "soBoChanDoan": best["chan_doan"],
                "soBoPhuongPhap": best["phuong_phap"],
                "goiYPtvChinh": alias_list[0],
                "goiYPhuMo": alias_list[1:4],
                "doKhopPhuongPhap": round(best_score, 2) if best_score is not None else None,
                "soUngVienCungTen": len(candidates),
            }
        )

    return ket_qua


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Đối chiếu dòng phauthuat thiếu PTV với dữ liệu ảnh.")
    parser.add_argument("--file", required=True, type=Path)
    parser.add_argument("--so-bo", required=True, type=Path, dest="so_bo")
    parser.add_argument("--sheet-so-bo", default=SHEET_SO_BO_MAC_DINH, dest="sheet_so_bo")
    return parser


def main() -> None:
    args = build_parser().parse_args()
    ket_qua = doi_chieu(args.file, args.so_bo, args.sheet_so_bo)
    print(json.dumps({"ok": True, "goiY": ket_qua, "count": len(ket_qua)}, ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
